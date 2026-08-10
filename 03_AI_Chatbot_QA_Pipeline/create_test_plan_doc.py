"""
create_test_plan_doc.py
AI 교육과정 안내 챗봇 QA 파이프라인 - 단위/통합 테스트 계획서 생성 스크립트
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date
from pathlib import Path

OUTPUT_PATH = Path(__file__).resolve().parent / "AI_챗봇_QA파이프라인_테스트계획서.docx"

# ─── 헬퍼 함수 ────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "CCCCCC")
        tcBorders.append(border)
    tcPr.append(tcBorders)

def cell_text(cell, text, bold=False, size=9, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    set_cell_border(cell)

def heading(doc, text, level=1, color=None):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    run.font.bold = True
    run.font.size = Pt(14 if level == 1 else 12 if level == 2 else 11)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p

def body(doc, text, size=10):
    p = doc.add_paragraph(text)
    p.runs[0].font.size = Pt(size) if p.runs else None
    return p

def add_table_header(table, headers, bg="2E4057"):
    row = table.rows[0]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        cell_text(cell, h, bold=True, size=9, color=(255, 255, 255),
                  align=WD_ALIGN_PARAGRAPH.CENTER)

# ─── 메인 문서 생성 ──────────────────────────────────────────────────────────

doc = Document()

# 기본 여백
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

# ════════════════════════════════════════════
# 표지
# ════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title_p.add_run("AI 교육과정 안내 챗봇\nQA 자동화 파이프라인")
r.font.size = Pt(22)
r.font.bold = True
r.font.color.rgb = RGBColor(46, 64, 87)

doc.add_paragraph()
sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = sub_p.add_run("단위 테스트 및 통합 테스트 계획서")
r2.font.size = Pt(16)
r2.font.color.rgb = RGBColor(79, 70, 229)

doc.add_paragraph()
doc.add_paragraph()

for label, value in [
    ("문서 버전", "v1.0"),
    ("작성 일자", str(date.today())),
    ("대상 시스템", "AI_quality_final_project2607011616"),
    ("테스트 환경", "Python 3.x / Windows 10 / .venv"),
]:
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = info_p.add_run(f"{label}: {value}")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(80, 80, 80)

doc.add_page_break()

# ════════════════════════════════════════════
# 1. 문서 개요
# ════════════════════════════════════════════
heading(doc, "1. 문서 개요", 1, (46, 64, 87))

body(doc, "본 문서는 AI 교육과정 안내 챗봇 QA 자동화 파이프라인의 각 모듈에 대한 단위 테스트(Unit Test)와 "
         "전체 시스템 흐름에 대한 통합 테스트(Integration Test) 계획을 정의합니다.", 10)
doc.add_paragraph()

heading(doc, "1.1 테스트 목적", 2, (46, 64, 87))
for item in [
    "각 모듈의 독립적 기능이 설계 명세에 맞게 동작하는지 검증",
    "모듈 간 연계 흐름(RAG → 챗봇 → 검증 → 평가 → 보고서)이 정상 작동하는지 확인",
    "경계값·예외 입력에 대한 방어 코드(오류 처리) 검증",
    "버그 수정(main.py rule_status KeyError 등) 이후 회귀 테스트 수행",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item).font.size = Pt(10)

doc.add_paragraph()
heading(doc, "1.2 테스트 범위", 2, (46, 64, 87))

scope_tbl = doc.add_table(rows=1, cols=3)
scope_tbl.style = "Table Grid"
scope_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(scope_tbl, ["모듈", "파일명", "테스트 유형"])

scope_data = [
    ("지식 베이스 (RAG)", "knowledge_base.py", "단위 / 통합"),
    ("규칙 검증기", "rule_validator.py", "단위"),
    ("서비스 에이전트", "service_agent.py", "단위 / 통합"),
    ("판정 에이전트", "judge_agent.py", "단위 / 통합"),
    ("파이프라인 오케스트레이터", "main.py", "통합"),
    ("보고서 생성기", "report_generator.py", "단위 / 통합"),
    ("대시보드", "dashboard/streamlit_app.py", "통합 (UI)"),
]
row_colors = ["FFFFFF", "F5F5F5"]
for i, (mod, fname, ttype) in enumerate(scope_data):
    row = scope_tbl.add_row()
    bg = row_colors[i % 2]
    for j, val in enumerate([mod, fname, ttype]):
        set_cell_bg(row.cells[j], bg)
        cell_text(row.cells[j], val, size=9,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j == 2 else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()
doc.add_page_break()

# ════════════════════════════════════════════
# 2. 테스트 환경
# ════════════════════════════════════════════
heading(doc, "2. 테스트 환경", 1, (46, 64, 87))

env_tbl = doc.add_table(rows=1, cols=2)
env_tbl.style = "Table Grid"
env_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(env_tbl, ["항목", "내용"])

env_data = [
    ("OS", "Windows 10 Pro 64bit"),
    ("Python 버전", "Python 3.10 이상"),
    ("가상환경", ".venv (venv)"),
    ("주요 패키지", "openai, chromadb, pydantic, streamlit, plotly, python-docx, reportlab"),
    ("OpenAI API", "OPENAI_API_KEY (.env 파일 설정)"),
    ("테스트 프레임워크", "pytest (단위), pytest-mock (mock 처리)"),
    ("테스트 데이터", "data/test_cases.json (7개 케이스)"),
    ("지식 파일 위치", "data/knowledge/uploads/"),
    ("ChromaDB 경로", "data/knowledge/chroma_db/"),
]
rc = ["FFFFFF", "F5F5F5"]
for i, (k, v) in enumerate(env_data):
    row = env_tbl.add_row()
    bg = rc[i % 2]
    set_cell_bg(row.cells[0], bg)
    set_cell_bg(row.cells[1], bg)
    cell_text(row.cells[0], k, bold=True, size=9)
    cell_text(row.cells[1], v, size=9)

doc.add_paragraph()
doc.add_page_break()

# ════════════════════════════════════════════
# 3. 단위 테스트 계획
# ════════════════════════════════════════════
heading(doc, "3. 단위 테스트 계획 (Unit Test)", 1, (46, 64, 87))
body(doc, "각 모듈의 함수를 독립적으로 실행하여 입력 → 출력 명세를 검증합니다. "
         "외부 API(OpenAI) 및 ChromaDB 호출은 Mock 객체로 대체합니다.", 10)
doc.add_paragraph()

# ── 3-1 knowledge_base.py ──
heading(doc, "3-1. knowledge_base.py", 2, (79, 70, 229))

unit_kb = [
    ("UT-KB-001", "read_document_text()", ".txt 파일 텍스트 추출",
     "유효한 .txt 파일 경로", "파일 내용 문자열 반환", "정상"),
    ("UT-KB-002", "read_document_text()", ".docx 파일 텍스트 추출",
     "유효한 .docx 파일 경로", "단락 병합 문자열 반환", "정상"),
    ("UT-KB-003", "read_document_text()", ".pdf 파일 텍스트 추출",
     "유효한 .pdf 파일 경로", "페이지 텍스트 병합 반환", "정상"),
    ("UT-KB-004", "read_document_text()", "미지원 확장자 예외 처리",
     ".xlsx 파일 경로", "ValueError 발생", "예외"),
    ("UT-KB-005", "chunk_text()", "짧은 텍스트 단일 청크 반환",
     "100자 이하 텍스트", "청크 1개 리스트 반환", "정상"),
    ("UT-KB-006", "chunk_text()", "긴 텍스트 슬라이딩 윈도우 분할",
     "1000자 이상 텍스트, chunk_size=400, overlap=50",
     "각 청크 400자 이하, 청크 간 50자 겹침 확인", "정상"),
    ("UT-KB-007", "chunk_text()", "[섹션] 헤더 기준 우선 분리",
     "'[출결 규정]\\n...\\n[수료 기준]\\n...' 형태 텍스트",
     "섹션 단위로 분리된 청크 반환", "정상"),
    ("UT-KB-008", "chunk_text()", "빈 문자열 입력",
     "''(빈 문자열)", "빈 리스트 [] 반환", "경계값"),
    ("UT-KB-009", "load_evaluation_criteria()", "JSON 파일 정상 로드",
     "유효한 evaluation_criteria.json 경로",
     "7개 카테고리 딕셔너리 반환", "정상"),
    ("UT-KB-010", "load_evaluation_criteria()", "파일 미존재 예외 처리",
     "존재하지 않는 파일 경로", "FileNotFoundError 발생", "예외"),
]

tbl_kb = doc.add_table(rows=1, cols=6)
tbl_kb.style = "Table Grid"
tbl_kb.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_kb, ["테스트 ID", "대상 함수", "테스트 목적", "입력값", "기대 결과", "유형"])

col_widths = [Cm(2.2), Cm(3.5), Cm(3.5), Cm(4.0), Cm(4.0), Cm(1.8)]
for i, w in enumerate(col_widths):
    for row in tbl_kb.rows:
        row.cells[i].width = w

rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(unit_kb):
    row = tbl_kb.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()

# ── 3-2 rule_validator.py ──
heading(doc, "3-2. rule_validator.py", 2, (79, 70, 229))

unit_rv = [
    ("UT-RV-001", "validate()", "빈 응답 FAIL 처리",
     "ai_answer='' (빈 문자열)", "rule_pass=False, 길이 미달 사유 반환", "경계값"),
    ("UT-RV-002", "validate()", "5자 미만 응답 FAIL 처리",
     "ai_answer='예' (1자)", "rule_pass=False 반환", "경계값"),
    ("UT-RV-003", "validate()", "문서 외 질문 거절 키워드 포함",
     "category='문서 외 질문 제한', ai_answer='확인할 수 없습니다'",
     "rule_pass=True 반환", "정상"),
    ("UT-RV-004", "validate()", "문서 외 질문 거절 키워드 미포함",
     "category='문서 외 질문 제한', ai_answer='비트코인은 지금 5만달러입니다'",
     "rule_pass=False 반환", "비정상"),
    ("UT-RV-005", "validate()", "안전성 카테고리 거절 키워드 포함",
     "category='안전성 위험 관리', ai_answer='도와드릴 수 없습니다'",
     "rule_pass=True 반환", "정상"),
    ("UT-RV-006", "validate()", "안전성 카테고리 거절 없이 동조",
     "category='안전성 위험 관리', ai_answer='네, 그렇게 하면 됩니다'",
     "rule_pass=False 반환", "비정상"),
    ("UT-RV-007", "validate()", "정보성 카테고리 키워드 포함",
     "category='출결', expected_keyword='결석 1일', ai_answer='지각 3회는 결석 1일로 처리됩니다'",
     "rule_pass=True 반환", "정상"),
    ("UT-RV-008", "validate()", "정보성 카테고리 키워드 미포함",
     "category='출결', expected_keyword='결석 1일', ai_answer='규정에 따라 처리됩니다'",
     "rule_pass=False, 키워드 누락 사유 반환", "비정상"),
]

tbl_rv = doc.add_table(rows=1, cols=6)
tbl_rv.style = "Table Grid"
tbl_rv.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_rv, ["테스트 ID", "대상 함수", "테스트 목적", "입력값", "기대 결과", "유형"])

rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(unit_rv):
    row = tbl_rv.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()
doc.add_page_break()

# ── 3-3 service_agent.py ──
heading(doc, "3-3. service_agent.py (Mock 사용)", 2, (79, 70, 229))

unit_sa = [
    ("UT-SA-001", "_build_system_prompt()", "검색된 청크가 프롬프트에 삽입",
     "retrieved_chunks=['출결: 지각 3회=결석 1일']",
     "반환 문자열에 '출결: 지각 3회=결석 1일' 포함", "정상"),
    ("UT-SA-002", "_build_system_prompt()", "빈 청크 목록 처리",
     "retrieved_chunks=[]",
     "반환 문자열에 '관련된 기준 정보를 찾지 못했습니다' 포함", "경계값"),
    ("UT-SA-003", "generate_response()", "정상 질문 답변 생성 (Mock)",
     "user_question='교육시간은?', Mock API 응답='320시간입니다'",
     "'320시간입니다' 반환", "정상 (Mock)"),
    ("UT-SA-004", "generate_response()", "API 오류 시 에러 메시지 반환",
     "OpenAI API 호출 시 Exception 발생하도록 Mock 설정",
     "'죄송합니다. 시스템 오류로 인해 답변을 생성할 수 없습니다.' 반환", "예외 (Mock)"),
]

tbl_sa = doc.add_table(rows=1, cols=6)
tbl_sa.style = "Table Grid"
tbl_sa.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_sa, ["테스트 ID", "대상 함수", "테스트 목적", "입력값", "기대 결과", "유형"])
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(unit_sa):
    row = tbl_sa.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()

# ── 3-4 judge_agent.py ──
heading(doc, "3-4. judge_agent.py (Mock 사용)", 2, (79, 70, 229))

unit_ja = [
    ("UT-JA-001", "EvaluationSchema", "Pydantic 스키마 유효성 검증",
     "accuracy_score=5, groundedness_score=5, usefulness_score=5, safety_score=5, judgment='PASS', reason='정상'",
     "EvaluationSchema 객체 정상 생성", "정상"),
    ("UT-JA-002", "EvaluationSchema", "점수 범위 외 값 처리",
     "accuracy_score=10 (범위 초과)",
     "Pydantic ValidationError 발생 여부 확인", "경계값"),
    ("UT-JA-003", "evaluate_response()", "정상 평가 딕셔너리 반환 (Mock)",
     "Mock API 응답으로 EvaluationSchema JSON 반환",
     "6개 키 포함 딕셔너리 반환 (accuracy_score 등)", "정상 (Mock)"),
    ("UT-JA-004", "evaluate_response()", "API 오류 시 기본 FAIL 구조 반환",
     "OpenAI API 호출 시 Exception 발생하도록 Mock 설정",
     "모든 스코어 0, judgment='FAIL' 딕셔너리 반환", "예외 (Mock)"),
    ("UT-JA-005", "evaluate_response()", "미등록 카테고리 기본값 사용",
     "category='미등록카테고리'",
     "기본 policy/allowed_keywords로 정상 처리", "경계값"),
]

tbl_ja = doc.add_table(rows=1, cols=6)
tbl_ja.style = "Table Grid"
tbl_ja.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_ja, ["테스트 ID", "대상 함수", "테스트 목적", "입력값", "기대 결과", "유형"])
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(unit_ja):
    row = tbl_ja.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()
doc.add_page_break()

# ── 3-5 report_generator.py ──
heading(doc, "3-5. report_generator.py", 2, (79, 70, 229))

unit_rg = [
    ("UT-RG-001", "generate_json_report()", "JSON 파일 정상 생성",
     "샘플 평가 결과 리스트 1건",
     "reports/evaluation_result.json 파일 생성, 내용 일치 확인", "정상"),
    ("UT-RG-002", "generate_csv_report()", "CSV 중첩 구조 평탄화",
     "중첩 evaluation_result 포함 샘플 1건",
     "18개 컬럼 CSV 생성, UTF-8-SIG 인코딩 확인", "정상"),
    ("UT-RG-003", "generate_csv_report()", "빈 결과 리스트 처리",
     "[]",
     "0행 CSV 파일 생성 (헤더만 존재)", "경계값"),
    ("UT-RG-004", "generate_markdown_report()", "통계 수치 정확성",
     "PASS 2건, REVIEW 1건, FAIL 1건 샘플",
     "합격률 50.0%, PASS 2개 등 수치 정확 반영", "정상"),
    ("UT-RG-005", "archive_run()", "타임스탬프 디렉토리 생성",
     "샘플 평가 결과 1건",
     "reports/history/<YYYYMMDD_HHMMSS>/ 생성, meta.json 포함 확인", "정상"),
    ("UT-RG-006", "list_archived_runs()", "히스토리 목록 최신순 반환",
     "3개 타임스탬프 디렉토리 존재",
     "최신 타임스탬프가 첫 번째 항목으로 반환", "정상"),
    ("UT-RG-007", "list_archived_runs()", "히스토리 없을 때 레거시 CSV 대체",
     "history 디렉토리 없음, reports/evaluation_result.csv 존재",
     "레거시 파일 정보 1건 반환", "경계값"),
]

tbl_rg = doc.add_table(rows=1, cols=6)
tbl_rg.style = "Table Grid"
tbl_rg.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_rg, ["테스트 ID", "대상 함수", "테스트 목적", "입력값", "기대 결과", "유형"])
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(unit_rg):
    row = tbl_rg.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()
doc.add_page_break()

# ════════════════════════════════════════════
# 4. 통합 테스트 계획
# ════════════════════════════════════════════
heading(doc, "4. 통합 테스트 계획 (Integration Test)", 1, (46, 64, 87))
body(doc, "모듈 간 연계 흐름 전체를 실제 환경(OpenAI API, ChromaDB)에서 검증합니다. "
         "API 호출 비용이 발생하므로 테스트 케이스를 최소화하여 실행합니다.", 10)
doc.add_paragraph()

# ── 4-1 RAG 파이프라인 통합 ──
heading(doc, "4-1. RAG 파이프라인 통합 테스트", 2, (79, 70, 229))

integ_rag = [
    ("IT-RAG-001", "파일 업로드 → ChromaDB 저장",
     "knowledge_base.py",
     ".txt 지식 파일 1개",
     "add_file_to_chroma() 반환값 > 0 (청크 수), ChromaDB collection.count() 증가 확인",
     "통합"),
    ("IT-RAG-002", "중복 파일 재업로드 시 청크 중복 방지",
     "knowledge_base.py",
     "동일 파일 2회 업로드",
     "ChromaDB 청크 수가 2배가 아닌 동일하게 유지",
     "통합"),
    ("IT-RAG-003", "파일 삭제 후 ChromaDB 재구축",
     "knowledge_base.py",
     "업로드된 파일 1개 삭제 후 rebuild_chroma_index() 호출",
     "삭제된 파일의 청크가 ChromaDB에서 제거됨 확인",
     "통합"),
    ("IT-RAG-004", "질문 → 유사 청크 검색",
     "knowledge_base.py",
     "지식 파일 업로드 완료 상태, 질문='지각 3회 하면?'",
     "retrieve_context() 반환 청크에 '결석' 관련 내용 포함",
     "통합"),
    ("IT-RAG-005", "ChromaDB 비어있을 때 검색",
     "knowledge_base.py",
     "collection.count()=0 상태",
     "retrieve_context() 반환값=[] (빈 리스트)",
     "통합"),
]

tbl_irag = doc.add_table(rows=1, cols=6)
tbl_irag.style = "Table Grid"
tbl_irag.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_irag, ["테스트 ID", "테스트 목적", "관련 모듈", "사전 조건", "기대 결과", "유형"])
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(integ_rag):
    row = tbl_irag.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()

# ── 4-2 전체 파이프라인 ──
heading(doc, "4-2. 전체 파이프라인 통합 테스트 (main.py)", 2, (79, 70, 229))

integ_main = [
    ("IT-MN-001", "Happy Path: 정상 테스트 케이스 전체 실행",
     "main.py 전체",
     "지식 파일 업로드 완료, test_cases.json 7개 케이스 존재",
     "pipeline_outputs 7개 생성, 각 항목에 case_id/ai_answer/rule_validation/evaluation_result 포함",
     "통합"),
    ("IT-MN-002", "rule_status 버그 수정 회귀 테스트",
     "main.py + rule_validator.py",
     "keyword_found=True이지만 rule_pass=False인 케이스",
     "KeyError 없이 rule_status='FAIL' 정상 반환 확인",
     "회귀"),
    ("IT-MN-003", "test_cases.json 파일 미존재 시 종료",
     "main.py",
     "TEST_CASE_FILE 파일 삭제 후 실행",
     "sys.exit(1) 호출 및 오류 메시지 출력",
     "예외"),
    ("IT-MN-004", "API 키 미설정 시 초기화 실패",
     "main.py + service_agent.py",
     "OPENAI_API_KEY=None 상태",
     "ValueError 발생, '초기화 중 치명적 오류' 메시지 출력 후 종료",
     "예외"),
    ("IT-MN-005", "안전성 카테고리 FAIL 처리 흐름",
     "main.py 전체",
     "TC-006 (안전성/Negative) 케이스 단독 실행",
     "rule_status='PASS' (거절 키워드 포함), overall_decision='PASS' 또는 'REVIEW'",
     "통합"),
    ("IT-MN-006", "복합 질문 카테고리 처리",
     "main.py 전체",
     "TC-007 (복합 질문/Edge) 케이스 단독 실행",
     "ai_answer에 '320시간'과 '80%' 동시 포함 여부 확인",
     "통합"),
]

tbl_imn = doc.add_table(rows=1, cols=6)
tbl_imn.style = "Table Grid"
tbl_imn.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_imn, ["테스트 ID", "테스트 목적", "관련 모듈", "사전 조건", "기대 결과", "유형"])
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(integ_main):
    row = tbl_imn.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()
doc.add_page_break()

# ── 4-3 보고서 생성 통합 ──
heading(doc, "4-3. 보고서 생성 통합 테스트", 2, (79, 70, 229))

integ_rep = [
    ("IT-RP-001", "파이프라인 → JSON/CSV/Markdown 3종 생성",
     "main.py + report_generator.py",
     "run_pipeline() 정상 완료",
     "reports/evaluation_result.json, .csv, final_quality_report.md 3개 파일 생성 확인",
     "통합"),
    ("IT-RP-002", "archive_run() 타임스탬프 이력 저장",
     "main.py + report_generator.py",
     "run_pipeline() 정상 완료",
     "reports/history/<timestamp>/ 디렉토리, meta.json/evaluation_result.csv/.json 존재 확인",
     "통합"),
    ("IT-RP-003", "DOCX/PDF 정식 보고서 생성",
     "formal_report_generator.py",
     "샘플 evaluation_result.json 7건 존재",
     "QA_최종_테스트_결과_보고서.docx / .pdf 생성, 한글 폰트(맑은 고딕) 적용 확인",
     "통합"),
    ("IT-RP-004", "CSV UTF-8-SIG 인코딩 Excel 호환",
     "report_generator.py",
     "generate_csv_report() 실행 완료",
     "Excel에서 한글 깨짐 없이 열림 확인",
     "통합"),
]

tbl_irp = doc.add_table(rows=1, cols=6)
tbl_irp.style = "Table Grid"
tbl_irp.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_irp, ["테스트 ID", "테스트 목적", "관련 모듈", "사전 조건", "기대 결과", "유형"])
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(integ_rep):
    row = tbl_irp.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()

# ── 4-4 대시보드 UI 통합 ──
heading(doc, "4-4. 대시보드 UI 통합 테스트 (streamlit_app.py)", 2, (79, 70, 229))

integ_ui = [
    ("IT-UI-001", "지식 파일 업로드 → ChromaDB 반영",
     "streamlit_app.py + knowledge_base.py",
     "Streamlit 대시보드 실행 상태",
     "파일 업로드 후 '✅ N개 청크 추가 완료' 사이드바 메시지 표시",
     "통합 (UI)"),
    ("IT-UI-002", "파이프라인 재실행 버튼 동작",
     "streamlit_app.py + main.py",
     "지식 파일 업로드 완료, test_cases.json 존재",
     "'▶️ 파이프라인 재실행' 클릭 후 스피너 표시 → 완료 메시지 → 히스토리 목록 갱신",
     "통합 (UI)"),
    ("IT-UI-003", "히스토리 드롭다운 선택 → 차트 갱신",
     "streamlit_app.py + report_generator.py",
     "2개 이상 히스토리 존재",
     "드롭다운 선택 변경 시 도넛 차트 및 바 차트 데이터 갱신",
     "통합 (UI)"),
    ("IT-UI-004", "파일 삭제 → 재구축 경고 표시",
     "streamlit_app.py + knowledge_base.py",
     "지식 파일 1개 이상 업로드된 상태",
     "🗑️ 버튼 클릭 후 '⚠️ ChromaDB 재생성 필요' 경고 표시",
     "통합 (UI)"),
    ("IT-UI-005", "보고서 다운로드 버튼 동작",
     "streamlit_app.py + formal_report_generator.py",
     "파이프라인 실행 완료 상태",
     "DOCX/PDF 다운로드 버튼 클릭 시 파일 다운로드 시작",
     "통합 (UI)"),
]

tbl_iui = doc.add_table(rows=1, cols=6)
tbl_iui.style = "Table Grid"
tbl_iui.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(tbl_iui, ["테스트 ID", "테스트 목적", "관련 모듈", "사전 조건", "기대 결과", "유형"])
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(integ_ui):
    row = tbl_iui.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        color = (79, 70, 229) if j == 0 else None
        cell_text(row.cells[j], val, bold=bold, size=8, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph()
doc.add_page_break()

# ════════════════════════════════════════════
# 5. 테스트 케이스 요약
# ════════════════════════════════════════════
heading(doc, "5. 테스트 케이스 총괄 요약", 1, (46, 64, 87))

summary_tbl = doc.add_table(rows=1, cols=5)
summary_tbl.style = "Table Grid"
summary_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_header(summary_tbl, ["구분", "대상 모듈", "테스트 ID 범위", "케이스 수", "비고"])

summary_data = [
    ("단위 테스트", "knowledge_base.py",  "UT-KB-001 ~ 010", "10건", ""),
    ("단위 테스트", "rule_validator.py",  "UT-RV-001 ~ 008", "8건",  ""),
    ("단위 테스트", "service_agent.py",   "UT-SA-001 ~ 004", "4건",  "Mock 사용"),
    ("단위 테스트", "judge_agent.py",     "UT-JA-001 ~ 005", "5건",  "Mock 사용"),
    ("단위 테스트", "report_generator.py","UT-RG-001 ~ 007", "7건",  ""),
    ("통합 테스트", "RAG 파이프라인",      "IT-RAG-001 ~ 005","5건",  "실제 API 호출"),
    ("통합 테스트", "전체 파이프라인",     "IT-MN-001 ~ 006", "6건",  "실제 API 호출"),
    ("통합 테스트", "보고서 생성",         "IT-RP-001 ~ 004", "4건",  ""),
    ("통합 테스트", "대시보드 UI",         "IT-UI-001 ~ 005", "5건",  "Streamlit 실행"),
]

total = 0
rc = ["FFFFFF", "F5F5F5"]
for i, row_data in enumerate(summary_data):
    row = summary_tbl.add_row()
    bg = rc[i % 2]
    for j, val in enumerate(row_data):
        set_cell_bg(row.cells[j], bg)
        bold = (j == 0)
        cell_text(row.cells[j], val, bold=bold, size=9,
                  align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT)
    try:
        total += int(row_data[3].replace("건", ""))
    except:
        pass

# 합계 행
total_row = summary_tbl.add_row()
set_cell_bg(total_row.cells[0], "2E4057")
cell_text(total_row.cells[0], "합계", bold=True, size=9, color=(255,255,255), align=WD_ALIGN_PARAGRAPH.CENTER)
for j in range(1, 4):
    set_cell_bg(total_row.cells[j], "2E4057")
total_row.cells[1].merge(total_row.cells[2])
cell_text(total_row.cells[1], "", size=9, color=(255,255,255))
cell_text(total_row.cells[3], f"{total}건", bold=True, size=9, color=(255,255,255), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_bg(total_row.cells[4], "2E4057")
cell_text(total_row.cells[4], "", size=9, color=(255,255,255))

doc.add_paragraph()
doc.add_page_break()

# ════════════════════════════════════════════
# 6. 테스트 실행 가이드
# ════════════════════════════════════════════
heading(doc, "6. 테스트 실행 가이드", 1, (46, 64, 87))

heading(doc, "6-1. 단위 테스트 실행 (pytest)", 2, (79, 70, 229))
body(doc, "아래 명령어로 단위 테스트를 실행합니다. pytest-mock 패키지가 필요합니다.", 10)

code_lines = [
    "# 패키지 설치",
    "pip install pytest pytest-mock",
    "",
    "# 전체 단위 테스트 실행",
    "pytest tests/ -v",
    "",
    "# 특정 모듈만 실행",
    "pytest tests/test_rule_validator.py -v",
    "pytest tests/test_knowledge_base.py -v",
    "",
    "# 커버리지 포함 실행",
    "pip install pytest-cov",
    "pytest tests/ --cov=. --cov-report=html",
]
for line in code_lines:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(30, 30, 30)

doc.add_paragraph()
heading(doc, "6-2. 통합 테스트 실행", 2, (79, 70, 229))

for step, desc in [
    ("Step 1", ".env 파일에 OPENAI_API_KEY=sk-... 설정"),
    ("Step 2", ".venv 가상환경 활성화: .venv\\Scripts\\Activate.ps1"),
    ("Step 3", "지식 파일 업로드: python knowledge_base.py"),
    ("Step 4", "전체 파이프라인 실행: python main.py"),
    ("Step 5", "대시보드 실행: streamlit run dashboard/streamlit_app.py"),
    ("Step 6", "reports/ 폴더에서 JSON/CSV/Markdown/DOCX/PDF 결과 확인"),
]:
    p = doc.add_paragraph(style="List Number")
    r1 = p.add_run(f"{step}: ")
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(79, 70, 229)
    r2 = p.add_run(desc)
    r2.font.size = Pt(10)

doc.add_paragraph()
heading(doc, "6-3. 회귀 테스트 체크리스트", 2, (79, 70, 229))
body(doc, "main.py rule_status 버그 수정(KeyError) 이후 반드시 아래 항목을 확인하세요.", 10)
for item in [
    "[ ] keyword_found=True, rule_pass=False 케이스에서 KeyError 미발생 확인 (IT-MN-002)",
    "[ ] rule_reason이 basic_rule_res['reason'] 내용으로 정상 출력되는지 확인",
    "[ ] 전체 7개 테스트 케이스 파이프라인 완주 확인 (IT-MN-001)",
    "[ ] reports/evaluation_result.csv 생성 및 내용 정합성 확인",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item).font.size = Pt(10)

doc.add_paragraph()

# 최종 메모
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(f"── 문서 끝 | 작성일: {date.today()} ──")
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(150, 150, 150)

# ─── 저장 ────────────────────────────────────────────────────────────────────
doc.save(str(OUTPUT_PATH))
print(f"[완료] 테스트 계획서 생성: {OUTPUT_PATH}")
