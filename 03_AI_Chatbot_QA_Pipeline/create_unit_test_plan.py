"""
create_unit_test_plan.py  — 단위 테스트 계획서 Word 생성
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date
from pathlib import Path

OUTPUT = Path(__file__).resolve().parent / "AI_챗봇_단위테스트_계획서.docx"

# ── 헬퍼 ──────────────────────────────────────────────────────────────────────

def shd(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_color)
    tcPr.append(s)

def border(cell, color="CCCCCC"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcB = OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        tcB.append(b)
    tcPr.append(tcB)

def cw(cell, text, bold=False, sz=9, color=None,
        align=WD_ALIGN_PARAGRAPH.LEFT, bg=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if bg:  shd(cell, bg)
    border(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(sz)
    r.font.bold = bold
    if color: r.font.color.rgb = RGBColor(*color)

def head_row(tbl, cols, bg="1E3A5F"):
    row = tbl.rows[0]
    for i, h in enumerate(cols):
        cw(row.cells[i], h, bold=True, sz=9, color=(255,255,255),
           align=WD_ALIGN_PARAGRAPH.CENTER, bg=bg)

def data_rows(tbl, data, stripe=("FFFFFF","F4F6FB"), id_col=0):
    for i, rd in enumerate(data):
        row = tbl.add_row()
        bg = stripe[i % 2]
        for j, val in enumerate(rd):
            is_id = (j == id_col)
            cw(row.cells[j], val, bold=is_id, sz=8,
               color=(30, 90, 200) if is_id else None,
               align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, len(rd)-1) else WD_ALIGN_PARAGRAPH.LEFT,
               bg=bg)

def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(30, 58, 95)
    return p

def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(79, 70, 229)
    return p

def body(doc, text, sz=10):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(sz)
    return p

def bullet(doc, items, sz=10):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item).font.size = Pt(sz)

def make_table(doc, cols):
    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    return tbl

# ── 문서 ──────────────────────────────────────────────────────────────────────

doc = Document()
for sec in doc.sections:
    sec.top_margin    = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin   = Cm(3.0)
    sec.right_margin  = Cm(2.0)

# ── 표지 ──────────────────────────────────────────────────────────────────────
for _ in range(4): doc.add_paragraph()

tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run("AI 교육과정 안내 챗봇\nQA 자동화 파이프라인")
r.font.size = Pt(22); r.font.bold = True
r.font.color.rgb = RGBColor(30, 58, 95)

doc.add_paragraph()
sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = sp.add_run("단위 테스트 계획서  (Unit Test Plan)")
r2.font.size = Pt(16); r2.font.color.rgb = RGBColor(79, 70, 229)

doc.add_paragraph(); doc.add_paragraph()

ct = doc.add_table(rows=4, cols=2)
ct.style = "Table Grid"
ct.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_info = [
    ("문서 구분",   "단위 테스트 계획서"),
    ("문서 버전",   "v1.0"),
    ("작성 일자",   str(date.today())),
    ("테스트 환경", "Python 3.x / pytest / pytest-mock / .venv"),
]
for i,(k,v) in enumerate(cover_info):
    cw(ct.rows[i].cells[0], k, bold=True, sz=10, bg="1E3A5F", color=(255,255,255), align=WD_ALIGN_PARAGRAPH.CENTER)
    cw(ct.rows[i].cells[1], v, sz=10, bg="F4F6FB")

doc.add_page_break()

# ── 1. 개요 ──────────────────────────────────────────────────────────────────
h1(doc, "1. 개요")
body(doc,
    "본 문서는 AI 교육과정 안내 챗봇 QA 파이프라인의 각 모듈을 독립적으로 검증하는 "
    "단위 테스트(Unit Test) 계획을 정의합니다. 외부 API(OpenAI) 및 ChromaDB 호출은 "
    "pytest-mock을 이용해 Mock 객체로 대체하여 비용 없이 로직만 검증합니다.")

doc.add_paragraph()
h2(doc, "1-1. 테스트 목적")
bullet(doc, [
    "각 함수·메서드가 설계 명세(입력 → 출력)에 따라 정확히 동작하는지 검증",
    "경계값(빈 문자열, 파일 미존재 등) 및 예외 입력에 대한 방어 코드 확인",
    "main.py KeyError 버그 수정 이후 rule_validator 연동 회귀 검증",
    "외부 의존성(OpenAI API, ChromaDB) 없이 빠른 피드백 확보",
])

doc.add_paragraph()
h2(doc, "1-2. 테스트 범위")

t = make_table(doc, ["모듈","파일명","테스트 ID 범위","케이스 수"])
head_row(t, ["모듈","파일명","테스트 ID 범위","케이스 수"])
data_rows(t, [
    ("지식 베이스 (RAG)", "knowledge_base.py", "UT-KB-001 ~ 010", "10건"),
    ("규칙 검증기",       "rule_validator.py",  "UT-RV-001 ~ 008", "8건"),
    ("서비스 에이전트",   "service_agent.py",   "UT-SA-001 ~ 004", "4건 (Mock)"),
    ("판정 에이전트",     "judge_agent.py",     "UT-JA-001 ~ 005", "5건 (Mock)"),
    ("보고서 생성기",     "report_generator.py","UT-RG-001 ~ 007", "7건"),
    ("",                 "합  계",              "",               "34건"),
])

doc.add_paragraph()
h2(doc, "1-3. 실행 환경 및 명령어")
for line in [
    "# 패키지 설치 (최초 1회)",
    "pip install pytest pytest-mock pytest-cov",
    "",
    "# 전체 단위 테스트 실행",
    "pytest tests/ -v",
    "",
    "# 특정 모듈만 실행",
    "pytest tests/test_rule_validator.py -v",
    "pytest tests/test_knowledge_base.py -v",
    "",
    "# 커버리지 HTML 보고서",
    "pytest tests/ --cov=. --cov-report=html",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(30, 30, 30)

doc.add_page_break()

# ── 2. knowledge_base.py ─────────────────────────────────────────────────────
h1(doc, "2. knowledge_base.py 단위 테스트")
body(doc, "파일 텍스트 추출, 청크 분할, ChromaDB 연동, 평가 기준 로드 함수를 검증합니다.")
doc.add_paragraph()

COLS_UNIT = ["테스트 ID","대상 함수","테스트 목적","입력값","기대 결과","유형"]

t_kb = make_table(doc, COLS_UNIT)
head_row(t_kb, COLS_UNIT)
data_rows(t_kb, [
    ("UT-KB-001","read_document_text()","`.txt` 파일 텍스트 추출",
     "유효한 .txt 파일 경로","파일 내용 문자열 반환","정상"),
    ("UT-KB-002","read_document_text()","`.docx` 파일 텍스트 추출",
     "유효한 .docx 파일 경로","단락 병합 문자열 반환","정상"),
    ("UT-KB-003","read_document_text()","`.pdf` 파일 텍스트 추출",
     "유효한 .pdf 파일 경로","페이지 텍스트 병합 반환","정상"),
    ("UT-KB-004","read_document_text()","미지원 확장자 예외 처리",
     ".xlsx 파일 경로","ValueError 발생","예외"),
    ("UT-KB-005","chunk_text()","짧은 텍스트 단일 청크",
     "100자 이하 텍스트","청크 1개 리스트 반환","정상"),
    ("UT-KB-006","chunk_text()","긴 텍스트 슬라이딩 윈도우 분할",
     "1000자 이상 텍스트, chunk_size=400, overlap=50",
     "각 청크 ≤400자, 청크 간 50자 겹침 확인","정상"),
    ("UT-KB-007","chunk_text()","[섹션] 헤더 기준 우선 분리",
     "'[출결 규정]\\n...\\n[수료 기준]\\n...'","섹션 단위 청크 반환","정상"),
    ("UT-KB-008","chunk_text()","빈 문자열 입력",
     "'' (빈 문자열)","빈 리스트 [] 반환","경계값"),
    ("UT-KB-009","load_evaluation_criteria()","JSON 정상 로드",
     "유효한 evaluation_criteria.json 경로",
     "7개 카테고리 딕셔너리 반환","정상"),
    ("UT-KB-010","load_evaluation_criteria()","파일 미존재 예외",
     "존재하지 않는 파일 경로","FileNotFoundError 발생","예외"),
])

doc.add_paragraph()
h2(doc, "pytest 예시 코드 — knowledge_base.py")
for line in [
    "def test_chunk_text_short():",
    "    result = chunk_text('안녕하세요')",
    "    assert len(result) == 1",
    "    assert result[0] == '안녕하세요'",
    "",
    "def test_chunk_text_empty():",
    "    assert chunk_text('') == []",
    "",
    "def test_read_unsupported_ext(tmp_path):",
    "    f = tmp_path / 'test.xlsx'",
    "    f.write_text('data')",
    "    with pytest.raises(ValueError):",
    "        read_document_text(f)",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"; r.font.size = Pt(9)

doc.add_page_break()

# ── 3. rule_validator.py ─────────────────────────────────────────────────────
h1(doc, "3. rule_validator.py 단위 테스트")
body(doc, "카테고리별 검증 분기(거절 키워드 / 필수 키워드 / 길이 체크)를 검증합니다.")
doc.add_paragraph()

t_rv = make_table(doc, COLS_UNIT)
head_row(t_rv, COLS_UNIT)
data_rows(t_rv, [
    ("UT-RV-001","validate()","빈 응답 FAIL",
     "ai_answer='' (빈 문자열)",
     "rule_pass=False, 길이 미달 사유 반환","경계값"),
    ("UT-RV-002","validate()","5자 미만 응답 FAIL",
     "ai_answer='예' (1자)","rule_pass=False 반환","경계값"),
    ("UT-RV-003","validate()","문서 외 질문 — 거절 키워드 포함",
     "category='문서 외 질문 제한', ai_answer='확인할 수 없습니다'",
     "rule_pass=True 반환","정상"),
    ("UT-RV-004","validate()","문서 외 질문 — 거절 키워드 없음",
     "category='문서 외 질문 제한', ai_answer='비트코인은 5만달러입니다'",
     "rule_pass=False 반환","비정상"),
    ("UT-RV-005","validate()","안전성 — 거절 키워드 포함",
     "category='안전성 위험 관리', ai_answer='도와드릴 수 없습니다'",
     "rule_pass=True 반환","정상"),
    ("UT-RV-006","validate()","안전성 — 거절 없이 동조",
     "category='안전성 위험 관리', ai_answer='네, 그렇게 하면 됩니다'",
     "rule_pass=False 반환","비정상"),
    ("UT-RV-007","validate()","정보성 카테고리 키워드 포함",
     "category='출결', expected_keyword='결석 1일', ai_answer='지각 3회는 결석 1일로 처리됩니다'",
     "rule_pass=True 반환","정상"),
    ("UT-RV-008","validate()","정보성 카테고리 키워드 미포함",
     "category='출결', expected_keyword='결석 1일', ai_answer='규정에 따라 처리됩니다'",
     "rule_pass=False, 키워드 누락 사유 반환","비정상"),
])

doc.add_paragraph()
h2(doc, "pytest 예시 코드 — rule_validator.py")
for line in [
    "from rule_validator import RuleValidator",
    "rv = RuleValidator()",
    "",
    "def test_empty_answer():",
    "    r = rv.validate({'category':'출결','expected_keyword':'결석 1일'}, '')",
    "    assert r['rule_pass'] == False",
    "",
    "def test_rejection_keyword_present():",
    "    case = {'category':'문서 외 질문 제한','expected_keyword':''}",
    "    r = rv.validate(case, '확인할 수 없습니다')",
    "    assert r['rule_pass'] == True",
    "",
    "def test_keyword_missing():",
    "    case = {'category':'출결','expected_keyword':'결석 1일'}",
    "    r = rv.validate(case, '규정에 따라 처리됩니다')",
    "    assert r['rule_pass'] == False",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"; r.font.size = Pt(9)

doc.add_page_break()

# ── 4. service_agent.py ───────────────────────────────────────────────────────
h1(doc, "4. service_agent.py 단위 테스트 (Mock)")
body(doc, "OpenAI API 및 ChromaDB를 Mock으로 대체하여 시스템 프롬프트 구성과 오류 처리 로직을 검증합니다.")
doc.add_paragraph()

t_sa = make_table(doc, COLS_UNIT)
head_row(t_sa, COLS_UNIT)
data_rows(t_sa, [
    ("UT-SA-001","_build_system_prompt()","검색 청크가 프롬프트에 삽입",
     "retrieved_chunks=['출결: 지각 3회=결석 1일']",
     "반환 문자열에 해당 내용 포함 확인","정상 (Mock)"),
    ("UT-SA-002","_build_system_prompt()","빈 청크 목록 처리",
     "retrieved_chunks=[]",
     "'관련된 기준 정보를 찾지 못했습니다' 포함","경계값"),
    ("UT-SA-003","generate_response()","정상 답변 생성",
     "Mock API 응답='320시간입니다', 질문='교육시간은?'",
     "'320시간입니다' 반환","정상 (Mock)"),
    ("UT-SA-004","generate_response()","API 오류 → 에러 메시지 반환",
     "OpenAI API 호출 시 Exception 발생 Mock",
     "'죄송합니다. 시스템 오류로 인해 답변을 생성할 수 없습니다.' 반환","예외 (Mock)"),
])

doc.add_paragraph()
h2(doc, "pytest 예시 코드 — service_agent.py")
for line in [
    "from unittest.mock import patch, MagicMock",
    "from service_agent import ServiceAgent",
    "",
    "@patch('service_agent.get_chroma_collection')",
    "@patch('service_agent.OpenAI')",
    "def test_generate_response_api_error(mock_openai, mock_chroma):",
    "    mock_openai.return_value.chat.completions.create.side_effect = Exception('API Error')",
    "    agent = ServiceAgent()",
    "    result = agent.generate_response('테스트 질문')",
    "    assert '시스템 오류' in result",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"; r.font.size = Pt(9)

doc.add_page_break()

# ── 5. judge_agent.py ─────────────────────────────────────────────────────────
h1(doc, "5. judge_agent.py 단위 테스트 (Mock)")
body(doc, "Pydantic EvaluationSchema 유효성과 API 오류 시 기본 FAIL 구조 반환 로직을 검증합니다.")
doc.add_paragraph()

t_ja = make_table(doc, COLS_UNIT)
head_row(t_ja, COLS_UNIT)
data_rows(t_ja, [
    ("UT-JA-001","EvaluationSchema","유효한 스키마 생성",
     "모든 필드 정상 값 입력",
     "EvaluationSchema 객체 정상 생성","정상"),
    ("UT-JA-002","EvaluationSchema","필수 필드 누락 시 ValidationError",
     "judgment 필드 제외 입력",
     "pydantic.ValidationError 발생","예외"),
    ("UT-JA-003","evaluate_response()","정상 평가 딕셔너리 반환",
     "Mock API → EvaluationSchema 구조 반환",
     "6개 키 포함 딕셔너리 반환","정상 (Mock)"),
    ("UT-JA-004","evaluate_response()","API 오류 → 기본 FAIL 반환",
     "OpenAI API Exception Mock",
     "모든 스코어=0, judgment='FAIL' 반환","예외 (Mock)"),
    ("UT-JA-005","evaluate_response()","미등록 카테고리 기본값 사용",
     "category='미등록카테고리'",
     "기본 policy로 정상 처리 (오류 없음)","경계값"),
])

doc.add_paragraph()
h2(doc, "pytest 예시 코드 — judge_agent.py")
for line in [
    "from judge_agent import EvaluationSchema",
    "from pydantic import ValidationError",
    "",
    "def test_schema_valid():",
    "    s = EvaluationSchema(accuracy_score=5, groundedness_score=5,",
    "                         usefulness_score=5, safety_score=5,",
    "                         judgment='PASS', reason='정상')",
    "    assert s.judgment == 'PASS'",
    "",
    "def test_schema_missing_field():",
    "    with pytest.raises(ValidationError):",
    "        EvaluationSchema(accuracy_score=5)  # 필수 필드 누락",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"; r.font.size = Pt(9)

doc.add_page_break()

# ── 6. report_generator.py ───────────────────────────────────────────────────
h1(doc, "6. report_generator.py 단위 테스트")
body(doc, "JSON·CSV·Markdown 파일 생성, 통계 수치 계산 정확성, 타임스탬프 아카이브 동작을 검증합니다.")
doc.add_paragraph()

t_rg = make_table(doc, COLS_UNIT)
head_row(t_rg, COLS_UNIT)
data_rows(t_rg, [
    ("UT-RG-001","generate_json_report()","JSON 파일 정상 생성",
     "샘플 평가 결과 1건",
     "reports/evaluation_result.json 생성, 내용 일치","정상"),
    ("UT-RG-002","generate_csv_report()","CSV 평탄화 및 인코딩",
     "중첩 evaluation_result 포함 샘플 1건",
     "UTF-8-SIG 인코딩 CSV, 19개 컬럼 정확 확인","정상"),
    ("UT-RG-003","generate_csv_report()","빈 결과 리스트",
     "[]",
     "0행 CSV 생성 (헤더만 존재)","경계값"),
    ("UT-RG-004","generate_markdown_report()","통계 수치 정확성",
     "PASS 2, REVIEW 1, FAIL 1 샘플",
     "합격률 50.0% 및 건수 정확 반영","정상"),
    ("UT-RG-005","archive_run()","타임스탬프 디렉토리 생성",
     "샘플 평가 결과 1건",
     "reports/history/<YYYYMMDD_HHMMSS>/ + meta.json 생성","정상"),
    ("UT-RG-006","list_archived_runs()","히스토리 최신순 반환",
     "3개 타임스탬프 디렉토리 존재",
     "최신 항목이 index[0]으로 반환","정상"),
    ("UT-RG-007","list_archived_runs()","히스토리 없을 때 레거시 대체",
     "history 없음, reports/evaluation_result.csv 존재",
     "레거시 파일 정보 1건 반환","경계값"),
])

doc.add_paragraph()
h2(doc, "pytest 예시 코드 — report_generator.py")
for line in [
    "import json, pytest",
    "from report_generator import ReportGenerator",
    "",
    "SAMPLE = [{",
    "  'case_id':'TC-001', 'category':'정확성',",
    "  'test_type':'Happy', 'user_question':'Q', 'ai_answer':'A',",
    "  'rule_validation':{'keyword_found':True,'rule_status':'PASS','rule_reason':'OK'},",
    "  'evaluation_result':{",
    "    'accuracy':{'score':5,'reason':'ok'},'groundedness':{'score':5,'reason':'ok'},",
    "    'helpfulness':{'score':5,'reason':'ok'},'safety':{'score':5,'reason':'ok'},",
    "    'overall_decision':'PASS','summary':'good'}}]",
    "",
    "def test_json_report(tmp_path):",
    "    rg = ReportGenerator()",
    "    rg.output_dir = tmp_path",
    "    path = rg.generate_json_report(SAMPLE)",
    "    data = json.loads(path.read_text(encoding='utf-8'))",
    "    assert data[0]['case_id'] == 'TC-001'",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"; r.font.size = Pt(9)

doc.add_paragraph()

# ── 7. 통과 기준 ─────────────────────────────────────────────────────────────
doc.add_page_break()
h1(doc, "7. 테스트 통과 기준 및 총괄")

doc.add_paragraph()
h2(doc, "7-1. 통과 기준")
t_pass = make_table(doc, ["항목","기준"])
head_row(t_pass, ["항목","기준"])
data_rows(t_pass, [
    ("최소 통과율",     "단위 테스트 전체의 95% 이상 통과"),
    ("예외 처리 테스트","Exception 발생 케이스 100% 통과 필수"),
    ("경계값 테스트",   "빈 입력·미존재 파일 등 전 경계값 케이스 통과"),
    ("회귀 테스트",     "KeyError 버그 수정 후 UT-RV-007~008 반드시 재실행"),
    ("코드 커버리지",   "rule_validator.py, report_generator.py 90% 이상 권장"),
])

doc.add_paragraph()
h2(doc, "7-2. 테스트 케이스 총괄")
t_sum = make_table(doc, ["모듈","케이스 수","Mock 사용","비고"])
head_row(t_sum, ["모듈","케이스 수","Mock 사용","비고"])
data_rows(t_sum, [
    ("knowledge_base.py",  "10건","X",       ""),
    ("rule_validator.py",  "8건", "X",       "회귀 포함"),
    ("service_agent.py",   "4건", "O",       "OpenAI + ChromaDB Mock"),
    ("judge_agent.py",     "5건", "O",       "OpenAI Mock"),
    ("report_generator.py","7건", "X (파일)","tmp_path 사용"),
    ("합  계",             "34건","",        ""),
])

# 끝
doc.add_paragraph()
ep = doc.add_paragraph()
ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ep.add_run(f"── 문서 끝  |  작성일: {date.today()} ──")
r.font.size = Pt(9); r.font.color.rgb = RGBColor(160,160,160)

doc.save(str(OUTPUT))
print(f"[완료] {OUTPUT}")
