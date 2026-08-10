"""
create_integration_test_plan.py  — 통합 테스트 계획서 Word 생성
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date
from pathlib import Path

OUTPUT = Path(__file__).resolve().parent / "AI_챗봇_통합테스트_계획서.docx"

# ── 헬퍼 ──────────────────────────────────────────────────────────────────────

def shd(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_color)
    tcPr.append(s)

def border(cell, color="CCCCCC"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcB = OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0"); b.set(qn("w:color"), color)
        tcB.append(b)
    tcPr.append(tcB)

def cw(cell, text, bold=False, sz=9, color=None,
        align=WD_ALIGN_PARAGRAPH.LEFT, bg=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if bg:  shd(cell, bg)
    border(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(sz); r.font.bold = bold
    if color: r.font.color.rgb = RGBColor(*color)

def head_row(tbl, cols, bg="0F3460"):
    row = tbl.rows[0]
    for i, h in enumerate(cols):
        cw(row.cells[i], h, bold=True, sz=9, color=(255,255,255),
           align=WD_ALIGN_PARAGRAPH.CENTER, bg=bg)

def data_rows(tbl, data, stripe=("FFFFFF","EEF2FF"), id_col=0):
    for i, rd in enumerate(data):
        row = tbl.add_row()
        bg = stripe[i % 2]
        for j, val in enumerate(rd):
            is_id = (j == id_col)
            cw(row.cells[j], val, bold=is_id, sz=8,
               color=(15, 52, 96) if is_id else None,
               align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, len(rd)-1) else WD_ALIGN_PARAGRAPH.LEFT,
               bg=bg)

def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text)
    r.font.bold = True; r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 52, 96)
    return p

def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(text)
    r.font.bold = True; r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(22, 101, 162)
    return p

def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(text)
    r.font.bold = True; r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(60, 60, 60)

def body(doc, text, sz=10):
    p = doc.add_paragraph()
    p.add_run(text).font.size = Pt(sz)

def bullet(doc, items, sz=10):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item).font.size = Pt(sz)

def numbered(doc, items, sz=10):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item).font.size = Pt(sz)

def make_table(doc, cols):
    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    return tbl

COLS_IT = ["테스트 ID","테스트 목적","관련 모듈","사전 조건","기대 결과","유형"]

# ── 문서 ──────────────────────────────────────────────────────────────────────

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.0)

# ── 표지 ──────────────────────────────────────────────────────────────────────
for _ in range(4): doc.add_paragraph()

tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run("AI 교육과정 안내 챗봇\nQA 자동화 파이프라인")
r.font.size = Pt(22); r.font.bold = True
r.font.color.rgb = RGBColor(15, 52, 96)

doc.add_paragraph()
sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = sp.add_run("통합 테스트 계획서  (Integration Test Plan)")
r2.font.size = Pt(16); r2.font.color.rgb = RGBColor(22, 101, 162); r2.font.italic = True

doc.add_paragraph(); doc.add_paragraph()

ct = doc.add_table(rows=5, cols=2)
ct.style = "Table Grid"
ct.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_info = [
    ("문서 구분",   "통합 테스트 계획서"),
    ("문서 버전",   "v1.0"),
    ("작성 일자",   str(date.today())),
    ("테스트 환경", "Python 3.x / OpenAI API / ChromaDB / Streamlit"),
    ("주의 사항",   "실제 OpenAI API 호출 포함 → API 비용 발생"),
]
for i,(k,v) in enumerate(cover_info):
    cw(ct.rows[i].cells[0], k, bold=True, sz=10, bg="0F3460", color=(255,255,255), align=WD_ALIGN_PARAGRAPH.CENTER)
    cw(ct.rows[i].cells[1], v, sz=10, bg="EEF2FF")

doc.add_page_break()

# ── 1. 개요 ──────────────────────────────────────────────────────────────────
h1(doc, "1. 개요")
body(doc,
    "본 문서는 AI 교육과정 안내 챗봇 QA 파이프라인의 모듈 간 연계 흐름을 실제 환경에서 검증하는 "
    "통합 테스트(Integration Test) 계획을 정의합니다. "
    "RAG 파이프라인, 전체 파이프라인 실행, 보고서 생성, 대시보드 UI의 4개 영역을 다룹니다.")

doc.add_paragraph()
h2(doc, "1-1. 테스트 목적")
bullet(doc, [
    "파일 업로드 → 청크 분할 → ChromaDB 저장 → 유사도 검색 RAG 흐름 검증",
    "ServiceAgent → RuleValidator → JudgeAgent → ReportGenerator 전체 파이프라인 연계 검증",
    "main.py KeyError 버그 수정 후 전체 파이프라인 회귀 테스트",
    "보고서 3종(JSON/CSV/Markdown) 및 DOCX/PDF 정식 보고서 생성 검증",
    "Streamlit 대시보드의 업로드·실행·히스토리·다운로드 UI 흐름 검증",
])

doc.add_paragraph()
h2(doc, "1-2. 테스트 범위")
t = make_table(doc, ["영역","관련 모듈","테스트 ID 범위","케이스 수"])
head_row(t, ["영역","관련 모듈","테스트 ID 범위","케이스 수"])
data_rows(t, [
    ("RAG 파이프라인",    "knowledge_base.py",               "IT-RAG-001 ~ 005","5건"),
    ("전체 파이프라인",   "main.py + 전 모듈",               "IT-MN-001 ~ 006", "6건"),
    ("보고서 생성",       "report_generator.py + formal_...", "IT-RP-001 ~ 004", "4건"),
    ("대시보드 UI",       "dashboard/streamlit_app.py",       "IT-UI-001 ~ 005", "5건"),
    ("",                  "합  계",                           "",                "20건"),
])

doc.add_paragraph()
h2(doc, "1-3. 테스트 실행 순서 (전체)")
numbered(doc, [
    ".env 파일에 OPENAI_API_KEY 설정 확인",
    ".venv 가상환경 활성화:  .venv\\Scripts\\Activate.ps1",
    "지식 파일 ChromaDB 등록:  python knowledge_base.py  (IT-RAG 영역)",
    "전체 파이프라인 실행:  python main.py  (IT-MN 영역)",
    "보고서 파일 생성 확인:  reports/ 폴더  (IT-RP 영역)",
    "대시보드 실행:  streamlit run dashboard/streamlit_app.py  (IT-UI 영역)",
])

doc.add_paragraph()
h2(doc, "1-4. 사전 조건 (공통)")
bullet(doc, [
    "OPENAI_API_KEY가 .env 파일에 설정되어 있어야 함",
    ".venv 가상환경에 requirements.txt 패키지 설치 완료",
    "data/knowledge/uploads/ 폴더에 지식 파일 1개 이상 존재",
    "data/test_cases.json 파일 존재",
])

doc.add_page_break()

# ── 2. RAG 파이프라인 통합 테스트 ────────────────────────────────────────────
h1(doc, "2. RAG 파이프라인 통합 테스트")
body(doc,
    "지식 파일 업로드부터 ChromaDB 벡터 검색까지 전체 RAG 흐름을 검증합니다. "
    "실제 OpenAI Embedding API(text-embedding-3-small)를 호출합니다.")
doc.add_paragraph()

t_rag = make_table(doc, COLS_IT)
head_row(t_rag, COLS_IT)
data_rows(t_rag, [
    ("IT-RAG-001","파일 업로드 → ChromaDB 저장",
     "knowledge_base.py",
     ".txt 지식 파일 1개 준비",
     "add_file_to_chroma() 반환값 > 0, ChromaDB collection.count() 증가 확인",
     "통합"),
    ("IT-RAG-002","중복 업로드 시 청크 중복 방지",
     "knowledge_base.py",
     "동일 파일 2회 연속 업로드",
     "ChromaDB 청크 수가 2배가 아닌 동일하게 유지",
     "통합"),
    ("IT-RAG-003","파일 삭제 후 ChromaDB 재구축",
     "knowledge_base.py",
     "파일 1개 삭제 후 rebuild_chroma_index() 호출",
     "삭제된 파일의 청크가 ChromaDB에서 제거됨 확인",
     "통합"),
    ("IT-RAG-004","질문 → 유사 청크 검색",
     "knowledge_base.py",
     "지식 파일 업로드 완료, 질문='지각 3회 하면?'",
     "retrieve_context() 반환 청크에 '결석' 관련 내용 포함",
     "통합"),
    ("IT-RAG-005","ChromaDB 비어있을 때 검색",
     "knowledge_base.py",
     "collection.count()=0 상태",
     "retrieve_context() 반환값=[] (빈 리스트, 오류 없음)",
     "경계값"),
])

doc.add_paragraph()
h2(doc, "확인 방법")
bullet(doc, [
    "python knowledge_base.py 실행 후 콘솔에서 '청크를 ChromaDB에 추가' 메시지 확인",
    "ChromaDB 폴더(data/knowledge/chroma_db/) 파일 생성 여부 확인",
    "IT-RAG-004: 콘솔 출력에서 '검색된 청크' 내용에 출결 키워드 포함 여부 육안 검증",
])

doc.add_page_break()

# ── 3. 전체 파이프라인 통합 테스트 ───────────────────────────────────────────
h1(doc, "3. 전체 파이프라인 통합 테스트 (main.py)")
body(doc,
    "ServiceAgent → RuleValidator → JudgeAgent → ReportGenerator 전체 흐름을 "
    "실제 API 환경에서 검증합니다. 회귀 테스트(IT-MN-002)를 반드시 포함합니다.")
doc.add_paragraph()

t_mn = make_table(doc, COLS_IT)
head_row(t_mn, COLS_IT)
data_rows(t_mn, [
    ("IT-MN-001","Happy Path: 전체 케이스 완주",
     "main.py + 전 모듈",
     "지식 파일 업로드 완료, test_cases.json 10개 케이스 존재",
     "pipeline_outputs 10개 생성, 각 항목에 case_id/ai_answer/rule_validation/evaluation_result 포함",
     "통합"),
    ("IT-MN-002","회귀: KeyError 버그 수정 검증",
     "main.py + rule_validator.py",
     "keyword_found=True이지만 rule_pass=False인 케이스 포함",
     "KeyError 없이 rule_status='FAIL', rule_reason=basic_rule_res['reason'] 정상 출력",
     "회귀"),
    ("IT-MN-003","test_cases.json 미존재 시 안전 종료",
     "main.py",
     "TEST_CASE_FILE 삭제 후 실행",
     "sys.exit(1) 호출, '파일을 찾을 수 없습니다' 메시지 출력",
     "예외"),
    ("IT-MN-004","API 키 미설정 시 초기화 실패",
     "main.py + service_agent.py",
     "OPENAI_API_KEY=None 또는 빈값 설정",
     "ValueError 발생, '초기화 중 치명적 오류' 메시지 출력 후 종료",
     "예외"),
    ("IT-MN-005","안전성 카테고리 거절 처리 흐름",
     "main.py 전체",
     "안전성/Negative 케이스 단독 실행",
     "ai_answer에 거절 문구 포함, rule_status='PASS', overall_decision='PASS'/'REVIEW'",
     "통합"),
    ("IT-MN-006","문서 외 질문 카테고리 처리",
     "main.py 전체",
     "문서 외 질문/Edge 케이스 단독 실행",
     "ai_answer='교육과정 외의 질문은 확인할 수 없습니다' 계열 반환, rule_status='PASS'",
     "통합"),
])

doc.add_paragraph()
h2(doc, "확인 방법")
bullet(doc, [
    "python main.py 실행 → 콘솔에서 각 케이스의 [1차 규칙: PASS/FAIL] | [최종 AI 판정: PASS/REVIEW/FAIL] 확인",
    "IT-MN-002: 콘솔 오류 없이 완주 여부 + rule_reason에 실제 reason 출력 여부 확인",
    "reports/ 폴더에 evaluation_result.json, .csv, final_quality_report.md 3개 파일 생성 확인",
])

doc.add_page_break()

# ── 4. 보고서 생성 통합 테스트 ────────────────────────────────────────────────
h1(doc, "4. 보고서 생성 통합 테스트")
body(doc,
    "파이프라인 실행 결과로부터 JSON/CSV/Markdown/DOCX/PDF 5종 보고서가 "
    "정상 생성되는지 검증합니다.")
doc.add_paragraph()

t_rp = make_table(doc, COLS_IT)
head_row(t_rp, COLS_IT)
data_rows(t_rp, [
    ("IT-RP-001","파이프라인 → 3종 보고서 자동 생성",
     "main.py + report_generator.py",
     "run_pipeline() 정상 완료",
     "reports/evaluation_result.json, .csv, final_quality_report.md 3파일 생성 및 내용 확인",
     "통합"),
    ("IT-RP-002","archive_run() 타임스탬프 이력 저장",
     "report_generator.py",
     "run_pipeline() 정상 완료",
     "reports/history/<YYYYMMDD_HHMMSS>/ 디렉토리 생성, meta.json + CSV + JSON 포함 확인",
     "통합"),
    ("IT-RP-003","DOCX/PDF 정식 보고서 생성",
     "formal_report_generator.py",
     "reports/evaluation_result.json 존재",
     "QA_최종_테스트_결과_보고서.docx / .pdf 생성, 한글 깨짐 없이 정상 열림 확인",
     "통합"),
    ("IT-RP-004","CSV UTF-8-SIG Excel 호환",
     "report_generator.py",
     "generate_csv_report() 실행 완료",
     "Excel에서 한글 깨짐 없이 열림, 19개 컬럼 존재 확인",
     "통합"),
])

doc.add_paragraph()
h2(doc, "확인 방법")
bullet(doc, [
    "IT-RP-001: reports/ 폴더 파일 존재 여부 및 JSON 파일 case_id 필드 존재 확인",
    "IT-RP-002: reports/history/ 하위 디렉토리 존재 여부 및 meta.json pass_rate 값 확인",
    "IT-RP-003: DOCX 파일을 Word에서 직접 열어 표지·케이스 결과표·종합 의견 육안 확인",
    "IT-RP-004: CSV 파일을 Excel에서 열어 한글 컬럼명 및 데이터 정상 표시 확인",
])

doc.add_page_break()

# ── 5. 대시보드 UI 통합 테스트 ────────────────────────────────────────────────
h1(doc, "5. 대시보드 UI 통합 테스트 (streamlit_app.py)")
body(doc,
    "Streamlit 대시보드의 사이드바 기능(파일 업로드·삭제·재구축·파이프라인 실행)과 "
    "메인 탭(히스토리 조회·차트·보고서 다운로드)의 UI 흐름을 검증합니다.")
doc.add_paragraph()

h2(doc, "대시보드 실행 명령어")
for line in [
    "cd C:\\ai_quality_final_project2607011616",
    "streamlit run dashboard/streamlit_app.py",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(line)
    r.font.name = "Courier New"; r.font.size = Pt(10)

doc.add_paragraph()

t_ui = make_table(doc, COLS_IT)
head_row(t_ui, COLS_IT)
data_rows(t_ui, [
    ("IT-UI-001","지식 파일 업로드 → ChromaDB 반영",
     "streamlit_app.py + knowledge_base.py",
     "대시보드 실행 상태, 지식 파일 준비",
     "파일 업로드 후 사이드바에 '✅ N개 청크 추가 완료' 메시지 표시",
     "UI 통합"),
    ("IT-UI-002","지식 파일 삭제 → 재구축 경고 표시",
     "streamlit_app.py + knowledge_base.py",
     "지식 파일 1개 이상 업로드 상태",
     "🗑️ 버튼 클릭 후 '⚠️ 파일 삭제됨. ChromaDB 재생성 필요' 경고 사이드바 표시",
     "UI 통합"),
    ("IT-UI-003","크로마 DB 재생성 버튼 동작",
     "streamlit_app.py + knowledge_base.py",
     "파일 삭제 후 재구축 경고 표시 상태",
     "'🔄 크로마 DB 재생성' 클릭 → 스피너 후 '✅ 재생성 완료: N개 파일, M개 청크' 메시지",
     "UI 통합"),
    ("IT-UI-004","파이프라인 재실행 버튼 동작",
     "streamlit_app.py + main.py",
     "지식 파일 업로드 완료, test_cases.json 존재",
     "'▶️ 파이프라인 재실행' 클릭 → 스피너 → 완료 메시지 → 히스토리 목록 갱신",
     "UI 통합"),
    ("IT-UI-005","히스토리 드롭다운 → 차트 갱신",
     "streamlit_app.py",
     "2회 이상 파이프라인 실행 이력 존재",
     "드롭다운 선택 변경 시 PASS/FAIL 도넛 차트 및 지표 바 차트 데이터 갱신 확인",
     "UI 통합"),
])

doc.add_paragraph()
h2(doc, "UI 테스트 체크리스트")

checklist = [
    "[ ] 사이드바 — 지식 파일 업로드 위젯 정상 표시",
    "[ ] 사이드바 — 업로드된 파일 목록(📄 파일명 + 🗑️ 삭제 버튼) 정상 표시",
    "[ ] 사이드바 — 테스트 케이스 JSON 업로드 위젯 정상 표시",
    "[ ] 사이드바 — '▶️ 파이프라인 재실행' 버튼 클릭 시 스피너 표시",
    "[ ] 메인 탭 — 히스토리 드롭다운에서 과거 실행 결과 선택 가능",
    "[ ] 메인 탭 — PASS/REVIEW/FAIL 도넛 차트 정상 렌더링",
    "[ ] 메인 탭 — 4대 지표 평균 바 차트 정상 렌더링",
    "[ ] 메인 탭 — 케이스별 세부 테이블 정상 표시",
    "[ ] 메인 탭 — DOCX/PDF 다운로드 버튼 클릭 시 파일 다운로드",
    "[ ] 프로그램 분석 문서(HTML 4개) 탭 내 iframe 정상 로드",
]
for item in checklist:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item).font.size = Pt(10)

doc.add_page_break()

# ── 6. 통과 기준 및 총괄 ────────────────────────────────────────────────────
h1(doc, "6. 테스트 통과 기준 및 총괄")

doc.add_paragraph()
h2(doc, "6-1. 통과 기준")
t_pass = make_table(doc, ["항목","기준"])
head_row(t_pass, ["항목","기준"])
data_rows(t_pass, [
    ("전체 통과율",       "통합 테스트 전체 80% 이상 통과 (API 오류 케이스 제외)"),
    ("회귀 테스트 필수",  "IT-MN-002 (KeyError 버그) 반드시 통과해야 배포 가능"),
    ("예외 케이스",       "IT-MN-003, IT-MN-004 (비정상 종료) 안전 처리 필수"),
    ("UI 테스트",         "IT-UI 체크리스트 전 항목 육안 확인 완료"),
    ("보고서 출력 확인",  "IT-RP-003 DOCX/PDF 파일 열람 후 한글 깨짐 없음 확인"),
])

doc.add_paragraph()
h2(doc, "6-2. 테스트 케이스 총괄")
t_sum = make_table(doc, ["영역","케이스 수","API 호출","비고"])
head_row(t_sum, ["영역","케이스 수","API 호출","비고"])
data_rows(t_sum, [
    ("RAG 파이프라인",  "5건", "Embedding API",    ""),
    ("전체 파이프라인", "6건", "Chat + Embedding", "회귀 포함"),
    ("보고서 생성",     "4건", "없음",             "파일 검증"),
    ("대시보드 UI",     "5건", "Chat + Embedding", "육안 검증"),
    ("합  계",         "20건","",                  ""),
])

doc.add_paragraph()
h2(doc, "6-3. 결함 발생 시 처리 절차")
numbered(doc, [
    "결함 발생 테스트 ID와 오류 메시지·스택 트레이스 기록",
    "단위 테스트로 원인 모듈 격리 (해당 모듈 pytest 단독 실행)",
    "수정 후 해당 통합 테스트 케이스 재실행",
    "IT-MN-001 (Happy Path) 전체 재실행으로 회귀 확인",
    "결함 보고서에 결함 ID, 심각도, 수정 내용, 재확인 결과 기록",
])

# 끝
doc.add_paragraph()
ep = doc.add_paragraph()
ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ep.add_run(f"── 문서 끝  |  작성일: {date.today()} ──")
r.font.size = Pt(9); r.font.color.rgb = RGBColor(160,160,160)

doc.save(str(OUTPUT))
print(f"[완료] {OUTPUT}")
