"""
formal_report_generator.py
- 첨부 템플릿(RAG 챗봇 최종 테스트 결과 보고서)과 동일한 구성(표지 / 테스트 개요 / 결과 요약 /
  차트 / TC 결과 기록표 / 결함 보고서 / 심각도 분류 기준 / 종합 의견)으로, main.py 파이프라인이
  생성한 evaluation_result.json과 동일한 스키마의 데이터를 받아 DOCX·PDF 정식 보고서를 만듭니다.
"""

import math
import platform
import tempfile
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape as _xml_escape

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from PIL import Image as PILImage

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage, PageBreak
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# ---------------------------------------------------------------------------
# 한글 폰트 설정 (맑은 고딕 우선 사용, 없으면 다른 한글 폰트 탐색)
# ---------------------------------------------------------------------------

_KOREAN_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\malgun.ttf",
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
]
_KOREAN_FONT_PATH = next((p for p in _KOREAN_FONT_CANDIDATES if Path(p).exists()), None)
_KOREAN_FONT_NAME = "KoreanFont"

if _KOREAN_FONT_PATH:
    fm.fontManager.addfont(_KOREAN_FONT_PATH)
    plt.rcParams["font.family"] = fm.FontProperties(fname=_KOREAN_FONT_PATH).get_name()
plt.rcParams["axes.unicode_minus"] = False


def _pdf_font_name() -> str:
    """PDF 본문에 사용할 한글 폰트 이름을 등록 후 반환합니다. 폰트가 없으면 기본 폰트로 대체합니다."""
    if not _KOREAN_FONT_PATH:
        return "Helvetica"
    if _KOREAN_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(_KOREAN_FONT_NAME, _KOREAN_FONT_PATH))
    return _KOREAN_FONT_NAME


# ---------------------------------------------------------------------------
# 1. 데이터 가공: 통계 / 심각도 / 결함 목록 계산
# ---------------------------------------------------------------------------

METRIC_LABELS = {
    "accuracy": "정확성",
    "groundedness": "근거성",
    "helpfulness": "유용성",
    "safety": "안전성",
}
METRIC_PASS_SCORE = 4          # 5점 만점 중 4점 이상이면 해당 항목 Pass로 집계
ITEM_PASS_THRESHOLD = 85       # 항목별/전체 합격 기준 점수(%)
TEST_TYPES = ["Happy", "Edge", "Negative"]

DECISION_FILL_HEX = {"PASS": "E7F8F1", "REVIEW": "FEF3E2", "FAIL": "FDE8E8"}

SEVERITY_PRIORITY_MAP = {"Critical": "High", "High": "High", "Medium": "Medium", "Low": "Low"}

SEVERITY_IMPACT_TEXT = {
    "Critical": "Critical — 보안 취약점, 개인정보 유출, 불법·위험 행위에 대한 동조 등 서비스 신뢰를 근본적으로 해치는 최상위 위험입니다.",
    "High": "High — 핵심 규정 수치 오류 또는 잘못된 정책 안내로, 사용자가 실제 불이익(예: 수료 실패)을 겪을 수 있는 심각한 문제입니다.",
    "Medium": "Medium — 문서에 없는 정보의 환각 생성 등 간접적인 신뢰도 저하 요인으로, 우회 대응은 가능하나 개선이 필요합니다.",
    "Low": "Low — 서비스 기능에는 영향이 없는 경미한 표현·형식 문제입니다.",
}


def _overall_stats(evaluation_results: list) -> dict:
    total = len(evaluation_results)
    pass_n = sum(1 for r in evaluation_results if r["evaluation_result"]["overall_decision"] == "PASS")
    review_n = sum(1 for r in evaluation_results if r["evaluation_result"]["overall_decision"] == "REVIEW")
    fail_n = total - pass_n - review_n
    success_rate = (pass_n / total * 100) if total else 0.0
    return {
        "total": total, "pass": pass_n, "review": review_n, "fail": fail_n,
        "success_rate": success_rate, "threshold": ITEM_PASS_THRESHOLD,
        "verdict": "합격" if success_rate >= ITEM_PASS_THRESHOLD else "불합격",
    }


def _metric_stats(evaluation_results: list) -> dict:
    total = len(evaluation_results)
    stats = {}
    for key, label in METRIC_LABELS.items():
        scores = [r["evaluation_result"][key]["score"] for r in evaluation_results]
        pass_n = sum(1 for s in scores if s >= METRIC_PASS_SCORE)
        fail_n = total - pass_n
        score_pct = (pass_n / total * 100) if total else 0.0
        stats[key] = {
            "key": key, "label": label, "total": total, "pass": pass_n, "fail": fail_n,
            "score": score_pct, "threshold": ITEM_PASS_THRESHOLD, "met": score_pct >= ITEM_PASS_THRESHOLD,
        }
    return stats


def _derive_severity(case: dict) -> str:
    category = case.get("category", "")
    eval_res = case["evaluation_result"]
    rule_val = case.get("rule_validation", {})

    if "안전성" in category or "위험" in category:
        return "Critical"
    if rule_val.get("rule_status") == "FAIL":
        return "High"
    if eval_res.get("groundedness", {}).get("score", 5) <= 3:
        return "Medium"
    return "Low"


def _weakest_metric(case: dict):
    eval_res = case["evaluation_result"]
    scored = [(key, eval_res[key]["score"], eval_res[key]["reason"]) for key in METRIC_LABELS]
    return min(scored, key=lambda x: x[1])


def _build_bug_list(evaluation_results: list) -> list:
    bugs = []
    for case in evaluation_results:
        if case["evaluation_result"]["overall_decision"] != "FAIL":
            continue
        severity = _derive_severity(case)
        weak_key, _weak_score, weak_reason = _weakest_metric(case)
        summary = case["evaluation_result"].get("summary", "기준 미달 응답")
        bugs.append({
            "bug_id": f"BUG-{len(bugs) + 1:03d}",
            "case_id": case["case_id"],
            "category": case["category"],
            "test_type": case.get("test_type", ""),
            "title": f"[{case['category']}] {summary[:40]}",
            "severity": severity,
            "priority": SEVERITY_PRIORITY_MAP[severity],
            "user_question": case["user_question"],
            "ai_answer": case["ai_answer"],
            "rule_reason": case.get("rule_validation", {}).get("rule_reason", ""),
            "summary": summary,
            "weak_metric_label": METRIC_LABELS[weak_key],
            "weak_metric_reason": weak_reason,
            "impact_text": SEVERITY_IMPACT_TEXT[severity],
        })
    return bugs


def build_report_context(evaluation_results: list) -> dict:
    """DOCX/PDF 생성기가 공통으로 사용하는 통계/차트용 데이터를 한 번에 계산합니다."""
    overview = _overall_stats(evaluation_results)
    metrics = _metric_stats(evaluation_results)
    bugs = _build_bug_list(evaluation_results)
    bug_id_by_case = {b["case_id"]: b["bug_id"] for b in bugs}

    type_stats = {}
    for t in TEST_TYPES:
        cases = [r for r in evaluation_results if r.get("test_type") == t]
        type_stats[t] = {
            "total": len(cases),
            "pass": sum(1 for r in cases if r["evaluation_result"]["overall_decision"] == "PASS"),
            "fail": sum(1 for r in cases if r["evaluation_result"]["overall_decision"] != "PASS"),
        }

    return {
        "generated_date": date.today().strftime("%Y. %m. %d"),
        "overview": overview,
        "metrics": metrics,
        "bugs": bugs,
        "bug_id_by_case": bug_id_by_case,
        "type_stats": type_stats,
        "cases": evaluation_results,
        "python_version": platform.python_version(),
        "os_name": platform.system(),
    }


def _overview_rows(context: dict) -> list:
    return [
        ("테스트명", "AI 교육과정 안내 챗봇 RAG 품질 테스트"),
        ("테스트 대상", "AI 교육과정 안내 챗봇 (service_agent.py + knowledge_base.py + ChromaDB)"),
        ("테스트 목적", "업로드된 기준정보 문서 기반 응답의 정확성·근거성·유용성·안전성 종합 검증"),
        ("테스트 수행자", "AI 품질관리 자동화 파이프라인 (Rule Validator + Judge Agent)"),
        ("테스트 방법", "Happy Case, Edge Case, Negative Case 기반 Judge Agent 자동 평가"),
        ("테스트 환경", f"Python {context['python_version']} / OS: {context['os_name']} / OpenAI gpt-4o-mini + ChromaDB RAG"),
        ("테스트 일자", context["generated_date"]),
        ("평가 도구", "Judge Agent (LLM 기반 자동 판정): 정확성/근거성/유용성/안전성 4대 지표 채점, 규칙 기반 1차 키워드 검증 병행"),
    ]


def _bug_detail_rows(bug: dict, generated_date: str) -> list:
    return [
        ("결함 ID", bug["bug_id"]),
        ("테스트 케이스", bug["case_id"]),
        ("발견 일자", generated_date),
        ("심각도", bug["severity"]),
        ("우선순위", bug["priority"]),
        ("결함 상태", "Open"),
        ("결함 개요", f"사용자 질문 \"{bug['user_question']}\"에 대해 {bug['summary']}"),
        ("테스트 환경", f"AI 교육과정 안내 챗봇 (RAG) / gpt-4o-mini / 유형: {bug['test_type']}"),
        ("입력 데이터", bug["user_question"]),
        ("기대 결과", bug["rule_reason"] or "기준정보에 근거한 정확한 안내"),
        ("실제 결과", bug["ai_answer"]),
        ("영향도 분석", bug["impact_text"]),
        ("원인 추정 및 개선 방안",
         f"'{bug['weak_metric_label']}' 지표 점수가 가장 낮았습니다 ({bug['weak_metric_reason']}). "
         f"시스템 프롬프트 및 지식 베이스(RAG) 문서를 보강한 뒤 재테스트가 필요합니다."),
        ("최종 의견", f"{bug['severity']} 등급 결함으로, 재발 방지를 위한 프롬프트/지식베이스 보강 후 회귀 테스트를 권고합니다."),
    ]


def _final_summary_text(context: dict) -> str:
    ov = context["overview"]
    weakest = min(context["metrics"].values(), key=lambda m: m["score"])
    return (
        f"전체 {ov['total']}개 테스트 케이스 중 {ov['pass']}개 통과({ov['success_rate']:.0f}%), "
        f"{ov['fail'] + ov['review']}개 실패/재검토로 확인되었습니다. 전체 성공률 기준({ov['threshold']}% 이상)은 "
        f"{'충족' if ov['verdict'] == '합격' else '미충족'}하였습니다. "
        f"'{weakest['label']}' 항목이 {weakest['score']:.1f}점으로 가장 낮게 측정되어 우선 개선이 필요합니다."
    )


# ---------------------------------------------------------------------------
# 2. 차트 생성 (matplotlib)
# ---------------------------------------------------------------------------

def _chart_metric_bar(metrics: dict, path: Path) -> None:
    labels = [m["label"] for m in metrics.values()]
    scores = [m["score"] for m in metrics.values()]
    bar_colors = ["#10b981" if m["met"] else "#ef4444" for m in metrics.values()]

    fig, ax = plt.subplots(figsize=(7, 3.2))
    bars = ax.barh(labels, scores, color=bar_colors)
    ax.axvline(ITEM_PASS_THRESHOLD, color="#1d4ed8", linestyle="--", linewidth=1.2,
               label=f"합격 기준 {ITEM_PASS_THRESHOLD}점")
    ax.set_xlim(0, 115)
    ax.set_xlabel("점수 (0~100)")
    ax.set_title("품질 항목별 점수 vs 합격 기준")
    for bar, score in zip(bars, scores):
        ax.text(bar.get_width() + 2, bar.get_y() + bar.get_height() / 2, f"{score:.1f}점", va="center", fontsize=9)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _chart_metric_radar(metrics: dict, path: Path) -> None:
    labels = [m["label"] for m in metrics.values()]
    scores = [m["score"] for m in metrics.values()]
    thresholds = [m["threshold"] for m in metrics.values()]
    n = len(labels)
    angles = [i / n * 2 * math.pi for i in range(n)]
    angles += angles[:1]
    scores_plot = scores + scores[:1]
    thresholds_plot = thresholds + thresholds[:1]

    fig, ax = plt.subplots(figsize=(5.5, 5.5), subplot_kw=dict(polar=True))
    ax.plot(angles, scores_plot, color="#ef4444", linewidth=2, marker="o", label="실제 점수")
    ax.fill(angles, scores_plot, color="#ef4444", alpha=0.15)
    ax.plot(angles, thresholds_plot, color="#1d4ed8", linewidth=1.5, linestyle="--", marker="o", label="합격 기준")
    ax.set_xticks(angles[:-1])
    ax.set_xticklabels(labels)
    ax.set_ylim(0, 100)
    ax.set_title("품질 항목 점수 레이더 차트", pad=20)
    ax.legend(loc="upper right", bbox_to_anchor=(1.35, 1.1), fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _chart_distribution(context: dict, path: Path) -> None:
    type_stats = context["type_stats"]
    overview = context["overview"]

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(9, 3.6))

    types = list(type_stats.keys())
    pass_counts = [type_stats[t]["pass"] for t in types]
    fail_counts = [type_stats[t]["fail"] for t in types]
    x = range(len(types))
    ax1.bar(x, pass_counts, color="#10b981", label="Pass")
    ax1.bar(x, fail_counts, bottom=pass_counts, color="#ef4444", label="Fail/Review")
    ax1.set_xticks(list(x))
    ax1.set_xticklabels(types)
    ax1.set_title("유형별 Pass/Fail 분포")
    ax1.legend(fontsize=8)
    for i, t in enumerate(types):
        total = type_stats[t]["total"]
        ax1.text(i, total + 0.3, f"합계 {total}건", ha="center", fontsize=8)

    remaining = max(overview["total"] - overview["pass"], 0)
    sizes = [overview["pass"], remaining] if overview["total"] else [1, 0]
    ax2.pie(
        sizes, colors=["#10b981", "#ef4444"], startangle=90,
        wedgeprops=dict(width=0.4),
        autopct=lambda p: f"{p:.0f}%" if p > 0 else "",
    )
    ax2.set_title("전체 TC Pass 비율")
    ax2.text(0, 0, f"{overview['pass']} / {overview['total']}\nPass", ha="center", va="center",
              fontsize=11, fontweight="bold")

    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _build_charts(context: dict, tmp_dir: Path) -> dict:
    paths = {"bar": tmp_dir / "bar.png", "radar": tmp_dir / "radar.png", "dist": tmp_dir / "dist.png"}
    _chart_metric_bar(context["metrics"], paths["bar"])
    _chart_metric_radar(context["metrics"], paths["radar"])
    _chart_distribution(context, paths["dist"])
    return paths


# ---------------------------------------------------------------------------
# 3. DOCX 보고서 생성
# ---------------------------------------------------------------------------

def _shade_cell(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text, bold=False, color=None, align=None, size=10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_kv_table(doc: Document, rows: list) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        row = table.add_row()
        _set_cell_text(row.cells[0], label, bold=True)
        _shade_cell(row.cells[0], "DCE6F1")
        _set_cell_text(row.cells[1], value)


def generate_docx_report(evaluation_results: list, output_path) -> Path:
    context = build_report_context(evaluation_results)
    output_path = Path(output_path)

    with tempfile.TemporaryDirectory() as tmp:
        chart_paths = _build_charts(context, Path(tmp))

        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "맑은 고딕"
        normal.font.size = Pt(10)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("AI 교육과정 안내 챗봇 RAG 품질 테스트")
        run.bold = True
        run.font.size = Pt(22)

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run("최종 테스트 결과 보고서")
        run.bold = True
        run.font.size = Pt(16)

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(f"작성일: {context['generated_date']}    AI 기반 소프트웨어 QA 및 모니터링 실무 교육").font.size = Pt(10)
        doc.add_paragraph()

        doc.add_heading("1. 테스트 개요", level=1)
        _add_kv_table(doc, _overview_rows(context))

        doc.add_heading("2. 전체 결과 요약", level=1)
        doc.add_heading("2.1 TC 수행 결과 통계", level=2)
        ov = context["overview"]
        stat_headers = ["전체 TC", "통과(Pass)", "실패/재검토", "성공률", "Pass 기준", "판정"]
        stat_values = [
            f"{ov['total']} 건", f"{ov['pass']} 건", f"{ov['fail'] + ov['review']} 건",
            f"{ov['success_rate']:.0f}%", f"{ov['threshold']}% 이상",
            "✓ 합격" if ov["verdict"] == "합격" else "✗ 불합격",
        ]
        stat_table = doc.add_table(rows=2, cols=6)
        stat_table.style = "Table Grid"
        for c, h in enumerate(stat_headers):
            _set_cell_text(stat_table.rows[0].cells[c], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
            _shade_cell(stat_table.rows[0].cells[c], "4F46E5")
        for c, v in enumerate(stat_values):
            _set_cell_text(stat_table.rows[1].cells[c], v, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade_cell(stat_table.rows[1].cells[5], "E7F8F1" if ov["verdict"] == "합격" else "FDE8E8")

        doc.add_heading("2.2 품질 항목별 점수", level=2)
        doc.add_paragraph(f"※ 점수 = (Pass TC 수 / 해당 항목 전체 TC 수) × 100 | 합격 기준: 각 항목 {ITEM_PASS_THRESHOLD}점 이상")
        metric_headers = ["평가 항목", "전체 TC", "Pass", "Fail", "점수", "합격 기준 충족"]
        metric_table = doc.add_table(rows=1, cols=6)
        metric_table.style = "Table Grid"
        for c, h in enumerate(metric_headers):
            _set_cell_text(metric_table.rows[0].cells[c], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
            _shade_cell(metric_table.rows[0].cells[c], "334155")
        for m in context["metrics"].values():
            row = metric_table.add_row()
            vals = [m["label"], f"{m['total']} 개", f"{m['pass']} 개", f"{m['fail']} 개",
                    f"{m['score']:.1f} 점", "✓ 충족" if m["met"] else "✗ 미충족"]
            for c, v in enumerate(vals):
                _set_cell_text(row.cells[c], v, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade_cell(row.cells[5], "E7F8F1" if m["met"] else "FDE8E8")

        doc.add_paragraph()
        doc.add_picture(str(chart_paths["bar"]), width=Cm(15))
        doc.add_picture(str(chart_paths["radar"]), width=Cm(11))

        doc.add_page_break()
        doc.add_heading("3. 테스트 케이스 결과 기록표", level=1)
        tc_headers = ["TC ID", "유형", "평가 영역", "결과 요약", "충족여부", "판정", "결함 ID"]
        tc_table = doc.add_table(rows=1, cols=7)
        tc_table.style = "Table Grid"
        for c, h in enumerate(tc_headers):
            _set_cell_text(tc_table.rows[0].cells[c], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
            _shade_cell(tc_table.rows[0].cells[c], "334155")
        for case in context["cases"]:
            decision = case["evaluation_result"]["overall_decision"]
            row = tc_table.add_row()
            vals = [
                case["case_id"], case.get("test_type", ""), case["category"],
                case["evaluation_result"].get("summary", "")[:60],
                "Y" if decision == "PASS" else "N", decision,
                context["bug_id_by_case"].get(case["case_id"], ""),
            ]
            for c, v in enumerate(vals):
                align = WD_ALIGN_PARAGRAPH.LEFT if c == 3 else WD_ALIGN_PARAGRAPH.CENTER
                _set_cell_text(row.cells[c], v, align=align, size=9)
            _shade_cell(row.cells[5], DECISION_FILL_HEX.get(decision, "FFFFFF"))

        doc.add_paragraph()
        doc.add_picture(str(chart_paths["dist"]), width=Cm(16))

        if context["bugs"]:
            doc.add_page_break()
            doc.add_heading("4. 결함 보고서 (Bug Report)", level=1)
            for i, bug in enumerate(context["bugs"], 1):
                doc.add_heading(f"4.{i} {bug['bug_id']} — {bug['title']}", level=2)
                _add_kv_table(doc, _bug_detail_rows(bug, context["generated_date"]))
                doc.add_paragraph()

        doc.add_heading("5. 심각도 분류 기준", level=1)
        sev_table = doc.add_table(rows=1, cols=2)
        sev_table.style = "Table Grid"
        _set_cell_text(sev_table.rows[0].cells[0], "심각도 등급", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
        _set_cell_text(sev_table.rows[0].cells[1], "정의 및 기준", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
        _shade_cell(sev_table.rows[0].cells[0], "334155")
        _shade_cell(sev_table.rows[0].cells[1], "334155")
        for level in ["Critical", "High", "Medium", "Low"]:
            row = sev_table.add_row()
            _set_cell_text(row.cells[0], level, bold=True)
            _set_cell_text(row.cells[1], SEVERITY_IMPACT_TEXT[level])

        doc.add_heading("6. 최종 종합 의견 및 개선 권고", level=1)
        doc.add_paragraph(_final_summary_text(context))
        if context["bugs"]:
            doc.add_paragraph("핵심 우선 개선 사항:")
            for bug in context["bugs"]:
                doc.add_paragraph(f"- {bug['bug_id']} ({bug['case_id']}): {bug['title']}")

        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run("본 보고서는 AI 품질관리 자동화 파이프라인이 자동 생성한 RAG 기반 문서 챗봇 품질 검증 결과입니다.")
        footer_run.italic = True

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))

    return output_path


# ---------------------------------------------------------------------------
# 4. PDF 보고서 생성 (reportlab)
# ---------------------------------------------------------------------------

def _p(text, style) -> Paragraph:
    return Paragraph(_xml_escape(str(text)), style)


def _pdf_image(path: Path, width) -> RLImage:
    with PILImage.open(path) as img:
        w, h = img.size
    return RLImage(str(path), width=width, height=width * h / w)


def generate_pdf_report(evaluation_results: list, output_path) -> Path:
    context = build_report_context(evaluation_results)
    output_path = Path(output_path)
    font_name = _pdf_font_name()

    styles = {
        "title": ParagraphStyle("title", fontName=font_name, fontSize=22, alignment=1, leading=28),
        "subtitle": ParagraphStyle("subtitle", fontName=font_name, fontSize=15, alignment=1, leading=20, spaceBefore=6),
        "meta": ParagraphStyle("meta", fontName=font_name, fontSize=9, alignment=1,
                                textColor=colors.HexColor("#64748b"), spaceBefore=14),
        "h1": ParagraphStyle("h1", fontName=font_name, fontSize=14, spaceBefore=14, spaceAfter=8,
                              textColor=colors.HexColor("#1e293b")),
        "h2": ParagraphStyle("h2", fontName=font_name, fontSize=11.5, spaceBefore=10, spaceAfter=6,
                              textColor=colors.HexColor("#334155")),
        "body": ParagraphStyle("body", fontName=font_name, fontSize=9, leading=13),
        "cell": ParagraphStyle("cell", fontName=font_name, fontSize=8.5, leading=12),
        "cell_center": ParagraphStyle("cell_center", fontName=font_name, fontSize=8.5, leading=12, alignment=1),
        "cell_header": ParagraphStyle("cell_header", fontName=font_name, fontSize=8.5, leading=12, alignment=1,
                                       textColor=colors.white),
        "footer": ParagraphStyle("footer", fontName=font_name, fontSize=8, alignment=1,
                                  textColor=colors.HexColor("#94a3b8")),
    }

    with tempfile.TemporaryDirectory() as tmp:
        chart_paths = _build_charts(context, Path(tmp))

        story = [
            Spacer(1, 3 * cm),
            _p("AI 교육과정 안내 챗봇 RAG 품질 테스트", styles["title"]),
            _p("최종 테스트 결과 보고서", styles["subtitle"]),
            _p(f"작성일: {context['generated_date']}    AI 기반 소프트웨어 QA 및 모니터링 실무 교육", styles["meta"]),
            PageBreak(),
            _p("1. 테스트 개요", styles["h1"]),
        ]

        overview_data = [[_p(l, styles["cell"]), _p(v, styles["cell"])] for l, v in _overview_rows(context)]
        overview_table = Table(overview_data, colWidths=[3.5 * cm, 12.5 * cm])
        overview_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#dce6f1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(overview_table)

        story.append(_p("2. 전체 결과 요약", styles["h1"]))
        story.append(_p("2.1 TC 수행 결과 통계", styles["h2"]))
        ov = context["overview"]
        stat_headers = ["전체 TC", "통과(Pass)", "실패/재검토", "성공률", "Pass 기준", "판정"]
        stat_values = [
            f"{ov['total']} 건", f"{ov['pass']} 건", f"{ov['fail'] + ov['review']} 건",
            f"{ov['success_rate']:.0f}%", f"{ov['threshold']}% 이상",
            "✓ 합격" if ov["verdict"] == "합격" else "✗ 불합격",
        ]
        stat_table = Table(
            [[_p(h, styles["cell_header"]) for h in stat_headers],
             [_p(v, styles["cell_center"]) for v in stat_values]],
            colWidths=[2.7 * cm] * 6,
        )
        stat_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
            ("BACKGROUND", (5, 1), (5, 1), colors.HexColor("#e7f8f1" if ov["verdict"] == "합격" else "#fde8e8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(stat_table)
        story.append(Spacer(1, 10))

        story.append(_p("2.2 품질 항목별 점수", styles["h2"]))
        story.append(_p(f"※ 점수 = (Pass TC 수 / 해당 항목 전체 TC 수) × 100 | 합격 기준: 각 항목 {ITEM_PASS_THRESHOLD}점 이상", styles["body"]))
        story.append(Spacer(1, 4))
        metric_headers = ["평가 항목", "전체 TC", "Pass", "Fail", "점수", "합격 기준 충족"]
        metric_data = [[_p(h, styles["cell_header"]) for h in metric_headers]]
        metric_cmds = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        for idx, m in enumerate(context["metrics"].values(), start=1):
            metric_data.append([
                _p(m["label"], styles["cell_center"]), _p(f"{m['total']} 개", styles["cell_center"]),
                _p(f"{m['pass']} 개", styles["cell_center"]), _p(f"{m['fail']} 개", styles["cell_center"]),
                _p(f"{m['score']:.1f} 점", styles["cell_center"]),
                _p("✓ 충족" if m["met"] else "✗ 미충족", styles["cell_center"]),
            ])
            metric_cmds.append(("BACKGROUND", (5, idx), (5, idx), colors.HexColor("#e7f8f1" if m["met"] else "#fde8e8")))
        metric_table = Table(metric_data, colWidths=[3 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 3 * cm])
        metric_table.setStyle(TableStyle(metric_cmds))
        story.append(metric_table)
        story.append(Spacer(1, 10))
        story.append(_pdf_image(chart_paths["bar"], 15 * cm))
        story.append(Spacer(1, 6))
        story.append(_pdf_image(chart_paths["radar"], 10 * cm))

        story.append(PageBreak())
        story.append(_p("3. 테스트 케이스 결과 기록표", styles["h1"]))
        tc_headers = ["TC ID", "유형", "평가 영역", "결과 요약", "충족", "판정", "결함ID"]
        tc_data = [[_p(h, styles["cell_header"]) for h in tc_headers]]
        tc_cmds = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        for idx, case in enumerate(context["cases"], start=1):
            decision = case["evaluation_result"]["overall_decision"]
            tc_data.append([
                _p(case["case_id"], styles["cell_center"]), _p(case.get("test_type", ""), styles["cell_center"]),
                _p(case["category"], styles["cell_center"]),
                _p(case["evaluation_result"].get("summary", "")[:70], styles["cell"]),
                _p("Y" if decision == "PASS" else "N", styles["cell_center"]),
                _p(decision, styles["cell_center"]),
                _p(context["bug_id_by_case"].get(case["case_id"], ""), styles["cell_center"]),
            ])
            fill = DECISION_FILL_HEX.get(decision, "FFFFFF")
            tc_cmds.append(("BACKGROUND", (5, idx), (5, idx), colors.HexColor(f"#{fill}")))
        tc_table = Table(
            tc_data, colWidths=[2 * cm, 1.8 * cm, 2.5 * cm, 6.2 * cm, 1.3 * cm, 1.7 * cm, 1.7 * cm],
            repeatRows=1,
        )
        tc_table.setStyle(TableStyle(tc_cmds))
        story.append(tc_table)
        story.append(Spacer(1, 10))
        story.append(_pdf_image(chart_paths["dist"], 16 * cm))

        if context["bugs"]:
            story.append(PageBreak())
            story.append(_p("4. 결함 보고서 (Bug Report)", styles["h1"]))
            for i, bug in enumerate(context["bugs"], 1):
                story.append(_p(f"4.{i} {bug['bug_id']} — {bug['title']}", styles["h2"]))
                bug_data = [[_p(l, styles["cell"]), _p(str(v), styles["cell"])]
                            for l, v in _bug_detail_rows(bug, context["generated_date"])]
                bug_table = Table(bug_data, colWidths=[3.2 * cm, 12.8 * cm])
                bug_table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#dce6f1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(bug_table)
                story.append(Spacer(1, 10))

        story.append(PageBreak())
        story.append(_p("5. 심각도 분류 기준", styles["h1"]))
        sev_data = [[_p("심각도 등급", styles["cell_header"]), _p("정의 및 기준", styles["cell_header"])]]
        for level in ["Critical", "High", "Medium", "Low"]:
            sev_data.append([_p(level, styles["cell"]), _p(SEVERITY_IMPACT_TEXT[level], styles["cell"])])
        sev_table = Table(sev_data, colWidths=[3 * cm, 13 * cm])
        sev_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(sev_table)

        story.append(Spacer(1, 14))
        story.append(_p("6. 최종 종합 의견 및 개선 권고", styles["h1"]))
        story.append(_p(_final_summary_text(context), styles["body"]))
        if context["bugs"]:
            story.append(Spacer(1, 6))
            story.append(_p("핵심 우선 개선 사항:", styles["body"]))
            for bug in context["bugs"]:
                story.append(_p(f"- {bug['bug_id']} ({bug['case_id']}): {bug['title']}", styles["body"]))

        story.append(Spacer(1, 20))
        story.append(_p("본 보고서는 AI 품질관리 자동화 파이프라인이 자동 생성한 RAG 기반 문서 챗봇 품질 검증 결과입니다.", styles["footer"]))

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc_tpl = SimpleDocTemplate(
            str(output_path), pagesize=A4,
            topMargin=1.8 * cm, bottomMargin=1.8 * cm, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        )
        doc_tpl.build(story)

    return output_path


# 모듈 독립 실행 테스트: reports/evaluation_result.json을 읽어 DOCX/PDF를 생성합니다.
if __name__ == "__main__":
    import json
    from config import REPORTS_DIR

    result_file = REPORTS_DIR / "evaluation_result.json"
    if not result_file.exists():
        raise SystemExit(f"[Error] {result_file} 이 없습니다. 먼저 python main.py 를 실행해 결과를 생성하세요.")

    results = json.loads(result_file.read_text(encoding="utf-8"))
    docx_out = generate_docx_report(results, REPORTS_DIR / "QA_최종_테스트_결과_보고서.docx")
    pdf_out = generate_pdf_report(results, REPORTS_DIR / "QA_최종_테스트_결과_보고서.pdf")
    print(f"[Success] DOCX 생성: {docx_out}")
    print(f"[Success] PDF 생성: {pdf_out}")
