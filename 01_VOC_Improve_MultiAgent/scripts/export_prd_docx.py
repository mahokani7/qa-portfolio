"""Export the as-built VOC_Improve PRD Markdown to a styled Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PRD_VOC_Improve_QA_Control_Center_20260716.md"
TARGET = ROOT / "docs" / "PRD_VOC_Improve_QA_Control_Center_20260716.docx"

NAVY = "123B66"
BLUE = "1F65AE"
TEAL = "0F8B8D"
LIGHT = "EAF2FA"
LINE = "C6D7E8"
MUTED = RGBColor(91, 111, 132)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(20, 41, 65)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for level, size, color in ((1, 18, NAVY), (2, 14, BLUE), (3, 11.5, TEAL)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "VOC_Improve · PRODUCT REQUIREMENTS DOCUMENT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "맑은 고딕"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("VOC_Improve QA Control Center · 2026-07-16")
    run.font.name = "맑은 고딕"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.autofit = True
    for row_index, values in enumerate(rows):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            cell.text = clean_inline(values[column_index]) if column_index < len(values) else ""
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                    run.font.size = Pt(8.5)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F8FC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "맑은 고딕"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
                    if row_index != 0:
                        run.font.size = Pt(8)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def export() -> Path:
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    title = clean_inline(source_lines[0].lstrip("# "))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(80)
    run = paragraph.add_run("VOC_Improve")
    run.font.name = "맑은 고딕"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("멀티 에이전트 QA Control Center")
    subtitle_run.font.name = "맑은 고딕"
    subtitle_run.font.size = Pt(18)
    subtitle_run.font.bold = True
    subtitle_run.font.color.rgb = RGBColor.from_string(TEAL)
    label = document.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_run = label.add_run("제품 요구사항 정의서(PRD) · As-built v1.0")
    label_run.font.name = "맑은 고딕"
    label_run.font.size = Pt(12)
    label_run.font.color.rgb = MUTED
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.add_run("기준일 2026-07-16")
    document.add_section(WD_SECTION.NEW_PAGE)

    index = 1
    in_code = False
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(document, table_rows)
            table_rows = []

    while index < len(source_lines):
        raw = source_lines[index]
        stripped = raw.strip()
        index += 1

        if stripped.startswith("```"):
            flush_table()
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.space_after = Pt(0)
            set_cell = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F0F4F8")
            set_cell.append(shading)
            run = paragraph.add_run(raw)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            values = [value.strip() for value in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", value.replace(" ", "")) for value in values):
                continue
            table_rows.append(values)
            continue

        flush_table()
        if not stripped or stripped == "---":
            continue
        if stripped.startswith("## "):
            document.add_heading(clean_inline(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            document.add_heading(clean_inline(stripped[4:]), level=2)
        elif stripped.startswith("#### "):
            document.add_heading(clean_inline(stripped[5:]), level=3)
        elif re.match(r"^\d+\. ", stripped):
            document.add_paragraph(clean_inline(re.sub(r"^\d+\. ", "", stripped)), style="List Number")
        elif stripped.startswith("- "):
            document.add_paragraph(clean_inline(stripped[2:]), style="List Bullet")
        elif stripped.startswith("> "):
            paragraph = document.add_paragraph(clean_inline(stripped[2:]))
            paragraph.paragraph_format.left_indent = Cm(0.6)
            for run in paragraph.runs:
                run.font.color.rgb = MUTED
                run.font.italic = True
        else:
            paragraph = document.add_paragraph(clean_inline(stripped))
            if "HOLD" in stripped or "배포 보류" in stripped:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(205, 55, 68)

    flush_table()
    properties = document.core_properties
    properties.title = title
    properties.subject = "현재 구현 기반 VOC_Improve 제품 요구사항 정의서"
    properties.author = "VOC_Improve QA Team"
    properties.keywords = "VOC, PRD, Multi-Agent, QA, OpenAI, Anthropic"
    document.save(TARGET)
    return TARGET


if __name__ == "__main__":
    print(export())
