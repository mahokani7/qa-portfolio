"""실행 이력 한 건을 감사 가능한 Word 종합 품질평가 보고서로 생성합니다."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = "12335F"
BLUE = "1D63AD"
LIGHT_BLUE = "EAF3FD"
GREEN = "087449"
ORANGE = "A85400"
GRAY = "5A6D82"


def _number(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _integer(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _safe_name(value: Any) -> str:
    text = re.sub(r"[^0-9A-Za-z가-힣_.-]+", "_", str(value or "run"))
    return text.strip("._")[:100] or "run"


def _set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margin(cell, value: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_text(cell, value: Any, *, bold: bool = False, color: str | None = None,
              size: float = 8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(value if value not in (None, "") else "-"))
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margin(cell)


def _add_table(document: Document, headers: Iterable[str], rows: Iterable[Iterable[Any]],
               widths: Iterable[float] | None = None):
    headers = list(headers)
    rows = [list(row) for row in rows]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _set_cell_shading(table.rows[0].cells[index], NAVY)
        _add_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF", size=8.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _add_text(cells[index], value, size=8)
            if row_index % 2:
                _set_cell_shading(cells[index], "F6F9FC")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(float(width))
    document.add_paragraph()
    return table


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Malgun Gothic"
    run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
    run.font.size = Pt(16 if level == 1 else 12)
    run.bold = True


def _add_body(document: Document, text: Any, *, bold: bool = False,
              color: str | None = None, size: float = 9.5) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(str(text))
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _status(run: dict[str, Any]) -> tuple[str, str, bool, bool]:
    total = _integer(run.get("total"))
    failed = _integer(run.get("failed"))
    score = _number(run.get("average_score"))
    threshold = _number(run.get("deployment_threshold"), 95.0)
    technical_pass = total > 0 and failed == 0
    final_approval = bool((run.get("release_approval") or {}).get("final_deployment_approved"))
    deployable = technical_pass and score >= threshold and final_approval
    if deployable:
        return "APPROVED", "배포 승인", technical_pass, True
    if technical_pass:
        return "HOLD", "기능 통과 · 운영 배포 보류", technical_pass, False
    return "FAIL", "품질 기준 미충족", technical_pass, False


def _trace_summary(run: dict[str, Any]) -> list[list[Any]]:
    stages: dict[str, dict[str, float]] = {}
    for case in run.get("cases") or []:
        for stage in case.get("trace") or []:
            agent = str(stage.get("agent") or stage.get("stage") or "Unknown")
            item = stages.setdefault(agent, {"count": 0, "duration": 0, "tokens": 0, "cost": 0})
            item["count"] += 1
            item["duration"] += _number(stage.get("duration_ms"))
            item["tokens"] += _integer(stage.get("total_tokens"))
            item["cost"] += _number(stage.get("cost"))
    return [
        [agent, int(item["count"]), round(item["duration"], 2), int(item["tokens"]), round(item["cost"], 6)]
        for agent, item in stages.items()
    ]


def _list_text(value: Any, default: str = "없음") -> str:
    if value in (None, "", [], {}):
        return default
    if isinstance(value, dict):
        return " · ".join(f"{key}: {_list_text(item, '-')}" for key, item in value.items())
    if isinstance(value, (list, tuple, set)):
        parts = []
        for item in value:
            if isinstance(item, dict):
                parts.append(str(item.get("item") or item.get("evidence") or item.get("sample") or _list_text(item)))
            else:
                parts.append(str(item))
        return ", ".join(part for part in parts if part) or default
    if isinstance(value, bool):
        return "예" if value else "아니오"
    return str(value)


def _percentage(value: Any) -> str:
    if value in (None, ""):
        return "-"
    number = _number(value, -1)
    return f"{number * 100:.1f}%" if 0 <= number <= 1 else f"{number:.1f}%"


def _case_parts(case: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    raw = case.get("raw") if isinstance(case.get("raw"), dict) else {}
    analysis = raw.get("analysis") if isinstance(raw.get("analysis"), dict) else {}
    quality = raw.get("quality") if isinstance(raw.get("quality"), dict) else {}
    metrics = case.get("metrics") if isinstance(case.get("metrics"), dict) else {}
    return analysis, quality, metrics


def _summary_and_policy(case: dict[str, Any]) -> tuple[str, str]:
    analysis, _, _ = _case_parts(case)
    summary = str(analysis.get("summary") or "").strip()
    policy = str(analysis.get("policy") or "").strip()
    output = str(case.get("output") or "").strip()
    if not summary and output:
        summary = output.split("정책:", 1)[0].replace("요약:", "", 1).strip()
    if not policy and "정책:" in output:
        policy = output.split("정책:", 1)[1].strip()
    return summary or "분석 결과 미기록", policy or "정책 개선안 미기록"


def _stage_result(stage: dict[str, Any]) -> str:
    output = stage.get("output") if isinstance(stage.get("output"), dict) else {}
    agent = str(stage.get("agent") or "")
    if agent == "Interpreter":
        return f"작업 {_list_text(output.get('task'), '-')} · 검색 조건 {_list_text(output.get('filters'))} · 최대 {_list_text(output.get('max_items'), '-')}건"
    if agent == "Retriever":
        samples = list(output.get("samples") or [])[:3]
        return f"관련 VOC {_integer(output.get('retrieved_count'))}건 검색 · 주요 근거: {_list_text(samples)}"
    if agent == "Summarizer":
        return "최종 요약: " + str(output.get("post_refine_summary") or output.get("pre_refine_summary") or "미기록")
    if agent == "Evaluator":
        return f"선정 후보 {_list_text(output.get('winner') or output.get('llm_winner'), '-')} · 후보 점수 {_list_text(output.get('scores'))}"
    if agent == "Critic":
        return f"보완 필요 · {_list_text(output.get('edits'))}" if output.get("need_refine") else "추가 보완 불필요 · 위험 지적 없음"
    if agent == "Improver":
        return str(output.get("policy") or "개선안 생성 생략")
    return str(stage.get("check") or stage.get("role") or "단계 수행 완료")


def _expectation_rows(case: dict[str, Any]) -> list[list[Any]]:
    _, quality, _ = _case_parts(case)
    checks = quality.get("checks") if isinstance(quality.get("checks"), dict) else {}
    intent = checks.get("intent") or {}
    keywords = checks.get("keywords") or {}
    required = checks.get("required_output") or {}
    prohibited = checks.get("prohibited_output") or {}
    return [
        ["기대 의도", intent.get("expected") or "미기록", "PASS" if intent.get("passed") else "FAIL",
         f"일치 {_percentage(intent.get('ratio'))} · 누락 {_list_text(intent.get('missing'))}"],
        ["필수 키워드", _list_text(keywords.get("matched")), "PASS" if keywords.get("passed") else "FAIL",
         f"누락 {_list_text(keywords.get('missing'))}"],
        ["필수 출력", _list_text(required.get("matched")), "PASS" if required.get("passed") else "FAIL",
         f"누락 {_list_text(required.get('missing'))}"],
        ["금지 출력", _list_text(prohibited.get("violations")), "PASS" if prohibited.get("passed") else "FAIL",
         "위반 없음" if prohibited.get("passed") else "금지 내용 검출"],
    ]


def _rubric_rows(case: dict[str, Any]) -> list[list[Any]]:
    _, quality, metrics = _case_parts(case)
    rubric = metrics.get("rubric") or quality.get("rubric") or {}
    rows = []
    for item in rubric.values():
        evidence = item.get("evidence") or {}
        evidence_text = "평가 기준 충족" if item.get("passed") else "평가 기준 미충족"
        if evidence.get("retrieved_count") is not None:
            evidence_text = f"검색 {evidence.get('retrieved_count')}건 중 관련 {evidence.get('relevant_count')}건"
        elif evidence.get("lexical_grounding_ratio") is not None:
            evidence_text = f"근거 일치 {_percentage(evidence.get('lexical_grounding_ratio'))} · 기대 키워드 {_percentage(evidence.get('expected_keyword_ratio'))}"
        elif evidence.get("winner"):
            evidence_text = f"선정 후보 {evidence.get('winner')} · 점수 {_list_text(evidence.get('scores'))}"
        elif evidence.get("need_refine") is not None:
            evidence_text = f"보완 필요 · {_list_text(evidence.get('edits'))}" if evidence.get("need_refine") else "추가 보완 불필요"
        elif evidence.get("missing") is not None:
            evidence_text = f"충족 {_list_text(evidence.get('matched'))} · 누락 {_list_text(evidence.get('missing'))}"
        elif evidence.get("duration_ms") is not None:
            evidence_text = f"처리시간 {evidence.get('duration_ms')} ms · 기준 {evidence.get('limit_ms')} ms"
        rows.append([
            item.get("label") or "평가 항목", f"{item.get('score', '-')}/{item.get('max_score', '-')}",
            "PASS" if item.get("passed") else "FAIL", evidence_text,
        ])
    return rows


def _rag_rows(case: dict[str, Any]) -> list[list[Any]]:
    _, quality, metrics = _case_parts(case)
    checks = quality.get("checks") or {}
    rag = metrics.get("rag") or checks.get("rag_metrics") or {}
    labels = (
        ("Context Precision", "context_precision"), ("Context Recall", "context_recall"),
        ("Faithfulness", "faithfulness"), ("Response Relevancy", "response_relevancy"),
        ("Noise Sensitivity", "noise_sensitivity"), ("Citation Coverage", "citation_coverage"),
    )
    rows: list[list[Any]] = []
    measured: list[float] = []
    for label, key in labels:
        value = rag.get(key)
        if value in (None, ""):
            rows.append([label, "-", "미측정"])
            continue
        number = max(0.0, min(1.0, _number(value)))
        measured.append(number)
        verdict = "양호" if number >= 0.8 else "주의" if number >= 0.6 else "개선 필요"
        rows.append([label, _percentage(number), verdict])
    aggregate_value = rag.get("ratio")
    aggregate = (
        max(0.0, min(1.0, _number(aggregate_value)))
        if aggregate_value not in (None, "")
        else (sum(measured) / len(measured) if measured else None)
    )
    if rag.get("skipped"):
        overall = "적용 제외"
    elif aggregate is None:
        overall = "측정 데이터 없음"
    else:
        passed = bool(rag.get("passed")) if "passed" in rag else aggregate >= 0.6
        overall = "PASS" if passed else "FAIL"
    rows.append(["RAG 6지표 종합", _percentage(aggregate), overall])
    return rows


def _trace_rows(case: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for stage in case.get("trace") or []:
        error = stage.get("error")
        status = "ERROR" if error else "완료"
        rows.append([
            stage.get("agent") or "-", stage.get("role") or "-", stage.get("check") or "-",
            f"{_number(stage.get('duration_ms')):.2f} ms", _integer(stage.get("total_tokens")),
            f"{stage.get('provider') or '-'} / {stage.get('model') or '-'}", status, _stage_result(stage),
        ])
    return rows


def _add_bullets(document: Document, values: Iterable[str]) -> None:
    for value in values:
        _add_body(document, "• " + str(value))


def build_history_word_report(run: dict[str, Any], output_dir: Path) -> Path:
    """첨부 PDF의 문서 구조를 실행 이력 한 건의 실제 데이터로 재구성합니다."""
    if not isinstance(run, dict) or not run.get("run_id"):
        raise ValueError("유효한 실행 이력 데이터가 필요합니다.")
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    status_code, status_label, technical_pass, deployable = _status(run)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = output_dir / (
        f"VOC_종합_품질평가_결과보고서_{_safe_name(run['run_id'])}_{stamp}.docx"
    )

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    styles = document.styles
    styles["Normal"].font.name = "Malgun Gothic"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    document.core_properties.title = "VOC 분석 및 개선안 생성 파이프라인 종합 품질평가 결과 보고서"
    document.core_properties.subject = f"실행 이력 {run['run_id']}"
    document.core_properties.author = "VOC QA Control Center"

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("VOC QA CONTROL CENTER · RUN HISTORY REPORT")
    header_run.font.name = "Malgun Gothic"
    header_run.font.size = Pt(7.5)
    header_run.font.color.rgb = RGBColor.from_string(GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"Run ID: {run['run_id']} · 생성 {datetime.now().isoformat(timespec='seconds')}")
    footer_run.font.name = "Malgun Gothic"
    footer_run.font.size = Pt(7.5)
    footer_run.font.color.rgb = RGBColor.from_string(GRAY)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(48)
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("VOC 분석 및 개선안 생성 파이프라인")
    title_run.font.name = "Malgun Gothic"
    title_run.font.size = Pt(13)
    title_run.font.color.rgb = RGBColor.from_string(BLUE)
    report_title = document.add_paragraph()
    report_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    report_title.paragraph_format.space_after = Pt(22)
    report_run = report_title.add_run("종합 품질평가 결과 보고서")
    report_run.font.name = "Malgun Gothic"
    report_run.font.size = Pt(25)
    report_run.bold = True
    report_run.font.color.rgb = RGBColor.from_string(NAVY)
    _add_table(document, ["평가 일시", "실행 유형", "평가 상태"], [[
        run.get("generated_at"), run.get("kind"), f"{status_code} · {status_label}"
    ]], [6.2, 5.2, 6.2])
    _add_body(document, "첨부된 품질평가 보고서 양식에 따라 실제 수행결과 값만 표와 설명으로 구성했습니다.", color=GRAY)

    cases = run.get("cases") or []
    blockers = []
    if not technical_pass:
        blockers.append(f"실패 {_integer(run.get('failed'))}건 또는 실행 데이터 미완료")
    if _number(run.get("average_score")) < _number(run.get("deployment_threshold"), 95):
        blockers.append(f"평균 {_number(run.get('average_score')):.1f}점이 배포 기준 {_number(run.get('deployment_threshold'), 95):.1f}점 미달")
    if not bool((run.get("release_approval") or {}).get("final_deployment_approved")):
        blockers.append("사람의 최종 배포 승인 미기록")

    _add_heading(document, "종합 결론: 이번 품질평가를 통해 얻은 것", 1)
    _add_body(document, f"[{status_code}] {status_label}", bold=True,
              color=GREEN if deployable else ORANGE, size=13)
    _add_body(document, f"이번 실행은 {_integer(run.get('total'))}개 테스트 중 {_integer(run.get('passed'))}개가 PASS하고 "
              f"{_integer(run.get('failed'))}개가 FAIL했습니다. 평균 품질점수는 {_number(run.get('average_score')):.1f}/100, "
              f"배포 기준은 {_number(run.get('deployment_threshold'), 95):.1f}/100입니다.")
    _add_table(document, ["전체", "PASS", "FAIL", "PASS율", "평균 점수", "배포 기준"], [[
        run.get("total"), run.get("passed"), run.get("failed"),
        f"{_number(run.get('pass_rate')):.1f}%", f"{_number(run.get('average_score')):.1f}/100",
        f"{_number(run.get('deployment_threshold'), 95):.1f}/100",
    ]], [2.8, 2.8, 2.8, 3.2, 3.2, 3.2])
    _add_body(document, "최종 판단 근거: " + (" · ".join(blockers) if blockers else "기능·점수·승인 기준을 모두 충족"), bold=True)
    _add_table(document, ["실행 ID", "도메인 / 모드", "모델", "데이터 / 프롬프트 버전"], [[
        run.get("run_id"), f"{run.get('domain') or '-'} / {run.get('mode') or '-'}",
        run.get("model") or run.get("provider") or "미기록",
        f"{run.get('dataset_version') or '-'} / {run.get('prompt_version') or '-'}",
    ]], [5.0, 4.0, 4.0, 5.0])
    document.add_page_break()

    _add_heading(document, "1. 우리가 구축한 품질평가 구조", 1)
    _add_body(document, "파이프라인의 전체 정보 흐름과 실제 실행에서 확인된 각 모듈의 역할은 다음과 같습니다.")
    stages = (
        ("01. Interpreter", "질문 의도, 분석 대상과 검색 조건을 구조화합니다."),
        ("02. Retriever", "고객 VOC에서 질문과 직접 관련된 근거 데이터를 검색합니다."),
        ("03. Summarizer", "반복 불만, 핵심 현상과 고객 영향을 요약합니다."),
        ("04. Evaluator", "요약 후보의 관련성과 타당성을 점수화하고 최종 후보를 선택합니다."),
        ("05. Critic", "근거 부족, 누락, 과장과 안전 위험을 비판적으로 점검합니다."),
        ("06. Refine", "Critic이 보완을 요구한 경우 요약과 정책을 교정합니다."),
        ("07. Improver", "원인과 연결된 담당자·기한·지표가 있는 정책 개선안을 생성합니다."),
        ("08. 독립 LLM Judge", "생성 주체와 분리된 모델이 정확성·충실성·유용성·안전성을 평가합니다."),
        ("09. 최종 판정", "기능 결과, 점수, 중대 결함과 사람 승인을 종합합니다."),
    )
    _add_bullets(document, [f"{name}: {description}" for name, description in stages])
    _add_body(document, "품질 평가의 핵심 차별점은 생성 Agent와 최종 평가 주체를 분리하여 자기평가 편향을 줄이는 것입니다.", bold=True, color=BLUE)
    document.add_page_break()

    _add_heading(document, "2. 고객 불만 사항에 기반한 개선 정책 도출 메커니즘", 1)
    _add_body(document, "각 테스트 케이스는 고객 질문 → VOC 근거 → 요약 → 내부 평가·비판 → 정책 개선안 → 정량 검증의 순서로 처리되었습니다.")
    if cases:
        _add_table(document, ["Case", "상태", "점수", "고객 질문", "정책 평가"], [
            [case.get("case_id"), case.get("status"), f"{_number(case.get('score')):.1f}",
             str(case.get("question") or "-")[:160],
             (_case_parts(case)[2].get("deployment") or _case_parts(case)[1].get("deployment") or {}).get("label", "-")]
            for case in cases
        ], [2.3, 1.8, 1.7, 7.5, 4.7])
        for index, case in enumerate(cases, 1):
            analysis, quality, metrics = _case_parts(case)
            summary, policy = _summary_and_policy(case)
            deployment = metrics.get("deployment") or quality.get("deployment") or {}
            _add_heading(document, f"2.{index} {case.get('case_id')} 수행결과", 2)
            _add_table(document, ["분석 단계", "실제 수행결과"], [
                ["고객 질문", case.get("question") or "-"],
                ["VOC 분석 결과", summary],
                ["정책 개선안", policy],
                ["품질 판정", f"{case.get('status')} · {_number(case.get('score')):.1f}/100"],
                ["배포 평가", f"{deployment.get('label') or deployment.get('code') or '-'} · 기준 {deployment.get('minimum_score', run.get('deployment_threshold', 95))}점 · 차이 {deployment.get('score_gap', '-')}점"],
                ["결함 / 중대 차단", f"{_list_text(metrics.get('defects'))} / {_list_text(metrics.get('hard_blockers'))}"],
            ], [4.0, 14.0])
            _add_heading(document, f"2.{index}.1 기대 결과 충족 여부", 2)
            _add_table(document, ["검증 기준", "확인된 수행결과", "판정", "근거"], _expectation_rows(case), [3.0, 6.0, 2.0, 7.0])
            _add_heading(document, f"2.{index}.2 9개 품질 평가 항목", 2)
            rubric_rows = _rubric_rows(case)
            if rubric_rows:
                _add_table(document, ["평가 항목", "점수", "판정", "평가 근거"], rubric_rows, [4.0, 2.3, 2.0, 9.7])
            _add_heading(document, f"2.{index}.3 RAG 전문 지표", 2)
            _add_table(document, ["RAG 지표", "측정값", "지표·종합 판정"], _rag_rows(case), [7.0, 4.0, 7.0])
            _add_heading(document, f"2.{index}.4 6-Agent 단계별 수행결과", 2)
            trace_rows = _trace_rows(case)
            if trace_rows:
                _add_table(document, ["Agent", "역할·점검", "시간·토큰", "모델", "상태", "수행결과"], [
                    [row[0], f"{row[1]} · {row[2]}", f"{row[3]} · {row[4]} tok", row[5], row[6], row[7]]
                    for row in trace_rows
                ], [2.3, 3.5, 2.4, 2.7, 1.5, 5.6])
    else:
        _add_body(document, "이 실행 유형에는 케이스 단위 수행결과가 없습니다.")
    document.add_page_break()

    _add_heading(document, "3. 왜 해당 개선안이 타당하다고 판단했는가", 1)
    rag_values = []
    actionability_scores = []
    for case in cases:
        _, quality, metrics = _case_parts(case)
        checks = quality.get("checks") or {}
        rag = metrics.get("rag") or checks.get("rag_metrics") or {}
        if rag:
            rag_values.append(rag)
        rubric = metrics.get("rubric") or quality.get("rubric") or {}
        if rubric.get("improver_actionability"):
            actionability_scores.append(_number(rubric["improver_actionability"].get("score")))
    avg_relevance = sum(_number(item.get("response_relevancy")) for item in rag_values) / len(rag_values) if rag_values else 0
    avg_faithfulness = sum(_number(item.get("faithfulness")) for item in rag_values) / len(rag_values) if rag_values else 0
    avg_actionability = sum(actionability_scores) / len(actionability_scores) if actionability_scores else 0
    _add_table(document, ["전문 평가 필터", "실제 측정 근거", "판단"], [
        ["3.1 불만 관련성", f"Response Relevancy 평균 {_percentage(avg_relevance)}", "고객 질문과 직접 연결된 답변인지 측정"],
        ["3.2 근본 원인 대응성", f"Faithfulness 평균 {_percentage(avg_faithfulness)}", "VOC 근거에서 벗어난 단정과 환각을 통제"],
        ["3.3 실행 현실성", f"Improver 실행 가능성 평균 {avg_actionability:.1f}/15", "담당자·기한·행동·지표·우선순위 포함 여부를 평가"],
        ["3.4 검증 가능성", f"처리시간 평균 {_number(run.get('average_duration_ms')):.2f} ms · P95 {_number(run.get('p95_duration_ms')):.2f} ms", "개선 전후를 점수와 성능 지표로 재측정 가능"],
    ], [4.0, 6.0, 8.0])
    document.add_page_break()

    _add_heading(document, "4. 프로세스 단계별 전문 판단 프로세스", 1)
    _add_body(document, "각 Agent가 실제 실행에서 수행한 판단과 결과를 대표 케이스 기준으로 요약했습니다.")
    representative = cases[0] if cases else {}
    stage_rows = _trace_rows(representative) if representative else []
    if stage_rows:
        _add_table(document, ["Agent", "전문 판단 역할", "점검 기준", "실제 수행결과"],
                   [[row[0], row[1], row[2], row[7]] for row in stage_rows], [3.0, 4.0, 4.5, 6.5])
    _add_table(document, ["평가 주체", "역할 구분", "배포 판단에서의 사용"], [
        ["Evaluator", "후보 타당성·구조·내부 점수", "후보 선택과 내부 품질 개선"],
        ["Critic", "누락·과장·근거·안전 위험", "Refine 필요 여부와 수정 지시"],
        ["독립 LLM Judge", "정확성·충실성·유용성·안전성", "생성 주체와 분리된 외부 감사 점수"],
        ["Human Review", "업무 타당성과 책임 있는 승인", "정식 운영 배포의 최종 승인"],
    ], [3.2, 7.0, 7.8])
    document.add_page_break()

    _add_heading(document, "5. 과정에서 점검한 기술적 사항", 1)
    _add_table(document, ["점검 항목", "실제 수행결과"], [
        ["기능 테스트", f"{_integer(run.get('passed'))}/{_integer(run.get('total'))} PASS · 실패 {_integer(run.get('failed'))}건"],
        ["성능", f"평균 {_number(run.get('average_duration_ms')):.2f} ms · P95 {_number(run.get('p95_duration_ms')):.2f} ms"],
        ["사용량·비용", f"토큰 {_integer(run.get('total_tokens'))} · 추정 비용 {_number(run.get('estimated_cost')):.6f}"],
        ["장애·재시도", f"API 429 {_integer(run.get('rate_limit_count'))}건 · 결함 {_integer(run.get('defects_count'))}건"],
        ["안전성", f"중대 위반 {_integer(run.get('critical_violations'))}건 · Hard Block {_integer(run.get('hard_blocked'))}건"],
        ["버전 추적", f"Git {run.get('git_commit') or '-'} · Dataset {run.get('dataset_version') or '-'} · Prompt {run.get('prompt_version') or '-'}"],
        ["실행 모드", f"{run.get('domain') or '-'} / {run.get('mode') or '-'} · Live 검증 {_list_text(run.get('live_verified'))}"],
    ], [5.0, 13.0])
    trace_summary = _trace_summary(run)
    if trace_summary:
        _add_heading(document, "5.1 Agent별 성능·사용량 집계", 2)
        _add_table(document, ["Agent", "단계 수", "총 처리시간(ms)", "토큰", "비용"], trace_summary, [4.0, 2.5, 4.0, 3.0, 3.0])
    document.add_page_break()

    _add_heading(document, "6. 품질전문가가 성공적인 품질평가라고 판단하는 근거", 1)
    _add_bullets(document, [
        f"요구사항 대비 추적성: {_integer(run.get('total'))}개 테스트가 질문·기대 결과·실제 출력·점수·판정으로 연결되었습니다.",
        "균형 잡힌 계층형 평가: 기능 결과와 9개 품질 항목, RAG 6지표, Agent 단계 결과를 함께 검증했습니다.",
        f"네거티브 경로 통제: 결함 {_integer(run.get('defects_count'))}건, 중대 위반 {_integer(run.get('critical_violations'))}건, 429 {_integer(run.get('rate_limit_count'))}건을 별도 기록했습니다.",
        "생성과 평가 분리: Summarizer·Improver의 생성 결과를 Evaluator·Critic 및 독립 Judge가 서로 다른 역할로 평가합니다.",
        "감사 가능성: 동일 Run ID로 실행 시각, 모델, 버전, 점수, Agent 결과와 사람 승인 상태가 연결됩니다.",
    ])
    all_rubric = []
    for case in cases:
        all_rubric.extend(_rubric_rows(case))
    if all_rubric:
        _add_heading(document, "6.1 요구사항-평가 결과 추적표", 2)
        _add_table(document, ["평가 요구사항", "측정 점수", "판정", "실제 근거"], all_rubric, [4.2, 2.3, 2.0, 9.5])
    document.add_page_break()

    _add_heading(document, "7. 이번 품질평가를 통해 실질적으로 얻은 성과", 1)
    _add_bullets(document, [
        "고객 불만을 질문, 근거 VOC, 요약, 정책 개선안, 검증 지표가 연결된 구조화된 품질 데이터로 전환했습니다.",
        "왜 해당 개선안을 제시했는지 기대 결과 충족 여부와 평가 근거를 통해 역추적할 수 있습니다.",
        "Evaluator → Critic → Refine → 독립 Judge의 다중 안전망으로 환각과 과도한 단정을 통제했습니다.",
        "모델·데이터·프롬프트 버전과 비용·성능을 같은 실행 이력에서 관리할 수 있습니다.",
        "테스트 완료와 동시에 화면 조회 및 제출용 Word 최종보고서를 생성해 품질 의사결정 시간을 단축했습니다.",
    ])

    _add_heading(document, "8. 성공 판정의 합리적 범위와 한계 명시", 1)
    risks = []
    if not technical_pass:
        risks.append(["기능 결함", "높음", "FAIL 케이스의 Agent 수행결과와 평가 근거를 개선한 뒤 동일 조건으로 재시험"])
    if _number(run.get("average_score")) < _number(run.get("deployment_threshold"), 95):
        risks.append(["배포 기준 미달", "높음", "낮은 점수 항목을 개선하고 기준 점수 이상인지 재평가"])
    if _integer(run.get("rate_limit_count")):
        risks.append(["API 429", "중간", "재시도·동시성·비용 제한을 조정하고 장애 회복 결과 확보"])
    if not deployable:
        risks.append(["사람 승인", "중간", "Judge와 사람 평가 차이를 검토하고 최종 배포 승인 기록"])
    risks.extend([
        ["운영 드리프트", "중간", "모델·프롬프트·데이터 버전별 점수와 P95 추세를 지속 감시"],
        ["보안·개인정보", "중간", "OWASP Red Team과 개인정보 마스킹 회귀를 배포 전 반복"],
    ])
    _add_table(document, ["잔여 위험", "수준", "운영 권고사항"], risks, [4.0, 2.5, 11.5])
    _add_heading(document, "8.1 QA Lead 최종 종합 평가 의견", 2)
    _add_body(document, f"[최종 종합 평가 판정: {status_code} · {status_label}]", bold=True,
              color=GREEN if deployable else ORANGE, size=12)
    _add_body(document, "본 판정은 이 보고서에 표시된 실행 ID, 데이터셋, 모델과 프롬프트 버전의 실제 수행결과에 적용됩니다. 다른 운영 부하와 새 데이터 분포까지 자동으로 보증하지는 않습니다.")
    _add_body(document, "정식 배포는 기능 PASS뿐 아니라 배포 기준 점수, 중대 위반 0건, 보안 진단과 사람 최종 승인을 함께 충족해야 합니다.")
    document.save(path)
    return path
