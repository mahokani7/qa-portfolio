
#이 파일은 문서 저장, PDF 읽기, 벡터DB 생성, 검색, 답변 생성을 담당합니다.
from pathlib import Path
from typing import List

from dotenv import load_dotenv

from langchain_community.document_loaders import TextLoader, PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_chroma import Chroma

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent

DOCUMENTS_DIR = BASE_DIR / "documents"
UPLOADS_DIR = BASE_DIR / "uploads"
CHROMA_DIR = BASE_DIR / "chroma_db"

COLLECTION_NAME = "rag_documents"


def ensure_directories():
    DOCUMENTS_DIR.mkdir(exist_ok=True)
    UPLOADS_DIR.mkdir(exist_ok=True)
    CHROMA_DIR.mkdir(exist_ok=True)


def get_embeddings():
    return OpenAIEmbeddings(
        model="text-embedding-3-small"
    )


def get_vector_db():
    return Chroma(
        persist_directory=str(CHROMA_DIR),
        embedding_function=get_embeddings(),
        collection_name=COLLECTION_NAME
    )


def load_docx(file_path: Path) -> List[Document]:
    from docx import Document as DocxDocument
    doc = DocxDocument(str(file_path))
    text = "\n".join(
        para.text for para in doc.paragraphs if para.text.strip()
    )
    return [Document(page_content=text, metadata={"source": file_path.name})]


def load_excel(file_path: Path) -> List[Document]:
    suffix = file_path.suffix.lower()
    if suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(file_path))
        rows = []
        for sheet in wb.sheets():
            rows.append(f"[시트: {sheet.name}]")
            for i in range(sheet.nrows):
                row_text = "\t".join(str(cell.value) for cell in sheet.row(i))
                if row_text.strip():
                    rows.append(row_text)
        text = "\n".join(rows)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        rows = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows.append(f"[시트: {sheet_name}]")
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join("" if cell is None else str(cell) for cell in row)
                if row_text.strip():
                    rows.append(row_text)
        text = "\n".join(rows)
    return [Document(page_content=text, metadata={"source": file_path.name})]


def load_file(file_path: Path):
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        loader = PyPDFLoader(str(file_path))
        documents = loader.load()

    elif suffix in [".txt", ".md"]:
        loader = TextLoader(str(file_path), encoding="utf-8")
        documents = loader.load()

    elif suffix == ".docx":
        documents = load_docx(file_path)

    elif suffix in [".xlsx", ".xls"]:
        documents = load_excel(file_path)

    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix}")

    for doc in documents:
        doc.metadata["source"] = file_path.name

    return documents


def load_all_documents():
    ensure_directories()

    all_documents = []

    target_dirs = [DOCUMENTS_DIR, UPLOADS_DIR]

    supported = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".xls"}

    for target_dir in target_dirs:
        for file_path in target_dir.glob("*"):
            if file_path.suffix.lower() not in supported:
                continue

            documents = load_file(file_path)
            all_documents.extend(documents)

    return all_documents


def rebuild_vector_db():
    ensure_directories()

    documents = load_all_documents()

    if not documents:
        raise ValueError(
            "documents 또는 uploads 폴더에 txt, md, pdf 문서를 넣어주세요."
        )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=700,
        chunk_overlap=120
    )

    split_documents = splitter.split_documents(documents)

    if CHROMA_DIR.exists():
        import shutil
        shutil.rmtree(CHROMA_DIR)

    CHROMA_DIR.mkdir(exist_ok=True)

    Chroma.from_documents(
        documents=split_documents,
        embedding=get_embeddings(),
        persist_directory=str(CHROMA_DIR),
        collection_name=COLLECTION_NAME
    )

    return {
        "original_document_count": len(documents),
        "chunk_count": len(split_documents)
    }


def save_uploaded_files(uploaded_files):
    ensure_directories()

    saved_files = []

    for uploaded_file in uploaded_files:
        save_path = UPLOADS_DIR / uploaded_file.name

        with open(save_path, "wb") as file:
            file.write(uploaded_file.getbuffer())

        saved_files.append(uploaded_file.name)

    return saved_files


def answer_question(question: str):
    vector_db = get_vector_db()

    retrieved_docs = vector_db.similarity_search(
        query=question,
        k=4
    )

    if not retrieved_docs:
        return {
            "answer": "관련 문서를 찾지 못했습니다.",
            "sources": [],
            "contexts": []
        }

    context = "\n\n".join(
        [
            f"[출처: {doc.metadata.get('source', '알 수 없음')}]\n"
            f"{doc.page_content}"
            for doc in retrieved_docs
        ]
    )

    llm = ChatOpenAI(
        model="gpt-4.1-mini",
        temperature=0
    )

    prompt = f"""
당신은 문서 기반 RAG 챗봇입니다.

반드시 아래 검색 문서에 포함된 내용만 근거로 답변하십시오.

문서에 없는 사실을 추측하거나 만들어 내지 마십시오.

답을 찾을 수 없으면 반드시 아래 문장으로 답하십시오.

"제공된 문서에서는 확인할 수 없습니다."

[검색 문서]
{context}

[사용자 질문]
{question}

[답변 원칙]
1. 한국어로 답변합니다.
2. 답변은 짧고 명확하게 작성합니다.
3. 문서 근거가 있으면 핵심 내용을 설명합니다.
4. 문서 밖의 내용은 절대 추가하지 않습니다.
"""

    response = llm.invoke(prompt)

    sources = list(
        {
            doc.metadata.get("source", "알 수 없음")
            for doc in retrieved_docs
        }
    )

    contexts = [
        {
            "source": doc.metadata.get("source", "알 수 없음"),
            "content": doc.page_content
        }
        for doc in retrieved_docs
    ]

    return {
        "answer": response.content,
        "sources": sources,
        "contexts": contexts
    }