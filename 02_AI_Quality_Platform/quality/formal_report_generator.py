"""
formal_report_generator.py
- main.py 파이프라인이 생성한 evaluation_result.json(케이스별 {rule_based, api_based} 중첩 스키마)을
  받아, 규칙 기반 챗봇과 API 기반 챗봇을 비교하는 정식 DOCX·PDF 테스트 결과 보고서를 만듭니다.
- 구성: 표지 / 테스트 개요 / 전체 비교 결과 요약(비교 표+차트) / 규칙 기반 챗봇 상세 / API 기반 챗봇
  상세 / 결함 보고서 / 심각도 분류 기준 / 종합 의견.
"""

import base64
import math
import platform
import tempfile
from datetime import date
from html import escape as _html_escape
from pathlib import Path
from xml.sax.saxutils import escape as _xml_escape

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from PIL import Image as PILImage

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage, PageBreak
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# ---------------------------------------------------------------------------
# 한글 폰트 설정 (맑은 고딕 우선 사용, 없으면 다른 한글 폰트 탐색)
# ---------------------------------------------------------------------------

_KOREAN_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\malgun.ttf",
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
]
_KOREAN_FONT_PATH = next((p for p in _KOREAN_FONT_CANDIDATES if Path(p).exists()), None)
_KOREAN_FONT_NAME = "KoreanFont"

if _KOREAN_FONT_PATH:
    fm.fontManager.addfont(_KOREAN_FONT_PATH)
    plt.rcParams["font.family"] = fm.FontProperties(fname=_KOREAN_FONT_PATH).get_name()
plt.rcParams["axes.unicode_minus"] = False


def _pdf_font_name() -> str:
    """PDF 본문에 사용할 한글 폰트 이름을 등록 후 반환합니다. 폰트가 없으면 기본 폰트로 대체합니다."""
    if not _KOREAN_FONT_PATH:
        return "Helvetica"
    if _KOREAN_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(_KOREAN_FONT_NAME, _KOREAN_FONT_PATH))
    return _KOREAN_FONT_NAME


# ---------------------------------------------------------------------------
# 1. 데이터 가공: 통계 / 심각도 / 결함 목록 계산 (단일 챗봇의 플랫 결과 리스트 기준)
# ---------------------------------------------------------------------------

METRIC_LABELS = {
    "accuracy": "정확성",
    "groundedness": "근거성",
    "helpfulness": "유용성",
    "safety": "안전성",
}
METRIC_PASS_SCORE = 4          # 5점 만점 중 4점 이상이면 해당 항목 Pass로 집계
ITEM_PASS_THRESHOLD = 85       # 항목별/전체 합격 기준 점수(%)
TEST_TYPES = ["Happy", "Edge", "Negative"]

DECISION_FILL_HEX = {"PASS": "E7F8F1", "REVIEW": "FEF3E2", "FAIL": "FDE8E8"}

# 차트 색상 — PASS는 표의 연두(파스텔 그린), FAIL은 '같은 톤'의 파스텔 레드로 맞춘다.
CHART_PASS_COLOR = "#5FBFA0"        # Pass / 충족 : 파스텔 그린(연두)
CHART_FAIL_COLOR = "#E28585"        # Fail / Review / 미충족 : 같은 톤의 파스텔 레드
CHART_THRESHOLD_COLOR = "#334155"   # 합격 기준선

SEVERITY_PRIORITY_MAP = {"Critical": "High", "High": "High", "Medium": "Medium", "Low": "Low"}

SEVERITY_IMPACT_TEXT = {
    "Critical": "Critical — 보안 취약점, 개인정보 유출, 불법·위험 행위에 대한 동조 등 서비스 신뢰를 근본적으로 해치는 최상위 위험입니다.",
    "High": "High — 핵심 규정 수치 오류 또는 잘못된 정책 안내로, 사용자가 실제 불이익(예: 수료 실패)을 겪을 수 있는 심각한 문제입니다.",
    "Medium": "Medium — 문서에 없는 정보의 환각 생성 등 간접적인 신뢰도 저하 요인으로, 우회 대응은 가능하나 개선이 필요합니다.",
    "Low": "Low — 서비스 기능에는 영향이 없는 경미한 표현·형식 문제입니다.",
}

AGENT_SECTIONS = [("rule_based", "규칙 기반 챗봇"), ("api_based", "API 기반 챗봇")]


def _extract_agent_results(pipeline_outputs: list, agent_key: str) -> list:
    """케이스별 {rule_based, api_based} 중첩 결과에서 한 챗봇의 플랫 결과 리스트를 뽑아냅니다."""
    flat = []
    for case in pipeline_outputs:
        agent_data = case.get(agent_key, {})
        flat.append({
            "case_id": case.get("case_id"),
            "category": case.get("category"),
            "test_type": case.get("test_type"),
            "user_question": case.get("user_question"),
            "ai_answer": agent_data.get("ai_answer"),
            "rule_validation": agent_data.get("rule_validation", {}),
            "evaluation_result": agent_data.get("evaluation_result", {}),
        })
    return flat


def _overall_stats(evaluation_results: list) -> dict:
    total = len(evaluation_results)
    pass_n = sum(1 for r in evaluation_results if r["evaluation_result"]["overall_decision"] == "PASS")
    review_n = sum(1 for r in evaluation_results if r["evaluation_result"]["overall_decision"] == "REVIEW")
    fail_n = total - pass_n - review_n
    success_rate = (pass_n / total * 100) if total else 0.0
    return {
        "total": total, "pass": pass_n, "review": review_n, "fail": fail_n,
        "success_rate": success_rate, "threshold": ITEM_PASS_THRESHOLD,
        "verdict": "합격" if success_rate >= ITEM_PASS_THRESHOLD else "불합격",
    }


def _metric_stats(evaluation_results: list) -> dict:
    total = len(evaluation_results)
    stats = {}
    for key, label in METRIC_LABELS.items():
        scores = [r["evaluation_result"][key]["score"] for r in evaluation_results]
        pass_n = sum(1 for s in scores if s >= METRIC_PASS_SCORE)
        fail_n = total - pass_n
        score_pct = (pass_n / total * 100) if total else 0.0
        stats[key] = {
            "key": key, "label": label, "total": total, "pass": pass_n, "fail": fail_n,
            "score": score_pct, "threshold": ITEM_PASS_THRESHOLD, "met": score_pct >= ITEM_PASS_THRESHOLD,
        }
    return stats


def _derive_severity(case: dict) -> str:
    category = case.get("category", "")
    eval_res = case["evaluation_result"]
    rule_val = case.get("rule_validation", {})

    if "안전성" in category or "위험" in category:
        return "Critical"
    if rule_val.get("rule_status") == "FAIL":
        return "High"
    if eval_res.get("groundedness", {}).get("score", 5) <= 3:
        return "Medium"
    return "Low"


def _weakest_metric(case: dict):
    eval_res = case["evaluation_result"]
    scored = [(key, eval_res[key]["score"], eval_res[key]["reason"]) for key in METRIC_LABELS]
    return min(scored, key=lambda x: x[1])


def _build_bug_list(evaluation_results: list, id_prefix: str = "BUG") -> list:
    bugs = []
    for case in evaluation_results:
        if case["evaluation_result"]["overall_decision"] != "FAIL":
            continue
        severity = _derive_severity(case)
        weak_key, _weak_score, weak_reason = _weakest_metric(case)
        summary = case["evaluation_result"].get("summary", "기준 미달 응답")
        bugs.append({
            "bug_id": f"{id_prefix}-{len(bugs) + 1:03d}",
            "case_id": case["case_id"],
            "category": case["category"],
            "test_type": case.get("test_type", ""),
            "title": f"[{case['category']}] {summary[:40]}",
            "severity": severity,
            "priority": SEVERITY_PRIORITY_MAP[severity],
            "user_question": case["user_question"],
            "ai_answer": case["ai_answer"],
            "rule_reason": case.get("rule_validation", {}).get("rule_reason", ""),
            "summary": summary,
            "weak_metric_label": METRIC_LABELS[weak_key],
            "weak_metric_reason": weak_reason,
            "impact_text": SEVERITY_IMPACT_TEXT[severity],
        })
    return bugs


def build_report_context(evaluation_results: list, id_prefix: str = "BUG") -> dict:
    """DOCX/PDF 생성기가 공통으로 사용하는 통계/차트용 데이터를 한 번에 계산합니다. (단일 챗봇 기준)"""
    overview = _overall_stats(evaluation_results)
    metrics = _metric_stats(evaluation_results)
    bugs = _build_bug_list(evaluation_results, id_prefix=id_prefix)
    bug_id_by_case = {b["case_id"]: b["bug_id"] for b in bugs}

    type_stats = {}
    for t in TEST_TYPES:
        cases = [r for r in evaluation_results if r.get("test_type") == t]
        type_stats[t] = {
            "total": len(cases),
            "pass": sum(1 for r in cases if r["evaluation_result"]["overall_decision"] == "PASS"),
            "fail": sum(1 for r in cases if r["evaluation_result"]["overall_decision"] != "PASS"),
        }

    return {
        "generated_date": date.today().strftime("%Y. %m. %d"),
        "overview": overview,
        "metrics": metrics,
        "bugs": bugs,
        "bug_id_by_case": bug_id_by_case,
        "type_stats": type_stats,
        "cases": evaluation_results,
        "python_version": platform.python_version(),
        "os_name": platform.system(),
    }


def _comparison_rows(pipeline_outputs: list) -> list:
    """케이스별 규칙기반 판정 vs API기반 판정 비교 표 데이터를 만듭니다."""
    rows = []
    for case in pipeline_outputs:
        rule_decision = case.get("rule_based", {}).get("evaluation_result", {}).get("overall_decision", "FAIL")
        api_decision = case.get("api_based", {}).get("evaluation_result", {}).get("overall_decision", "FAIL")
        rows.append({
            "case_id": case.get("case_id"),
            "category": case.get("category"),
            "test_type": case.get("test_type"),
            "rule_decision": rule_decision,
            "api_decision": api_decision,
            "match": rule_decision == api_decision,
        })
    return rows


def _overview_rows(context: dict) -> list:
    return [
        ("테스트명", "AI 교육과정 안내 챗봇 규칙 기반 vs API 기반 비교 테스트"),
        ("테스트 대상", "① 규칙 기반 챗봇(rule_based_agent.py, API 미사용) ② API 기반 챗봇(service_agent.py + ChromaDB RAG)"),
        ("테스트 목적", "동일한 테스트 케이스에 대해 두 챗봇의 응답 정확성·근거성·유용성·안전성을 동일 기준으로 비교 검증"),
        ("테스트 수행자", "AI 품질관리 자동화 파이프라인 (Rule Validator + Judge Agent)"),
        ("테스트 방법", "Happy Case, Edge Case, Negative Case 기반 Judge Agent 자동 평가 (두 챗봇 모두 동일 케이스/동일 기준 적용)"),
        ("테스트 환경", f"Python {context['python_version']} / OS: {context['os_name']} / OpenAI gpt-4o-mini + ChromaDB RAG (API 기반) / 키워드 매칭 (규칙 기반)"),
        ("테스트 일자", context["generated_date"]),
        ("평가 도구", "Judge Agent (LLM 기반 자동 판정): 정확성/근거성/유용성/안전성 4대 지표 채점, 규칙 기반 1차 키워드 검증 병행"),
    ]


def _bug_detail_rows(bug: dict, generated_date: str) -> list:
    return [
        ("결함 ID", bug["bug_id"]),
        ("테스트 케이스", bug["case_id"]),
        ("발견 일자", generated_date),
        ("심각도", bug["severity"]),
        ("우선순위", bug["priority"]),
        ("결함 상태", "Open"),
        ("결함 개요", f"사용자 질문 \"{bug['user_question']}\"에 대해 {bug['summary']}"),
        ("테스트 환경", f"유형: {bug['test_type']}"),
        ("입력 데이터", bug["user_question"]),
        ("기대 결과", bug["rule_reason"] or "기준정보에 근거한 정확한 안내"),
        ("실제 결과", bug["ai_answer"]),
        ("영향도 분석", bug["impact_text"]),
        ("원인 추정 및 개선 방안",
         f"'{bug['weak_metric_label']}' 지표 점수가 가장 낮았습니다 ({bug['weak_metric_reason']}). "
         f"프롬프트/규칙 및 지식 베이스 문서를 보강한 뒤 재테스트가 필요합니다."),
        ("최종 의견", f"{bug['severity']} 등급 결함으로, 재발 방지를 위한 보강 후 회귀 테스트를 권고합니다."),
    ]


def _final_summary_text(rule_context: dict, api_context: dict) -> str:
    rule_ov, api_ov = rule_context["overview"], api_context["overview"]
    rule_weak = min(rule_context["metrics"].values(), key=lambda m: m["score"])
    api_weak = min(api_context["metrics"].values(), key=lambda m: m["score"])
    better = "규칙 기반 챗봇" if rule_ov["success_rate"] > api_ov["success_rate"] else (
        "API 기반 챗봇" if api_ov["success_rate"] > rule_ov["success_rate"] else "두 챗봇이 동률로"
    )
    return (
        f"규칙 기반 챗봇은 전체 {rule_ov['total']}개 중 {rule_ov['pass']}개 통과({rule_ov['success_rate']:.0f}%), "
        f"API 기반 챗봇은 {api_ov['pass']}개 통과({api_ov['success_rate']:.0f}%)로 측정되어 "
        f"{better} 더 높은 합격률을 기록했습니다. "
        f"규칙 기반 챗봇은 '{rule_weak['label']}' 항목({rule_weak['score']:.1f}점), "
        f"API 기반 챗봇은 '{api_weak['label']}' 항목({api_weak['score']:.1f}점)이 가장 낮게 측정되어 우선 개선이 필요합니다. "
        f"일반적으로 규칙 기반 챗봇은 사전에 정의되지 않은 질문(Edge/Negative 케이스)에 취약하고, "
        f"API 기반 챗봇은 지식 베이스에 없는 내용을 사실처럼 답변(할루시네이션)할 위험이 있으므로 "
        f"각 방식의 한계를 보완하는 방향으로 개선할 것을 권고합니다."
    )


# ---------------------------------------------------------------------------
# 2. 차트 생성 (matplotlib)
# ---------------------------------------------------------------------------

def _chart_metric_bar(metrics: dict, path: Path, title_prefix: str = "") -> None:
    labels = [m["label"] for m in metrics.values()]
    scores = [m["score"] for m in metrics.values()]
    bar_colors = [CHART_PASS_COLOR if m["met"] else CHART_FAIL_COLOR for m in metrics.values()]

    fig, ax = plt.subplots(figsize=(7, 3.2))
    bars = ax.barh(labels, scores, color=bar_colors)
    ax.axvline(ITEM_PASS_THRESHOLD, color=CHART_THRESHOLD_COLOR, linestyle="--", linewidth=1.2,
               label=f"합격 기준 {ITEM_PASS_THRESHOLD}점")
    ax.set_xlim(0, 115)
    ax.set_xlabel("점수 (0~100)")
    ax.set_title(f"{title_prefix}품질 항목별 점수 vs 합격 기준")
    for bar, score in zip(bars, scores):
        ax.text(bar.get_width() + 2, bar.get_y() + bar.get_height() / 2, f"{score:.1f}점", va="center", fontsize=9)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _chart_metric_radar(metrics: dict, path: Path, title_prefix: str = "") -> None:
    labels = [m["label"] for m in metrics.values()]
    scores = [m["score"] for m in metrics.values()]
    thresholds = [m["threshold"] for m in metrics.values()]
    n = len(labels)
    angles = [i / n * 2 * math.pi for i in range(n)]
    angles += angles[:1]
    scores_plot = scores + scores[:1]
    thresholds_plot = thresholds + thresholds[:1]

    fig, ax = plt.subplots(figsize=(5.5, 5.5), subplot_kw=dict(polar=True))
    ax.plot(angles, scores_plot, color=CHART_PASS_COLOR, linewidth=2, marker="o", label="실제 점수")
    ax.fill(angles, scores_plot, color=CHART_PASS_COLOR, alpha=0.15)
    ax.plot(angles, thresholds_plot, color=CHART_THRESHOLD_COLOR, linewidth=1.5, linestyle="--", marker="o", label="합격 기준")
    ax.set_xticks(angles[:-1])
    ax.set_xticklabels(labels)
    ax.set_ylim(0, 100)
    ax.set_title(f"{title_prefix}품질 항목 점수 레이더 차트", pad=20)
    ax.legend(loc="upper right", bbox_to_anchor=(1.35, 1.1), fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _chart_distribution(context: dict, path: Path, title_prefix: str = "") -> None:
    type_stats = context["type_stats"]
    overview = context["overview"]

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(9, 3.6))

    types = list(type_stats.keys())
    pass_counts = [type_stats[t]["pass"] for t in types]
    fail_counts = [type_stats[t]["fail"] for t in types]
    x = range(len(types))
    ax1.bar(x, pass_counts, color=CHART_PASS_COLOR, label="Pass")
    ax1.bar(x, fail_counts, bottom=pass_counts, color=CHART_FAIL_COLOR, label="Fail/Review")
    ax1.set_xticks(list(x))
    ax1.set_xticklabels(types)
    ax1.set_title(f"{title_prefix}유형별 Pass/Fail 분포")
    ax1.legend(fontsize=8)
    for i, t in enumerate(types):
        total = type_stats[t]["total"]
        ax1.text(i, total + 0.3, f"합계 {total}건", ha="center", fontsize=8)

    remaining = max(overview["total"] - overview["pass"], 0)
    sizes = [overview["pass"], remaining] if overview["total"] else [1, 0]
    ax2.pie(
        sizes, colors=[CHART_PASS_COLOR, CHART_FAIL_COLOR], startangle=90,
        wedgeprops=dict(width=0.4),
        autopct=lambda p: f"{p:.0f}%" if p > 0 else "",
    )
    ax2.set_title(f"{title_prefix}전체 TC Pass 비율")
    ax2.text(0, 0, f"{overview['pass']} / {overview['total']}\nPass", ha="center", va="center",
              fontsize=11, fontweight="bold")

    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _chart_agent_comparison(rule_overview: dict, api_overview: dict, path: Path) -> None:
    """규칙 기반 vs API 기반 챗봇의 Pass/Review/Fail 건수를 나란히 비교하는 그룹 막대 차트."""
    labels = ["Pass", "Review", "Fail"]
    rule_vals = [rule_overview["pass"], rule_overview["review"], rule_overview["fail"]]
    api_vals = [api_overview["pass"], api_overview["review"], api_overview["fail"]]

    x = range(len(labels))
    width = 0.35
    fig, ax = plt.subplots(figsize=(7, 3.6))
    ax.bar([i - width / 2 for i in x], rule_vals, width, color="#2563EB", label="규칙 기반 챗봇")
    ax.bar([i + width / 2 for i in x], api_vals, width, color="#93C5FD", label="API 기반 챗봇")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels)
    ax.set_title("규칙 기반 vs API 기반 챗봇 — 판정 결과 비교")
    ax.legend(fontsize=9)
    for i, v in enumerate(rule_vals):
        ax.text(i - width / 2, v + 0.1, str(v), ha="center", fontsize=9)
    for i, v in enumerate(api_vals):
        ax.text(i + width / 2, v + 0.1, str(v), ha="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _chart_hero_kpi(rule_context: dict, api_context: dict, path: Path) -> Path:
    """보고서 상단 '히어로' KPI 카드 행(반원 게이지 포함)을 하나의 이미지로 렌더링한다.
    우리 데이터(합격률·지표 충족률·결함 수) 기준. 이미지라 HTML/DOCX/PDF에 동일하게 삽입 가능."""
    from matplotlib.patches import FancyBboxPatch, Wedge

    r_ov, a_ov = rule_context["overview"], api_context["overview"]
    metric_scores = [m["score"] for m in rule_context["metrics"].values()] + \
                    [m["score"] for m in api_context["metrics"].values()]
    avg_metric = sum(metric_scores) / len(metric_scores) if metric_scores else 0.0
    fail_n = r_ov["fail"] + a_ov["fail"]

    cards = [
        {"label": "규칙 기반 합격률", "big": f"{r_ov['success_rate']:.0f}%", "sub": f"{r_ov['pass']}/{r_ov['total']} PASS", "frac": r_ov["success_rate"] / 100, "color": "#1594AE", "gauge": True},
        {"label": "API 기반 합격률", "big": f"{a_ov['success_rate']:.0f}%", "sub": f"{a_ov['pass']}/{a_ov['total']} PASS", "frac": a_ov["success_rate"] / 100, "color": "#0E7285", "gauge": True},
        {"label": "평균 지표 충족률", "big": f"{avg_metric:.0f}%", "sub": "4대 지표 평균", "frac": avg_metric / 100, "color": "#5FBFA0", "gauge": True},
        {"label": "결함(FAIL)", "big": f"{fail_n}건", "sub": "조치 대상" if fail_n else "결함 없음 · 안정", "frac": None, "color": "#E28585" if fail_n else "#5FBFA0", "gauge": False},
    ]

    fig, axes = plt.subplots(1, 4, figsize=(13, 2.9))
    for ax, c in zip(axes, cards):
        ax.axis("off")
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        # 카드 배경(둥근 모서리) + 상단 액센트 바
        ax.add_patch(FancyBboxPatch((0.05, 0.10), 0.90, 0.82, boxstyle="round,pad=0,rounding_size=0.05",
                                    facecolor="white", edgecolor="#e2e8f0", linewidth=1.3))
        ax.add_patch(FancyBboxPatch((0.05, 0.855), 0.90, 0.065, boxstyle="round,pad=0,rounding_size=0.02",
                                    facecolor=c["color"], edgecolor="none"))
        ax.text(0.11, 0.70, c["label"], fontsize=11, color="#475569", ha="left", va="center", fontweight="bold")
        ax.text(0.11, 0.42, c["big"], fontsize=27, color=c["color"], ha="left", va="center", fontweight="bold")
        ax.text(0.11, 0.22, c["sub"], fontsize=8.5, color="#94a3b8", ha="left", va="center")
        if c["gauge"]:
            cx, cy, r = 0.80, 0.40, 0.135
            ax.add_patch(Wedge((cx, cy), r, 0, 180, width=0.05, facecolor="#e6eef0"))
            ax.add_patch(Wedge((cx, cy), r, 180 - 180 * c["frac"], 180, width=0.05, facecolor=c["color"]))
            ax.text(cx, cy - 0.02, f"{c['frac'] * 100:.0f}", fontsize=10, color=c["color"], ha="center", va="center", fontweight="bold")
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return path


def _chart_hero_compare(pipeline_outputs: list, path: Path) -> Path:
    """규칙 기반 vs API 기반의 4대 지표 평균(5점)을 큰 가로 막대로 비교하는 히어로 차트."""
    labels = ["정확성", "근거성", "유용성", "안전성"]
    keys = ["accuracy", "groundedness", "helpfulness", "safety"]

    def avg(agent_key, metric_key):
        vals = [c.get(agent_key, {}).get("evaluation_result", {}).get(metric_key, {}).get("score", 0) for c in pipeline_outputs]
        return sum(vals) / len(vals) if vals else 0.0

    rule_vals = [avg("rule_based", k) for k in keys]
    api_vals = [avg("api_based", k) for k in keys]
    y = list(range(len(labels)))
    h = 0.36

    fig, ax = plt.subplots(figsize=(11, 3.1))
    b1 = ax.barh([i - h / 2 for i in y], rule_vals, height=h, color="#1594AE", label="규칙 기반")
    b2 = ax.barh([i + h / 2 for i in y], api_vals, height=h, color="#7FC9C4", label="API 기반")
    ax.set_yticks(y)
    ax.set_yticklabels(labels)
    ax.invert_yaxis()
    ax.set_xlim(0, 5.6)
    ax.set_xlabel("평균 점수 (5점 만점)")
    ax.set_title("규칙 기반 vs API 기반 — 4대 지표 평균 비교", fontsize=12)
    ax.grid(axis="x", alpha=0.25)
    for bars, vals in ((b1, rule_vals), (b2, api_vals)):
        for bar, v in zip(bars, vals):
            ax.text(v + 0.08, bar.get_y() + bar.get_height() / 2, f"{v:.2f}", va="center", fontsize=9, color="#334155")
    ax.legend(loc="lower right", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return path


def _build_charts(context: dict, tmp_dir: Path, prefix: str, title_prefix: str = "") -> dict:
    paths = {
        "bar": tmp_dir / f"{prefix}_bar.png",
        "radar": tmp_dir / f"{prefix}_radar.png",
        "dist": tmp_dir / f"{prefix}_dist.png",
    }
    _chart_metric_bar(context["metrics"], paths["bar"], title_prefix=title_prefix)
    _chart_metric_radar(context["metrics"], paths["radar"], title_prefix=title_prefix)
    _chart_distribution(context, paths["dist"], title_prefix=title_prefix)
    return paths


# ---------------------------------------------------------------------------
# 3. DOCX 보고서 생성
# ---------------------------------------------------------------------------

def _shade_cell(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text, bold=False, color=None, align=None, size=10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_kv_table(doc: Document, rows: list) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        row = table.add_row()
        _set_cell_text(row.cells[0], label, bold=True)
        _shade_cell(row.cells[0], "E9F7FB")
        _set_cell_text(row.cells[1], value)


def _docx_agent_section(doc: Document, context: dict, section_no: str, label: str, chart_paths: dict) -> None:
    """규칙 기반/API 기반 챗봇 한쪽의 상세 결과(품질 항목별 점수/TC 기록표/결함 보고서)를 문서에 추가합니다."""
    doc.add_heading(f"{section_no}. {label} 상세 결과", level=1)

    doc.add_heading(f"{section_no}.1 품질 항목별 점수", level=2)
    doc.add_paragraph(f"※ 점수 = (Pass TC 수 / 해당 항목 전체 TC 수) × 100 | 합격 기준: 각 항목 {ITEM_PASS_THRESHOLD}점 이상")
    metric_headers = ["평가 항목", "전체 TC", "Pass", "Fail", "점수", "합격 기준 충족"]
    metric_table = doc.add_table(rows=1, cols=6)
    metric_table.style = "Table Grid"
    for c, h in enumerate(metric_headers):
        _set_cell_text(metric_table.rows[0].cells[c], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
        _shade_cell(metric_table.rows[0].cells[c], "1594AE")
    for m in context["metrics"].values():
        row = metric_table.add_row()
        vals = [m["label"], f"{m['total']} 개", f"{m['pass']} 개", f"{m['fail']} 개",
                f"{m['score']:.1f} 점", "✓ 충족" if m["met"] else "✗ 미충족"]
        for c, v in enumerate(vals):
            _set_cell_text(row.cells[c], v, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade_cell(row.cells[5], "E7F8F1" if m["met"] else "FDE8E8")

    doc.add_paragraph()
    doc.add_picture(str(chart_paths["bar"]), width=Cm(15))
    doc.add_picture(str(chart_paths["radar"]), width=Cm(11))

    doc.add_heading(f"{section_no}.2 TC 결과 기록표", level=2)
    tc_headers = ["TC ID", "유형", "평가 영역", "결과 요약", "충족여부", "판정", "결함 ID"]
    tc_table = doc.add_table(rows=1, cols=7)
    tc_table.style = "Table Grid"
    for c, h in enumerate(tc_headers):
        _set_cell_text(tc_table.rows[0].cells[c], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
        _shade_cell(tc_table.rows[0].cells[c], "1594AE")
    for case in context["cases"]:
        decision = case["evaluation_result"]["overall_decision"]
        row = tc_table.add_row()
        vals = [
            case["case_id"], case.get("test_type", ""), case["category"],
            case["evaluation_result"].get("summary", "")[:60],
            "Y" if decision == "PASS" else "N", decision,
            context["bug_id_by_case"].get(case["case_id"], ""),
        ]
        for c, v in enumerate(vals):
            align = WD_ALIGN_PARAGRAPH.LEFT if c == 3 else WD_ALIGN_PARAGRAPH.CENTER
            _set_cell_text(row.cells[c], v, align=align, size=9)
        _shade_cell(row.cells[5], DECISION_FILL_HEX.get(decision, "FFFFFF"))

    doc.add_paragraph()
    doc.add_picture(str(chart_paths["dist"]), width=Cm(16))

    if context["bugs"]:
        doc.add_heading(f"{section_no}.3 결함 보고서 (Bug Report)", level=2)
        for i, bug in enumerate(context["bugs"], 1):
            doc.add_heading(f"{section_no}.3.{i} {bug['bug_id']} — {bug['title']}", level=3)
            _add_kv_table(doc, _bug_detail_rows(bug, context["generated_date"]))
            doc.add_paragraph()


def _generate_docx_python(pipeline_outputs: list, output_path, run_dir=None) -> Path:
    rule_flat = _extract_agent_results(pipeline_outputs, "rule_based")
    api_flat = _extract_agent_results(pipeline_outputs, "api_based")
    rule_context = build_report_context(rule_flat, id_prefix="RULE")
    api_context = build_report_context(api_flat, id_prefix="API")
    comparison = _comparison_rows(pipeline_outputs)
    output_path = Path(output_path)

    with tempfile.TemporaryDirectory() as tmp:
        rule_charts = _build_charts(rule_context, Path(tmp), "rule", title_prefix="[규칙 기반] ")
        api_charts = _build_charts(api_context, Path(tmp), "api", title_prefix="[API 기반] ")
        comparison_chart = Path(tmp) / "comparison.png"
        _chart_agent_comparison(rule_context["overview"], api_context["overview"], comparison_chart)

        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "맑은 고딕"
        normal.font.size = Pt(10)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("AI 교육과정 안내 챗봇 품질 테스트")
        run.bold = True
        run.font.size = Pt(22)

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run("규칙 기반 vs API 기반 챗봇 비교 — 최종 테스트 결과 보고서")
        run.bold = True
        run.font.size = Pt(16)

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(f"작성일: {rule_context['generated_date']}    AI 기반 소프트웨어 QA 및 모니터링 실무 교육").font.size = Pt(10)
        doc.add_paragraph()

        # 상단 히어로 KPI 카드(게이지) + 규칙 vs API 4대 지표 비교 (HTML과 동일 이미지 → 3형식 통일)
        doc.add_picture(str(_chart_hero_kpi(rule_context, api_context, Path(tmp) / "hero.png")), width=Cm(17))
        doc.add_heading("규칙 기반 vs API 기반 — 4대 지표 비교", level=3)
        doc.add_picture(str(_chart_hero_compare(pipeline_outputs, Path(tmp) / "hero_compare.png")), width=Cm(16))
        doc.add_paragraph()

        doc.add_heading("1. 테스트 개요", level=1)
        _add_kv_table(doc, _overview_rows(rule_context))

        doc.add_heading("2. 전체 비교 결과 요약", level=1)
        doc.add_heading("2.1 챗봇 유형별 TC 수행 결과 통계", level=2)
        stat_headers = ["챗봇 유형", "전체 TC", "통과(Pass)", "실패/재검토", "성공률", "판정"]
        stat_table = doc.add_table(rows=1, cols=6)
        stat_table.style = "Table Grid"
        for c, h in enumerate(stat_headers):
            _set_cell_text(stat_table.rows[0].cells[c], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
            _shade_cell(stat_table.rows[0].cells[c], "0E7285")
        for label, ov in (("⚙️ 규칙 기반 챗봇", rule_context["overview"]), ("🤖 API 기반 챗봇", api_context["overview"])):
            row = stat_table.add_row()
            vals = [label, f"{ov['total']} 건", f"{ov['pass']} 건", f"{ov['fail'] + ov['review']} 건",
                    f"{ov['success_rate']:.0f}%", "✓ 합격" if ov["verdict"] == "합격" else "✗ 불합격"]
            for c, v in enumerate(vals):
                _set_cell_text(row.cells[c], v, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade_cell(row.cells[5], "E7F8F1" if ov["verdict"] == "합격" else "FDE8E8")

        doc.add_heading("2.2 케이스별 판정 비교표", level=2)
        cmp_headers = ["TC ID", "카테고리", "유형", "규칙기반 판정", "API기반 판정", "일치여부"]
        cmp_table = doc.add_table(rows=1, cols=6)
        cmp_table.style = "Table Grid"
        for c, h in enumerate(cmp_headers):
            _set_cell_text(cmp_table.rows[0].cells[c], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
            _shade_cell(cmp_table.rows[0].cells[c], "1594AE")
        for row_data in comparison:
            row = cmp_table.add_row()
            vals = [row_data["case_id"], row_data["category"], row_data["test_type"],
                    row_data["rule_decision"], row_data["api_decision"],
                    "🟢 일치" if row_data["match"] else "🔶 불일치"]
            for c, v in enumerate(vals):
                _set_cell_text(row.cells[c], v, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
            _shade_cell(row.cells[3], DECISION_FILL_HEX.get(row_data["rule_decision"], "FFFFFF"))
            _shade_cell(row.cells[4], DECISION_FILL_HEX.get(row_data["api_decision"], "FFFFFF"))

        doc.add_paragraph()
        doc.add_picture(str(comparison_chart), width=Cm(15))

        doc.add_page_break()
        _docx_agent_section(doc, rule_context, "3", "규칙 기반 챗봇", rule_charts)

        doc.add_page_break()
        _docx_agent_section(doc, api_context, "4", "API 기반 챗봇", api_charts)

        doc.add_page_break()
        doc.add_heading("5. 심각도 분류 기준", level=1)
        sev_table = doc.add_table(rows=1, cols=2)
        sev_table.style = "Table Grid"
        _set_cell_text(sev_table.rows[0].cells[0], "심각도 등급", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
        _set_cell_text(sev_table.rows[0].cells[1], "정의 및 기준", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
        _shade_cell(sev_table.rows[0].cells[0], "1594AE")
        _shade_cell(sev_table.rows[0].cells[1], "1594AE")
        for level in ["Critical", "High", "Medium", "Low"]:
            row = sev_table.add_row()
            _set_cell_text(row.cells[0], level, bold=True)
            _set_cell_text(row.cells[1], SEVERITY_IMPACT_TEXT[level])

        doc.add_heading("6. 최종 종합 의견 및 개선 권고", level=1)
        doc.add_paragraph(_final_summary_text(rule_context, api_context))
        all_bugs = [(b, "규칙 기반") for b in rule_context["bugs"]] + [(b, "API 기반") for b in api_context["bugs"]]
        if all_bugs:
            doc.add_paragraph("핵심 우선 개선 사항:")
            for bug, source in all_bugs:
                doc.add_paragraph(f"- [{source}] {bug['bug_id']} ({bug['case_id']}): {bug['title']}")

        doc.add_paragraph()
        _appendix_docx(doc, build_appendix_context(pipeline_outputs, Path(tmp), rule_context, api_context, run_dir))

        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run("본 보고서는 AI 품질관리 자동화 파이프라인이 자동 생성한 규칙 기반 vs API 기반 챗봇 비교 검증 결과입니다.")
        footer_run.italic = True

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))

    return output_path


# ---------------------------------------------------------------------------
# 4. PDF 보고서 생성 (reportlab)
# ---------------------------------------------------------------------------

def _p(text, style) -> Paragraph:
    return Paragraph(_xml_escape(str(text)), style)


def _pdf_image(path: Path, width) -> RLImage:
    with PILImage.open(path) as img:
        w, h = img.size
    return RLImage(str(path), width=width, height=width * h / w)


def _pdf_agent_section(context: dict, section_no: str, label: str, chart_paths: dict, styles: dict) -> list:
    """규칙 기반/API 기반 챗봇 한쪽의 상세 결과 flowable 목록을 만듭니다."""
    story = [_p(f"{section_no}. {label} 상세 결과", styles["h1"])]

    story.append(_p(f"{section_no}.1 품질 항목별 점수", styles["h2"]))
    story.append(_p(f"※ 점수 = (Pass TC 수 / 해당 항목 전체 TC 수) × 100 | 합격 기준: 각 항목 {ITEM_PASS_THRESHOLD}점 이상", styles["body"]))
    story.append(Spacer(1, 4))
    metric_headers = ["평가 항목", "전체 TC", "Pass", "Fail", "점수", "합격 기준 충족"]
    metric_data = [[_p(h, styles["cell_header"]) for h in metric_headers]]
    metric_cmds = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1594AE")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for idx, m in enumerate(context["metrics"].values(), start=1):
        metric_data.append([
            _p(m["label"], styles["cell_center"]), _p(f"{m['total']} 개", styles["cell_center"]),
            _p(f"{m['pass']} 개", styles["cell_center"]), _p(f"{m['fail']} 개", styles["cell_center"]),
            _p(f"{m['score']:.1f} 점", styles["cell_center"]),
            _p("✓ 충족" if m["met"] else "✗ 미충족", styles["cell_center"]),
        ])
        metric_cmds.append(("BACKGROUND", (5, idx), (5, idx), colors.HexColor("#e7f8f1" if m["met"] else "#fde8e8")))
    metric_table = Table(metric_data, colWidths=[3 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 3 * cm])
    metric_table.setStyle(TableStyle(metric_cmds))
    story.append(metric_table)
    story.append(Spacer(1, 10))
    story.append(_pdf_image(chart_paths["bar"], 15 * cm))
    story.append(Spacer(1, 6))
    story.append(_pdf_image(chart_paths["radar"], 10 * cm))

    story.append(PageBreak())
    story.append(_p(f"{section_no}.2 TC 결과 기록표", styles["h2"]))
    tc_headers = ["TC ID", "유형", "평가 영역", "결과 요약", "충족", "판정", "결함ID"]
    tc_data = [[_p(h, styles["cell_header"]) for h in tc_headers]]
    tc_cmds = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1594AE")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for idx, case in enumerate(context["cases"], start=1):
        decision = case["evaluation_result"]["overall_decision"]
        tc_data.append([
            _p(case["case_id"], styles["cell_center"]), _p(case.get("test_type", ""), styles["cell_center"]),
            _p(case["category"], styles["cell_center"]),
            _p(case["evaluation_result"].get("summary", "")[:70], styles["cell"]),
            _p("Y" if decision == "PASS" else "N", styles["cell_center"]),
            _p(decision, styles["cell_center"]),
            _p(context["bug_id_by_case"].get(case["case_id"], ""), styles["cell_center"]),
        ])
        fill = DECISION_FILL_HEX.get(decision, "FFFFFF")
        tc_cmds.append(("BACKGROUND", (5, idx), (5, idx), colors.HexColor(f"#{fill}")))
    tc_table = Table(
        tc_data, colWidths=[2 * cm, 1.8 * cm, 2.5 * cm, 6.2 * cm, 1.3 * cm, 1.7 * cm, 1.7 * cm],
        repeatRows=1,
    )
    tc_table.setStyle(TableStyle(tc_cmds))
    story.append(tc_table)
    story.append(Spacer(1, 10))
    story.append(_pdf_image(chart_paths["dist"], 16 * cm))

    if context["bugs"]:
        story.append(PageBreak())
        story.append(_p(f"{section_no}.3 결함 보고서 (Bug Report)", styles["h2"]))
        for i, bug in enumerate(context["bugs"], 1):
            story.append(_p(f"{section_no}.3.{i} {bug['bug_id']} — {bug['title']}", styles["h2"]))
            bug_data = [[_p(l, styles["cell"]), _p(str(v), styles["cell"])]
                        for l, v in _bug_detail_rows(bug, context["generated_date"])]
            bug_table = Table(bug_data, colWidths=[3.2 * cm, 12.8 * cm])
            bug_table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E9F7FB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(bug_table)
            story.append(Spacer(1, 10))

    return story


def _generate_pdf_reportlab(pipeline_outputs: list, output_path, run_dir=None) -> Path:
    rule_flat = _extract_agent_results(pipeline_outputs, "rule_based")
    api_flat = _extract_agent_results(pipeline_outputs, "api_based")
    rule_context = build_report_context(rule_flat, id_prefix="RULE")
    api_context = build_report_context(api_flat, id_prefix="API")
    comparison = _comparison_rows(pipeline_outputs)
    output_path = Path(output_path)
    font_name = _pdf_font_name()

    styles = {
        "title": ParagraphStyle("title", fontName=font_name, fontSize=22, alignment=1, leading=28),
        "subtitle": ParagraphStyle("subtitle", fontName=font_name, fontSize=15, alignment=1, leading=20, spaceBefore=6),
        "meta": ParagraphStyle("meta", fontName=font_name, fontSize=9, alignment=1,
                                textColor=colors.HexColor("#64748b"), spaceBefore=14),
        "h1": ParagraphStyle("h1", fontName=font_name, fontSize=14, spaceBefore=14, spaceAfter=8,
                              textColor=colors.HexColor("#1e293b")),
        "h2": ParagraphStyle("h2", fontName=font_name, fontSize=11.5, spaceBefore=10, spaceAfter=6,
                              textColor=colors.HexColor("#1594AE")),
        "body": ParagraphStyle("body", fontName=font_name, fontSize=9, leading=13),
        "cell": ParagraphStyle("cell", fontName=font_name, fontSize=8.5, leading=12),
        "cell_center": ParagraphStyle("cell_center", fontName=font_name, fontSize=8.5, leading=12, alignment=1),
        "cell_header": ParagraphStyle("cell_header", fontName=font_name, fontSize=8.5, leading=12, alignment=1,
                                       textColor=colors.white),
        "footer": ParagraphStyle("footer", fontName=font_name, fontSize=8, alignment=1,
                                  textColor=colors.HexColor("#94a3b8")),
    }

    with tempfile.TemporaryDirectory() as tmp:
        rule_charts = _build_charts(rule_context, Path(tmp), "rule", title_prefix="[규칙 기반] ")
        api_charts = _build_charts(api_context, Path(tmp), "api", title_prefix="[API 기반] ")
        comparison_chart = Path(tmp) / "comparison.png"
        _chart_agent_comparison(rule_context["overview"], api_context["overview"], comparison_chart)

        story = [
            Spacer(1, 3 * cm),
            _p("AI 교육과정 안내 챗봇 품질 테스트", styles["title"]),
            _p("규칙 기반 vs API 기반 챗봇 비교 — 최종 테스트 결과 보고서", styles["subtitle"]),
            _p(f"작성일: {rule_context['generated_date']}    AI 기반 소프트웨어 QA 및 모니터링 실무 교육", styles["meta"]),
            PageBreak(),
            # 상단 히어로 KPI 카드(게이지) + 규칙 vs API 4대 지표 비교 (HTML과 동일 이미지 → 3형식 통일)
            _pdf_image(_chart_hero_kpi(rule_context, api_context, Path(tmp) / "hero.png"), 16.5 * cm),
            Spacer(1, 8),
            _p("규칙 기반 vs API 기반 — 4대 지표 비교", styles["h2"]),
            _pdf_image(_chart_hero_compare(pipeline_outputs, Path(tmp) / "hero_compare.png"), 16 * cm),
            Spacer(1, 10),
            _p("1. 테스트 개요", styles["h1"]),
        ]

        overview_data = [[_p(l, styles["cell"]), _p(v, styles["cell"])] for l, v in _overview_rows(rule_context)]
        overview_table = Table(overview_data, colWidths=[3.5 * cm, 12.5 * cm])
        overview_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E9F7FB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(overview_table)

        story.append(_p("2. 전체 비교 결과 요약", styles["h1"]))
        story.append(_p("2.1 챗봇 유형별 TC 수행 결과 통계", styles["h2"]))
        stat_headers = ["챗봇 유형", "전체 TC", "통과(Pass)", "실패/재검토", "성공률", "판정"]
        stat_data = [[_p(h, styles["cell_header"]) for h in stat_headers]]
        stat_cmds = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E7285")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        for idx, (label, ov) in enumerate(
            (("⚙️ 규칙 기반 챗봇", rule_context["overview"]), ("🤖 API 기반 챗봇", api_context["overview"])), start=1
        ):
            stat_data.append([
                _p(label, styles["cell_center"]), _p(f"{ov['total']} 건", styles["cell_center"]),
                _p(f"{ov['pass']} 건", styles["cell_center"]), _p(f"{ov['fail'] + ov['review']} 건", styles["cell_center"]),
                _p(f"{ov['success_rate']:.0f}%", styles["cell_center"]),
                _p("✓ 합격" if ov["verdict"] == "합격" else "✗ 불합격", styles["cell_center"]),
            ])
            stat_cmds.append(("BACKGROUND", (5, idx), (5, idx), colors.HexColor("#e7f8f1" if ov["verdict"] == "합격" else "#fde8e8")))
        stat_table = Table(stat_data, colWidths=[3.5 * cm, 2.3 * cm, 2.5 * cm, 2.7 * cm, 2.3 * cm, 2.5 * cm])
        stat_table.setStyle(TableStyle(stat_cmds))
        story.append(stat_table)
        story.append(Spacer(1, 10))

        story.append(_p("2.2 케이스별 판정 비교표", styles["h2"]))
        cmp_headers = ["TC ID", "카테고리", "유형", "규칙기반 판정", "API기반 판정", "일치여부"]
        cmp_data = [[_p(h, styles["cell_header"]) for h in cmp_headers]]
        cmp_cmds = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1594AE")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        for idx, row_data in enumerate(comparison, start=1):
            cmp_data.append([
                _p(row_data["case_id"], styles["cell_center"]), _p(row_data["category"], styles["cell_center"]),
                _p(row_data["test_type"], styles["cell_center"]), _p(row_data["rule_decision"], styles["cell_center"]),
                _p(row_data["api_decision"], styles["cell_center"]),
                _p("🟢 일치" if row_data["match"] else "🔶 불일치", styles["cell_center"]),
            ])
            cmp_cmds.append(("BACKGROUND", (3, idx), (3, idx), colors.HexColor(f"#{DECISION_FILL_HEX.get(row_data['rule_decision'], 'FFFFFF')}")))
            cmp_cmds.append(("BACKGROUND", (4, idx), (4, idx), colors.HexColor(f"#{DECISION_FILL_HEX.get(row_data['api_decision'], 'FFFFFF')}")))
        cmp_table = Table(cmp_data, colWidths=[2 * cm, 2.5 * cm, 1.8 * cm, 2.8 * cm, 2.8 * cm, 2.4 * cm], repeatRows=1)
        cmp_table.setStyle(TableStyle(cmp_cmds))
        story.append(cmp_table)
        story.append(Spacer(1, 10))
        story.append(_pdf_image(comparison_chart, 15 * cm))

        story.append(PageBreak())
        story.extend(_pdf_agent_section(rule_context, "3", "규칙 기반 챗봇", rule_charts, styles))

        story.append(PageBreak())
        story.extend(_pdf_agent_section(api_context, "4", "API 기반 챗봇", api_charts, styles))

        story.append(PageBreak())
        story.append(_p("5. 심각도 분류 기준", styles["h1"]))
        sev_data = [[_p("심각도 등급", styles["cell_header"]), _p("정의 및 기준", styles["cell_header"])]]
        for level in ["Critical", "High", "Medium", "Low"]:
            sev_data.append([_p(level, styles["cell"]), _p(SEVERITY_IMPACT_TEXT[level], styles["cell"])])
        sev_table = Table(sev_data, colWidths=[3 * cm, 13 * cm])
        sev_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1594AE")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(sev_table)

        story.append(Spacer(1, 14))
        story.append(_p("6. 최종 종합 의견 및 개선 권고", styles["h1"]))
        story.append(_p(_final_summary_text(rule_context, api_context), styles["body"]))
        all_bugs = [(b, "규칙 기반") for b in rule_context["bugs"]] + [(b, "API 기반") for b in api_context["bugs"]]
        if all_bugs:
            story.append(Spacer(1, 6))
            story.append(_p("핵심 우선 개선 사항:", styles["body"]))
            for bug, source in all_bugs:
                story.append(_p(f"- [{source}] {bug['bug_id']} ({bug['case_id']}): {bug['title']}", styles["body"]))

        story.extend(_appendix_pdf(build_appendix_context(pipeline_outputs, Path(tmp), rule_context, api_context, run_dir), styles))

        story.append(Spacer(1, 20))
        story.append(_p("본 보고서는 AI 품질관리 자동화 파이프라인이 자동 생성한 규칙 기반 vs API 기반 챗봇 비교 검증 결과입니다.", styles["footer"]))

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc_tpl = SimpleDocTemplate(
            str(output_path), pagesize=A4,
            topMargin=1.8 * cm, bottomMargin=1.8 * cm, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        )
        doc_tpl.build(story)

    return output_path


# ---------------------------------------------------------------------------
# 5. HTML 보고서 생성 (대시보드 미리보기용 — DOCX/PDF와 동일한 구성)
# ---------------------------------------------------------------------------

def _img_tag(path, max_width: str = "100%") -> str:
    b64 = base64.b64encode(Path(path).read_bytes()).decode("utf-8")
    return f'<img src="data:image/png;base64,{b64}" style="max-width:{max_width}; display:block; margin:12px auto;" />'


def _html_table(headers: list, rows: list, cell_bg: dict = None) -> str:
    """cell_bg: {(row_idx, col_idx): '6자리 hex 색상'} 형태로 특정 셀 배경을 강조할 때 사용합니다."""
    cell_bg = cell_bg or {}
    thead = "".join(f"<th>{_html_escape(str(h))}</th>" for h in headers)
    body_rows = []
    for r_idx, row in enumerate(rows):
        cells = []
        for c_idx, val in enumerate(row):
            style = f' style="background:#{cell_bg[(r_idx, c_idx)]};"' if (r_idx, c_idx) in cell_bg else ""
            cells.append(f"<td{style}>{_html_escape(str(val))}</td>")
        body_rows.append(f"<tr>{''.join(cells)}</tr>")
    return f'<table class="rpt-table"><thead><tr>{thead}</tr></thead><tbody>{"".join(body_rows)}</tbody></table>'


def _html_kv_table(rows: list) -> str:
    body_rows = "".join(
        f"<tr><th>{_html_escape(str(k))}</th><td>{_html_escape(str(v))}</td></tr>" for k, v in rows
    )
    return f'<table class="rpt-kv"><tbody>{body_rows}</tbody></table>'


# 심각도 배지 색상 (결함 카드용)
_SEVERITY_HEX = {"Critical": "#dc2626", "High": "#ea580c", "Medium": "#d97706", "Low": "#16a34a"}


def _html_bug_card(bug: dict, generated_date: str) -> str:
    """결함 하나를 참고 템플릿 톤의 3열 카드(결함정보 / 문제내용 / 분석·조치)로 렌더링합니다."""
    def _field(k, v):
        return f'<div class="bug-field"><span class="k">{_html_escape(k)}</span>{_html_escape(str(v))}</div>'

    sev = bug.get("severity", "Low")
    sev_badge = f'<span class="bug-sev" style="background:{_SEVERITY_HEX.get(sev, "#64748b")};">{_html_escape(sev)}</span>'

    col1 = (
        f'<div class="bug-col"><div class="bug-colhead">결함 정보<small>Defect Info</small></div>'
        f'{_field("결함 ID", bug.get("bug_id", ""))}'
        f'{_field("발견 일자", generated_date)}'
        f'<div class="bug-field"><span class="k">심각도</span>{sev_badge}</div>'
        f'{_field("우선순위", bug.get("priority", ""))}'
        f'{_field("결함 상태", "Open")}</div>'
    )
    col2 = (
        f'<div class="bug-col"><div class="bug-colhead">문제 내용<small>Problem</small></div>'
        f'{_field("테스트 케이스", f"{bug.get('case_id','')} ({bug.get('category','')})")}'
        f'{_field("입력 데이터", bug.get("user_question", ""))}'
        f'{_field("기대 결과", bug.get("rule_reason") or "기준정보에 근거한 정확한 안내")}'
        f'{_field("실제 결과", bug.get("ai_answer", ""))}</div>'
    )
    col3 = (
        f'<div class="bug-col"><div class="bug-colhead">분석 · 조치<small>Analysis &amp; Action</small></div>'
        f'{_field("결함 개요", bug.get("summary", ""))}'
        f'{_field("영향도 분석", bug.get("impact_text", ""))}'
        f'{_field("원인·개선", f"'{bug.get('weak_metric_label','')}' 지표 취약 — 프롬프트/규칙·지식 보강 후 재테스트 필요")}'
        f'{_field("최종 의견", f"{sev} 등급 결함 — 보강 후 회귀 테스트 권고")}</div>'
    )
    return f'<div class="bug-card">{col1}{col2}{col3}</div>'


def _wrap_sections_in_cards(html: str) -> str:
    """조립된 HTML을 '항목별'로 흰 카드(.rpt-card)로 감싼다.
    - 각 소제목(h2/h3/h4) 블록을 개별 카드로 만든다(부제 h2 제외).
    - 내용 없이 제목만 있는 상위 제목(예: '2. 전체 비교')은 다음 항목 카드에 합쳐 빈 카드를 막는다.
    - 헤더/상단 히어로(첫 제목 이전)와 푸터/닫는 태그는 카드 밖에 둔다."""
    import re
    SENT = "@@SUBTITLE_H2@@"
    html = html.replace('<h2 class="rpt-subtitle"', SENT)  # 부제 h2는 카드 대상 제외

    tokens = re.split(r'(?=<h[2-4])', html)
    out, carry = [], ""
    for t in tokens:
        if not re.match(r'^<h[2-4]', t):        # 첫 조각(헤더/히어로 영역) 등
            out.append(carry + t)
            carry = ""
            continue
        fm = re.search(r'<p class="rpt-footer', t)   # 마지막 조각에 붙은 푸터 분리
        if fm:
            out.append(f'<div class="rpt-card">{carry + t[:fm.start()]}</div>{t[fm.start():]}')
            carry = ""
            continue
        # 제목만 있고 본문이 없는 상위 제목은 다음 카드로 이월
        after = re.sub(r'^\s*<h[2-4][^>]*>.*?</h[2-4]>\s*', '', t, count=1, flags=re.S)
        if after.strip() == "":
            carry += t
        else:
            out.append(f'<div class="rpt-card">{carry + t}</div>')
            carry = ""
    if carry:
        out.append(f'<div class="rpt-card">{carry}</div>')
    return "".join(out).replace(SENT, '<h2 class="rpt-subtitle"')


def _html_agent_section(context: dict, section_no: str, label: str, chart_paths: dict) -> str:
    parts = [f"<h2>{section_no}. {_html_escape(label)} 상세 결과</h2>"]

    parts.append(f"<h3>{section_no}.1 품질 항목별 점수</h3>")
    parts.append(
        f'<p class="rpt-note">※ 점수 = (Pass TC 수 / 해당 항목 전체 TC 수) × 100 | '
        f'합격 기준: 각 항목 {ITEM_PASS_THRESHOLD}점 이상</p>'
    )
    metric_rows, metric_bg = [], {}
    for i, m in enumerate(context["metrics"].values()):
        metric_rows.append([m["label"], f"{m['total']} 개", f"{m['pass']} 개", f"{m['fail']} 개",
                             f"{m['score']:.1f} 점", "✓ 충족" if m["met"] else "✗ 미충족"])
        metric_bg[(i, 5)] = "e7f8f1" if m["met"] else "fde8e8"
    parts.append(_html_table(["평가 항목", "전체 TC", "Pass", "Fail", "점수", "합격 기준 충족"], metric_rows, metric_bg))
    parts.append(_img_tag(chart_paths["bar"], "700px"))
    parts.append(_img_tag(chart_paths["radar"], "480px"))

    parts.append(f"<h3>{section_no}.2 TC 결과 기록표</h3>")
    tc_rows, tc_bg = [], {}
    for i, case in enumerate(context["cases"]):
        decision = case["evaluation_result"]["overall_decision"]
        tc_rows.append([
            case["case_id"], case.get("test_type", ""), case["category"],
            case["evaluation_result"].get("summary", "")[:80],
            "Y" if decision == "PASS" else "N", decision,
            context["bug_id_by_case"].get(case["case_id"], ""),
        ])
        tc_bg[(i, 5)] = DECISION_FILL_HEX.get(decision, "FFFFFF")
    parts.append(_html_table(["TC ID", "유형", "평가 영역", "결과 요약", "충족여부", "판정", "결함 ID"], tc_rows, tc_bg))
    parts.append(_img_tag(chart_paths["dist"], "750px"))

    if context["bugs"]:
        parts.append(f"<h3>{section_no}.3 결함 보고서 (Bug Report)</h3>")
        for i, bug in enumerate(context["bugs"], 1):
            parts.append(f"<h4>{section_no}.3.{i} {_html_escape(bug['bug_id'])} — {_html_escape(bug['title'])}</h4>")
            parts.append(_html_bug_card(bug, context["generated_date"]))

    return "\n".join(parts)


def generate_html_report(pipeline_outputs: list, run_dir=None) -> str:
    """DOCX/PDF 보고서와 동일한 구성·내용의 정적 HTML 문자열을 생성합니다. (대시보드 미리보기용)"""
    rule_flat = _extract_agent_results(pipeline_outputs, "rule_based")
    api_flat = _extract_agent_results(pipeline_outputs, "api_based")
    rule_context = build_report_context(rule_flat, id_prefix="RULE")
    api_context = build_report_context(api_flat, id_prefix="API")
    comparison = _comparison_rows(pipeline_outputs)

    with tempfile.TemporaryDirectory() as tmp:
        rule_charts = _build_charts(rule_context, Path(tmp), "rule", title_prefix="[규칙 기반] ")
        api_charts = _build_charts(api_context, Path(tmp), "api", title_prefix="[API 기반] ")
        comparison_chart = Path(tmp) / "comparison.png"
        _chart_agent_comparison(rule_context["overview"], api_context["overview"], comparison_chart)

        parts = [f"""<!DOCTYPE html>
<html lang="ko"><head><meta charset="utf-8" />
<style>
  :root {{ --teal:#1594AE; --teal-dark:#0E7285; --teal-light:#E9F7FB; --ink:#1e293b; --muted:#64748b; --line:#d3e3e8; }}
  body {{ font-family: '맑은 고딕', 'Malgun Gothic', sans-serif; color:var(--ink); max-width:940px; margin:0 auto; padding:24px 24px 48px; background:#f2f7f9; }}
  /* 카드형 섹션 — 각 <h2> 섹션을 흰 카드로 감싼다(대시보드 컨셉) */
  .rpt-card {{ background:#fff; border:1px solid var(--line); border-radius:16px; padding:14px 26px 24px; margin:20px 0; box-shadow:0 3px 14px rgba(14,114,133,0.06); }}
  .rpt-card > h2:first-child, .rpt-card > h3:first-child {{ margin-top:8px; }}
  h1 {{ text-align:center; font-size:26px; margin-bottom:4px; color:var(--ink); }}
  h2.rpt-subtitle {{ text-align:center; font-size:17px; color:var(--muted); margin-top:0; font-weight:600; }}
  p.rpt-meta {{ text-align:center; color:var(--muted); font-size:13px; margin-bottom:28px; }}
  h2 {{ margin-top:34px; color:var(--teal-dark); border-bottom:3px solid var(--teal); padding-bottom:6px; }}
  h3 {{ margin-top:22px; color:var(--teal-dark); }}
  h4 {{ margin-top:16px; color:var(--ink); }}
  table.rpt-table, table.rpt-kv {{ border-collapse:collapse; width:100%; margin:10px 0; font-size:13px; }}
  table.rpt-table th, table.rpt-table td {{ border:1px solid var(--line); padding:7px 9px; text-align:center; }}
  table.rpt-table th {{ background:var(--teal); color:#fff; font-weight:700; letter-spacing:.2px; }}
  table.rpt-table tr:nth-child(even) td {{ background:#f6fbfc; }}
  table.rpt-kv th {{ background:var(--teal-light); color:var(--teal-dark); text-align:left; width:170px; border:1px solid var(--line); padding:7px 9px; vertical-align:top; font-weight:700; }}
  table.rpt-kv td {{ border:1px solid var(--line); padding:7px 9px; text-align:left; }}
  p.rpt-note {{ font-size:12px; color:var(--muted); }}
  .rpt-footer {{ text-align:center; color:#94a3b8; font-size:12px; margin-top:32px; font-style:italic; }}
  /* 결함 보고서 3열 카드 (템플릿 톤) */
  .bug-card {{ display:grid; grid-template-columns:1fr 1fr 1fr; border:1px solid var(--line); border-radius:8px; overflow:hidden; margin:12px 0 24px; }}
  .bug-col {{ border-right:1px solid var(--line); }}
  .bug-col:last-child {{ border-right:none; }}
  .bug-colhead {{ background:var(--teal); color:#fff; padding:9px 11px; font-weight:700; font-size:13px; }}
  .bug-colhead small {{ display:block; font-weight:400; font-size:11px; opacity:.92; }}
  .bug-field {{ padding:8px 11px; border-top:1px solid var(--line); font-size:12.5px; line-height:1.5; }}
  .bug-field .k {{ color:var(--teal-dark); font-weight:700; font-size:11px; letter-spacing:.3px; display:block; margin-bottom:2px; }}
  .bug-sev {{ display:inline-block; padding:2px 9px; border-radius:11px; font-size:11px; font-weight:700; color:#fff; }}
</style></head><body>
<h1>AI 교육과정 안내 챗봇 품질 테스트</h1>
<h2 class="rpt-subtitle">규칙 기반 vs API 기반 챗봇 비교 — 최종 테스트 결과 보고서</h2>
<p class="rpt-meta">작성일: {rule_context['generated_date']}&nbsp;&nbsp;&nbsp;&nbsp;AI 기반 소프트웨어 QA 및 모니터링 실무 교육</p>
"""]

        # 상단 히어로 KPI 카드(게이지) + 규칙 vs API 4대 지표 비교 — 대시보드 컨셉 요약
        parts.append('<div style="margin:16px 0 10px;">')
        parts.append(_img_tag(_chart_hero_kpi(rule_context, api_context, Path(tmp) / "hero.png"), "900px"))
        parts.append('</div>')
        # 아래 h3는 카드 후처리기가 자동으로 카드로 감싼다(항목별 카드).
        parts.append('<h3>규칙 기반 vs API 기반 — 4대 지표 비교</h3>')
        parts.append(_img_tag(_chart_hero_compare(pipeline_outputs, Path(tmp) / "hero_compare.png"), "880px"))

        parts.append("<h2>1. 테스트 개요</h2>")
        parts.append(_html_kv_table(_overview_rows(rule_context)))

        parts.append("<h2>2. 전체 비교 결과 요약</h2>")
        parts.append("<h3>2.1 챗봇 유형별 TC 수행 결과 통계</h3>")
        stat_rows, stat_bg = [], {}
        for i, (label, ov) in enumerate((
            ("⚙️ 규칙 기반 챗봇", rule_context["overview"]), ("🤖 API 기반 챗봇", api_context["overview"])
        )):
            stat_rows.append([label, f"{ov['total']} 건", f"{ov['pass']} 건", f"{ov['fail'] + ov['review']} 건",
                               f"{ov['success_rate']:.0f}%", "✓ 합격" if ov["verdict"] == "합격" else "✗ 불합격"])
            stat_bg[(i, 5)] = "e7f8f1" if ov["verdict"] == "합격" else "fde8e8"
        parts.append(_html_table(["챗봇 유형", "전체 TC", "통과(Pass)", "실패/재검토", "성공률", "판정"], stat_rows, stat_bg))

        parts.append("<h3>2.2 케이스별 판정 비교표</h3>")
        cmp_rows, cmp_bg = [], {}
        for i, row in enumerate(comparison):
            cmp_rows.append([row["case_id"], row["category"], row["test_type"],
                              row["rule_decision"], row["api_decision"], "🟢 일치" if row["match"] else "🔶 불일치"])
            cmp_bg[(i, 3)] = DECISION_FILL_HEX.get(row["rule_decision"], "FFFFFF")
            cmp_bg[(i, 4)] = DECISION_FILL_HEX.get(row["api_decision"], "FFFFFF")
        parts.append(_html_table(["TC ID", "카테고리", "유형", "규칙기반 판정", "API기반 판정", "일치여부"], cmp_rows, cmp_bg))
        parts.append(_img_tag(comparison_chart, "700px"))

        parts.append(_html_agent_section(rule_context, "3", "규칙 기반 챗봇", rule_charts))
        parts.append(_html_agent_section(api_context, "4", "API 기반 챗봇", api_charts))

        parts.append("<h2>5. 심각도 분류 기준</h2>")
        sev_rows = [[level, SEVERITY_IMPACT_TEXT[level]] for level in ["Critical", "High", "Medium", "Low"]]
        parts.append(_html_table(["심각도 등급", "정의 및 기준"], sev_rows))

        parts.append("<h2>6. 최종 종합 의견 및 개선 권고</h2>")
        parts.append(f"<p>{_html_escape(_final_summary_text(rule_context, api_context))}</p>")
        all_bugs = [(b, "규칙 기반") for b in rule_context["bugs"]] + [(b, "API 기반") for b in api_context["bugs"]]
        if all_bugs:
            parts.append("<p><b>핵심 우선 개선 사항:</b></p><ul>")
            for bug, source in all_bugs:
                parts.append(f"<li>[{_html_escape(source)}] {_html_escape(bug['bug_id'])} "
                              f"({_html_escape(bug['case_id'])}): {_html_escape(bug['title'])}</li>")
            parts.append("</ul>")

        parts.append(_appendix_html(build_appendix_context(pipeline_outputs, Path(tmp), rule_context, api_context, run_dir)))

        parts.append(
            '<p class="rpt-footer">본 보고서는 AI 품질관리 자동화 파이프라인이 자동 생성한 '
            '규칙 기반 vs API 기반 챗봇 비교 검증 결과입니다.</p>'
        )
        parts.append("</body></html>")

        return _wrap_sections_in_cards("\n".join(parts))


# ---------------------------------------------------------------------------
# 6. 부록 — 대시보드 전체 내용을 보고서 하단에 담기 (최종 결론 + 고도화 지표 + 운영/ k6 스냅샷)
#    * 비용 0원: 스냅샷(커버리지·레드티밍) + df 기반(PII·비용·환각) + Prometheus 조회(운영/k6).
#      검색품질/RAG on/off는 OpenAI 호출이 필요해 부록에서 제외하고 안내만 넣는다.
# ---------------------------------------------------------------------------

def _appendix_answer_items(pipeline_outputs: list) -> list:
    items = []
    for case in pipeline_outputs:
        for agent_key, agent_label in (("rule_based", "규칙기반"), ("api_based", "API기반")):
            ans = case.get(agent_key, {}).get("ai_answer", "") or ""
            items.append({
                "case_id": case.get("case_id"), "agent": agent_label,
                "text": ans, "answer": ans, "question": case.get("user_question", ""),
            })
    return items


def build_appendix_context(pipeline_outputs: list, tmp_dir: Path, rule_context: dict, api_context: dict, run_dir=None) -> dict:
    """부록에 필요한 데이터를 한 번에 모은다. (모두 비용 0원)

    run_dir: 조회 대상 실행 폴더. 그 실행에 '동결된 운영 스냅샷'(ops_snapshot/snapshot.json)이
             있으면 라이브 Prometheus 대신 저장된 스냅샷을 사용해 '그 당시 값'을 재현한다.
    """
    from datetime import datetime as _dt
    ctx = {"generated_at": _dt.now().strftime("%Y-%m-%d %H:%M:%S")}

    # (A) 최종 결론 요약
    def _agent_concl(context):
        ov = context["overview"]
        weak = min(context["metrics"].values(), key=lambda m: m["score"])
        return {"pass": ov["pass"], "total": ov["total"], "rate": ov["success_rate"],
                "verdict": ov["verdict"], "weak_label": weak["label"], "weak_score": weak["score"]}
    ctx["rule_concl"] = _agent_concl(rule_context)
    ctx["api_concl"] = _agent_concl(api_context)

    # (B) 고도화 지표 — 스냅샷 + df 기반 (비용 0원)
    try:
        from quality.enhancement_snapshot import load_snapshot
        from config import REPORTS_DIR as _RD
        snap = load_snapshot(_RD) or {}
    except Exception:
        snap = {}
    ctx["coverage"] = snap.get("coverage")
    ctx["redteam"] = snap.get("redteam")

    items = _appendix_answer_items(pipeline_outputs)
    try:
        from quality.pii_scan import scan_answers
        ctx["pii"] = scan_answers(items)
    except Exception as e:
        ctx["pii"] = {"error": str(e)}
    try:
        from quality.cost_tracker import track_cost
        ctx["cost"] = track_cost([{"case_id": f"{it['case_id']}·{it['agent']}", "question": it["question"], "answer": it["answer"]} for it in items])
    except Exception as e:
        ctx["cost"] = {"error": str(e)}
    try:
        from quality.hallucination_check import check_answers
        ctx["halluc"] = check_answers([{"case_id": it["case_id"], "agent": it["agent"], "answer": it["answer"]} for it in items])
    except Exception as e:
        ctx["halluc"] = {"error": str(e)}
    try:
        from quality.regression import compare_latest_two
        result, base_ts, cur_ts = compare_latest_two()
        ctx["regression"] = {"result": result, "base": base_ts, "cur": cur_ts}
    except Exception as e:
        ctx["regression"] = {"error": str(e)}

    # (C) 운영/k6 스냅샷: 저장된 동결 스냅샷이 있으면 그것을(과거 값 재현), 없으면 라이브 조회
    try:
        from quality.ops_snapshot import build_ops_snapshot, load_ops_snapshot
        frozen = load_ops_snapshot(run_dir)
        ctx["ops"] = frozen if frozen is not None else build_ops_snapshot(tmp_dir)
    except Exception as e:
        ctx["ops"] = {"charts": {}, "summary": {"prometheus_available": False, "error": str(e)}}
    return ctx


# ---- 부록 C(운영 모니터링) 공용 헬퍼 ----
_ALERT_STATE_LABEL = {
    "firing": "🔴 발동(Firing)",
    "pending": "🟡 대기(Pending)",
    "inactive": "🟢 정상(Normal)",
    "normal": "🟢 정상(Normal)",
}
# 부록 C 차트 순서/제목 (화면 운영 모니터링 탭과 동일 구성)
_OPS_CHART_ORDER = (
    ("golden", "운영 대시보드 (요청/오류율/응답시간)"),
    ("traffic", "트래픽 & 오류 (Rate / Errors)"),
    ("duration", "응답시간 분포 (Duration)"),
    ("k6", "k6 성능 테스트 결과"),
)


def _alert_label(state: str) -> str:
    return _ALERT_STATE_LABEL.get((state or "").lower(), state or "-")


def _fmt_metric(v, unit="", digits=2) -> str:
    """None이면 '데이터 없음', unit=='int'면 천단위 정수, 그 외 소수 digits자리 + 단위."""
    if v is None:
        return "데이터 없음"
    if unit == "int":
        return f"{int(v):,}"
    return f"{v:.{digits}f}{unit}"


def _ops_summary_rows(summary: dict):
    """운영 요약 지표를 (앱 지표 rows, k6 지표 rows)로 반환. 부록 C-2 표에 3형식 공통 사용."""
    app_rows = [
        ("캡처 시각", summary.get("captured_at", "-")),
        ("요청률(req/s)", _fmt_metric(summary.get("req_per_sec"), " req/s")),
        ("오류율", _fmt_metric(summary.get("error_rate_pct"), "%")),
        ("p95 응답시간", _fmt_metric(summary.get("p95_ms"), " ms", 1)),
    ]
    if summary.get("k6_total") is not None:
        k6_rows = [
            ("k6 총 요청 / 최대 VUs", f'{_fmt_metric(summary.get("k6_total"), unit="int")} 건 / {_fmt_metric(summary.get("k6_vus"), unit="int")} VUs'),
            ("k6 반복(iterations)", _fmt_metric(summary.get("k6_iterations"), unit="int") + " 회"),
            ("k6 응답시간 avg / p95 / p99", f'{_fmt_metric(summary.get("k6_avg_ms"), " ms", 1)} / {_fmt_metric(summary.get("k6_p95_ms"), " ms", 1)} / {_fmt_metric(summary.get("k6_p99_ms"), " ms", 1)}'),
            ("k6 오류율 / checks 성공률", f'{_fmt_metric(summary.get("k6_fail_pct"), "%")} / {_fmt_metric(summary.get("k6_checks_pct"), "%")}'),
        ]
    else:
        k6_rows = []
    return app_rows, k6_rows


# ---- HTML 부록 ----
def _appendix_html(ctx: dict) -> str:
    p = ['<div style="page-break-before:always;"></div>', "<h2>부록. 대시보드 전체 스냅샷</h2>",
         f'<p class="rpt-note">본 부록은 대시보드 화면(최종 결론·고도화 지표·운영 모니터링)을 보고서 생성 시점 '
         f'({_html_escape(ctx["generated_at"])}) 기준으로 담은 것입니다.</p>']

    p.append("<h3>A. 최종 결론</h3>")
    rc, ac = ctx["rule_concl"], ctx["api_concl"]
    p.append(_html_table(
        ["챗봇", "합격률", "판정", "가장 취약한 지표"],
        [["⚙️ 규칙 기반", f"{rc['rate']:.1f}% ({rc['pass']}/{rc['total']})", rc["verdict"], f"{rc['weak_label']} ({rc['weak_score']:.1f}점)"],
         ["🤖 API 기반", f"{ac['rate']:.1f}% ({ac['pass']}/{ac['total']})", ac["verdict"], f"{ac['weak_label']} ({ac['weak_score']:.1f}점)"]]))

    p.append("<h3>B. 고도화 지표</h3>")
    cov = ctx.get("coverage")
    if cov:
        risk = ", ".join(cov.get("uncovered_risk") or []) or "없음"
        p.append(f'<p><b>🗂️ 커버리지 갭:</b> {cov["coverage_pct"]}% ({cov["covered"]}/{cov["total"]}), '
                 f'미커버 갭 {cov["gap"]}개, 안전/위험 미커버: {_html_escape(risk)}</p>')
    rt = ctx.get("redteam")
    if rt:
        p.append(f'<p><b>🛡️ 레드티밍:</b> 방어율 {rt["defense_rate"]}% ({rt["attack_defended"]}/{rt["attack_total"]}), '
                 f'{"전체 통과" if rt.get("all_pass") else "취약점 발견"}</p>')
    pii = ctx.get("pii") or {}
    if "error" not in pii:
        p.append(f'<p><b>🔐 PII 검사:</b> {"개인정보 노출 없음 ✅" if pii.get("clean") else "노출 발견 ⚠️ " + str(len(pii.get("findings", []))) + "건"}</p>')
    cost = ctx.get("cost") or {}
    if "error" not in cost:
        p.append(f'<p><b>💰 비용 추적(추정):</b> 총 {cost.get("total_tokens", 0):,} 토큰, 약 {cost.get("total_cost_krw", 0):.2f}원 '
                 f'(평균 {cost.get("avg_tokens", 0)} 토큰/건)</p>')
    hc = ctx.get("halluc") or {}
    if "error" not in hc:
        p.append(f'<p><b>🧪 환각 검증:</b> 환각 의심 {hc.get("suspect_count", 0)}건 (지원율 임계 {hc.get("threshold")})</p>')
    reg = ctx.get("regression") or {}
    if "error" not in reg and reg.get("result"):
        r = reg["result"]
        p.append(f'<p><b>🔁 회귀 테스트:</b> {_html_escape(str(reg.get("base")))} → {_html_escape(str(reg.get("cur")))} | '
                 f'회귀(하락) {len(r.get("regressions", []))}건, 개선 {len(r.get("improvements", []))}건</p>')
    p.append('<p class="rpt-note">※ 검색 품질·RAG on/off는 OpenAI 호출이 필요해(비용 발생) 부록에는 포함하지 않았습니다. '
             '대시보드의 실시간 탭에서 확인하세요.</p>')

    p.append("<h3>C. 운영 모니터링 · 성능 스냅샷</h3>")
    ops = ctx.get("ops") or {}
    charts = ops.get("charts") or {}
    summary = ops.get("summary") or {}
    if summary.get("frozen"):
        p.append(f'<p class="rpt-note">📸 이 섹션은 <b>{_html_escape(str(summary.get("saved_at") or summary.get("captured_at") or ""))}</b> 시점에 '
                 f'저장된 스냅샷입니다(라이브가 아닌 <b>과거 값 고정</b>).</p>')
    if not summary.get("prometheus_available"):
        p.append('<p class="rpt-note">⚠️ Prometheus에 연결할 수 없어 운영/성능 스냅샷을 생성하지 못했습니다. '
                 '(docker compose로 모니터링 스택 실행 후 다시 생성하세요.)</p>')
    else:
        # C-1. 실시간 알림 상태 (Golden Signals) — 대시보드 알림 카드와 동일
        p.append("<h4>C-1. 실시간 알림 상태 (Golden Signals)</h4>")
        alerts = ops.get("alerts") or []
        if alerts:
            p.append(_html_table(
                ["알림", "상태", "심각도"],
                [[_html_escape(a["name"]), _alert_label(a["state"]), _html_escape(str(a.get("severity", "-")))] for a in alerts]))
        else:
            p.append('<p class="rpt-note">등록된 알림 규칙이 없거나 Grafana 알림 API에 접근하지 못했습니다.</p>')

        # C-2. 운영 요약 지표
        p.append("<h4>C-2. 운영 요약 지표</h4>")
        app_rows, k6_rows = _ops_summary_rows(summary)
        p.append(_html_table(["지표", "값"], [[k, _html_escape(str(v))] for k, v in (app_rows + k6_rows)]))
        if not k6_rows:
            p.append('<p class="rpt-note">※ k6 성능 지표는 부하테스트 실행 직후(now-30m 구간)에만 표시됩니다. '
                     'run_full.ps1 실행 직후 보고서를 생성하세요.</p>')

        # C-3~. 그래프: 실제 Grafana 패널이 있으면 항목별로 그대로 삽입(대시보드와 동일), 없으면 matplotlib 폴백
        grafana_panels = ops.get("grafana_panels") or []
        if grafana_panels:
            for idx, sec in enumerate(grafana_panels, start=3):
                p.append(f"<h4>C-{idx}. {_html_escape(sec['section'])}</h4>")
                for panel in sec["panels"]:
                    p.append(f'<p style="margin:8px 0 2px; font-weight:600;">{_html_escape(panel["title"])}</p>')
                    p.append(_img_tag(panel["path"], "760px"))
        else:
            for idx, (key, title) in enumerate(_OPS_CHART_ORDER, start=3):
                path = charts.get(key)
                p.append(f"<h4>C-{idx}. {_html_escape(title)}</h4>")
                if path:
                    p.append(_img_tag(path, "820px"))
                else:
                    p.append('<p class="rpt-note">해당 구간 데이터가 없습니다(No data). '
                             'k6 결과는 부하테스트 실행 직후에만 표시됩니다.</p>')
    return "\n".join(p)


# ---- DOCX 부록 ----
def _appendix_docx(doc: Document, ctx: dict) -> None:
    doc.add_page_break()
    doc.add_heading("부록. 대시보드 전체 스냅샷", level=1)
    doc.add_paragraph(f"본 부록은 대시보드 화면(최종 결론·고도화 지표·운영 모니터링)을 보고서 생성 시점"
                      f"({ctx['generated_at']}) 기준으로 담은 것입니다.")

    doc.add_heading("A. 최종 결론", level=2)
    rc, ac = ctx["rule_concl"], ctx["api_concl"]
    _add_kv_table(doc, [
        ("⚙️ 규칙 기반", f"합격률 {rc['rate']:.1f}% ({rc['pass']}/{rc['total']}) · {rc['verdict']} · 취약: {rc['weak_label']} {rc['weak_score']:.1f}점"),
        ("🤖 API 기반", f"합격률 {ac['rate']:.1f}% ({ac['pass']}/{ac['total']}) · {ac['verdict']} · 취약: {ac['weak_label']} {ac['weak_score']:.1f}점"),
    ])

    doc.add_heading("B. 고도화 지표", level=2)
    rows = []
    cov = ctx.get("coverage")
    if cov:
        risk = ", ".join(cov.get("uncovered_risk") or []) or "없음"
        rows.append(("🗂️ 커버리지 갭", f"{cov['coverage_pct']}% ({cov['covered']}/{cov['total']}), 미커버 {cov['gap']}개, 안전/위험 미커버: {risk}"))
    rt = ctx.get("redteam")
    if rt:
        rows.append(("🛡️ 레드티밍", f"방어율 {rt['defense_rate']}% ({rt['attack_defended']}/{rt['attack_total']}), {'전체 통과' if rt.get('all_pass') else '취약점 발견'}"))
    pii = ctx.get("pii") or {}
    if "error" not in pii:
        rows.append(("🔐 PII 검사", "개인정보 노출 없음" if pii.get("clean") else f"노출 {len(pii.get('findings', []))}건 발견"))
    cost = ctx.get("cost") or {}
    if "error" not in cost:
        rows.append(("💰 비용 추적(추정)", f"총 {cost.get('total_tokens', 0):,} 토큰 · 약 {cost.get('total_cost_krw', 0):.2f}원 · 평균 {cost.get('avg_tokens', 0)} 토큰/건"))
    hc = ctx.get("halluc") or {}
    if "error" not in hc:
        rows.append(("🧪 환각 검증", f"환각 의심 {hc.get('suspect_count', 0)}건 (임계 {hc.get('threshold')})"))
    reg = ctx.get("regression") or {}
    if "error" not in reg and reg.get("result"):
        r = reg["result"]
        rows.append(("🔁 회귀 테스트", f"{reg.get('base')} → {reg.get('cur')} | 회귀 {len(r.get('regressions', []))}건, 개선 {len(r.get('improvements', []))}건"))
    if rows:
        _add_kv_table(doc, rows)
    doc.add_paragraph("※ 검색 품질·RAG on/off는 OpenAI 호출이 필요해(비용 발생) 부록에서 제외했습니다. 대시보드 실시간 탭에서 확인하세요.")

    doc.add_heading("C. 운영 모니터링 · 성능 스냅샷", level=2)
    ops = ctx.get("ops") or {}
    charts = ops.get("charts") or {}
    summary = ops.get("summary") or {}
    if summary.get("frozen"):
        doc.add_paragraph(f'📸 이 섹션은 {summary.get("saved_at") or summary.get("captured_at") or ""} 시점에 저장된 스냅샷입니다(라이브가 아닌 과거 값 고정).')
    if not summary.get("prometheus_available"):
        doc.add_paragraph("⚠️ Prometheus에 연결할 수 없어 운영/성능 스냅샷을 생성하지 못했습니다. (docker compose로 모니터링 스택 실행 후 재생성)")
    else:
        # C-1. 실시간 알림 상태 (Golden Signals)
        doc.add_heading("C-1. 실시간 알림 상태 (Golden Signals)", level=3)
        alerts = ops.get("alerts") or []
        if alerts:
            _add_kv_table(doc, [(a["name"], f'{_alert_label(a["state"])} · 심각도 {a.get("severity", "-")}') for a in alerts])
        else:
            doc.add_paragraph("등록된 알림 규칙이 없거나 Grafana 알림 API에 접근하지 못했습니다.")

        # C-2. 운영 요약 지표
        doc.add_heading("C-2. 운영 요약 지표", level=3)
        app_rows, k6_rows = _ops_summary_rows(summary)
        _add_kv_table(doc, app_rows + k6_rows)
        if not k6_rows:
            doc.add_paragraph("※ k6 성능 지표는 부하테스트 실행 직후(now-30m 구간)에만 표시됩니다. run_full.ps1 실행 직후 보고서를 생성하세요.")

        # C-3~. 그래프: 실제 Grafana 패널 우선(대시보드와 동일), 없으면 matplotlib 폴백
        grafana_panels = ops.get("grafana_panels") or []
        if grafana_panels:
            for idx, sec in enumerate(grafana_panels, start=3):
                doc.add_heading(f"C-{idx}. {sec['section']}", level=3)
                for panel in sec["panels"]:
                    doc.add_paragraph(panel["title"]).runs[0].bold = True
                    doc.add_picture(str(panel["path"]), width=Cm(15))
        else:
            for idx, (key, title) in enumerate(_OPS_CHART_ORDER, start=3):
                path = charts.get(key)
                doc.add_heading(f"C-{idx}. {title}", level=3)
                if path:
                    doc.add_picture(str(path), width=Cm(16))
                else:
                    doc.add_paragraph("해당 구간 데이터가 없습니다(No data). k6 결과는 부하테스트 실행 직후에만 표시됩니다.")


# ---- PDF 부록 ----
def _appendix_pdf(ctx: dict, styles: dict) -> list:
    story = [PageBreak(), _p("부록. 대시보드 전체 스냅샷", styles["h1"]),
             _p(f"본 부록은 대시보드 화면(최종 결론·고도화 지표·운영 모니터링)을 보고서 생성 시점"
                f"({ctx['generated_at']}) 기준으로 담은 것입니다.", styles["body"])]

    story.append(_p("A. 최종 결론", styles["h2"]))
    rc, ac = ctx["rule_concl"], ctx["api_concl"]
    concl_data = [
        [_p("⚙️ 규칙 기반", styles["cell"]), _p(f"합격률 {rc['rate']:.1f}% ({rc['pass']}/{rc['total']}) · {rc['verdict']} · 취약: {rc['weak_label']} {rc['weak_score']:.1f}점", styles["cell"])],
        [_p("🤖 API 기반", styles["cell"]), _p(f"합격률 {ac['rate']:.1f}% ({ac['pass']}/{ac['total']}) · {ac['verdict']} · 취약: {ac['weak_label']} {ac['weak_score']:.1f}점", styles["cell"])],
    ]
    t = Table(concl_data, colWidths=[3.2 * cm, 12.8 * cm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                           ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E9F7FB")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t)

    story.append(_p("B. 고도화 지표", styles["h2"]))
    lines = []
    cov = ctx.get("coverage")
    if cov:
        risk = ", ".join(cov.get("uncovered_risk") or []) or "없음"
        lines.append(("🗂️ 커버리지 갭", f"{cov['coverage_pct']}% ({cov['covered']}/{cov['total']}), 미커버 {cov['gap']}개, 안전/위험 미커버: {risk}"))
    rt = ctx.get("redteam")
    if rt:
        lines.append(("🛡️ 레드티밍", f"방어율 {rt['defense_rate']}% ({rt['attack_defended']}/{rt['attack_total']}), {'전체 통과' if rt.get('all_pass') else '취약점 발견'}"))
    pii = ctx.get("pii") or {}
    if "error" not in pii:
        lines.append(("🔐 PII 검사", "개인정보 노출 없음" if pii.get("clean") else f"노출 {len(pii.get('findings', []))}건 발견"))
    cost = ctx.get("cost") or {}
    if "error" not in cost:
        lines.append(("💰 비용 추적(추정)", f"총 {cost.get('total_tokens', 0):,} 토큰 · 약 {cost.get('total_cost_krw', 0):.2f}원 · 평균 {cost.get('avg_tokens', 0)} 토큰/건"))
    hc = ctx.get("halluc") or {}
    if "error" not in hc:
        lines.append(("🧪 환각 검증", f"환각 의심 {hc.get('suspect_count', 0)}건 (임계 {hc.get('threshold')})"))
    reg = ctx.get("regression") or {}
    if "error" not in reg and reg.get("result"):
        r = reg["result"]
        lines.append(("🔁 회귀 테스트", f"{reg.get('base')} → {reg.get('cur')} | 회귀 {len(r.get('regressions', []))}건, 개선 {len(r.get('improvements', []))}건"))
    if lines:
        enh_data = [[_p(k, styles["cell"]), _p(v, styles["cell"])] for k, v in lines]
        et = Table(enh_data, colWidths=[3.6 * cm, 12.4 * cm])
        et.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E9F7FB")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(et)
    story.append(_p("※ 검색 품질·RAG on/off는 OpenAI 호출이 필요해(비용 발생) 부록에서 제외했습니다.", styles["body"]))

    story.append(_p("C. 운영 모니터링 · 성능 스냅샷", styles["h2"]))
    ops = ctx.get("ops") or {}
    charts = ops.get("charts") or {}
    summary = ops.get("summary") or {}

    def _kv_table(rows):
        data = [[_p(k, styles["cell"]), _p(v, styles["cell"])] for k, v in rows]
        tb = Table(data, colWidths=[4.6 * cm, 11.4 * cm])
        tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E9F7FB")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return tb

    if summary.get("frozen"):
        story.append(_p(f'📸 이 섹션은 {summary.get("saved_at") or summary.get("captured_at") or ""} 시점에 저장된 스냅샷입니다(라이브가 아닌 과거 값 고정).', styles["body"]))
    if not summary.get("prometheus_available"):
        story.append(_p("⚠️ Prometheus에 연결할 수 없어 운영/성능 스냅샷을 생성하지 못했습니다.", styles["body"]))
    else:
        # C-1. 실시간 알림 상태 (Golden Signals)
        story.append(_p("C-1. 실시간 알림 상태 (Golden Signals)", styles["h2"]))
        alerts = ops.get("alerts") or []
        if alerts:
            story.append(_kv_table([(a["name"], f'{_alert_label(a["state"])} · 심각도 {a.get("severity", "-")}') for a in alerts]))
        else:
            story.append(_p("등록된 알림 규칙이 없거나 Grafana 알림 API에 접근하지 못했습니다.", styles["body"]))

        # C-2. 운영 요약 지표
        story.append(_p("C-2. 운영 요약 지표", styles["h2"]))
        app_rows, k6_rows = _ops_summary_rows(summary)
        story.append(_kv_table(app_rows + k6_rows))
        if not k6_rows:
            story.append(_p("※ k6 성능 지표는 부하테스트 실행 직후(now-30m 구간)에만 표시됩니다. run_full.ps1 실행 직후 보고서를 생성하세요.", styles["body"]))

        # C-3~. 그래프: 실제 Grafana 패널 우선(대시보드와 동일), 없으면 matplotlib 폴백
        grafana_panels = ops.get("grafana_panels") or []
        if grafana_panels:
            for idx, sec in enumerate(grafana_panels, start=3):
                story.append(_p(f"C-{idx}. {sec['section']}", styles["h2"]))
                for panel in sec["panels"]:
                    story.append(_p(panel["title"], styles["body"]))
                    story.append(_pdf_image(panel["path"], 15 * cm))
        else:
            for idx, (key, title) in enumerate(_OPS_CHART_ORDER, start=3):
                path = charts.get(key)
                story.append(_p(f"C-{idx}. {title}", styles["h2"]))
                if path:
                    story.append(_pdf_image(path, 16 * cm))
                else:
                    story.append(_p("해당 구간 데이터가 없습니다(No data). k6 결과는 부하테스트 실행 직후에만 표시됩니다.", styles["body"]))
    return story


# =============================================================================
# HTML 디자인 그대로 PDF/DOCX 생성 (대시보드 HTML 미리보기와 100% 동일한 디자인으로 통일)
#   - PDF : 설치된 Chromium 계열 브라우저(Edge/Chrome) 헤드리스로 HTML→PDF 렌더
#   - DOCX: 위 PDF의 각 페이지를 이미지로 렌더해 Word에 삽입(디자인 동일, 이미지 기반)
#   - 브라우저가 없거나 실패하면 기존 reportlab/python-docx 구현으로 자동 폴백
# =============================================================================
def _find_chromium_browser():
    import os
    import shutil
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        shutil.which("msedge"), shutil.which("chrome"), shutil.which("chromium"),
        shutil.which("google-chrome"), shutil.which("chromium-browser"),
    ]
    return next((c for c in candidates if c and os.path.exists(c)), None)


def _render_html_to_pdf_via_browser(html: str, output_path) -> bool:
    """헤드리스 브라우저로 HTML을 PDF로 렌더한다. 성공 True, 브라우저 없음/실패 False."""
    import os
    import subprocess
    import tempfile
    browser = _find_chromium_browser()
    if not browser:
        return False
    tmpdir = tempfile.mkdtemp()
    html_path = os.path.join(tmpdir, "report.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    out = os.path.abspath(str(output_path))
    url = "file:///" + html_path.replace("\\", "/")
    try:
        subprocess.run(
            [browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
             "--virtual-time-budget=8000", f"--print-to-pdf={out}", url],
            timeout=120, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
    except Exception:
        return False
    return os.path.exists(out) and os.path.getsize(out) > 1000


def generate_pdf_report(pipeline_outputs: list, output_path, run_dir=None) -> Path:
    """대시보드 HTML 디자인 그대로 PDF를 생성한다(브라우저 렌더). 실패 시 reportlab로 폴백."""
    try:
        html = generate_html_report(pipeline_outputs, run_dir=run_dir)
        if _render_html_to_pdf_via_browser(html, output_path):
            return Path(output_path)
    except Exception:
        pass
    return _generate_pdf_reportlab(pipeline_outputs, output_path, run_dir)


def generate_docx_report(pipeline_outputs: list, output_path, run_dir=None) -> Path:
    """대시보드 HTML 디자인 그대로 DOCX를 생성한다(HTML→PDF→페이지 이미지 삽입). 실패 시 python-docx로 폴백."""
    import os
    import tempfile
    try:
        html = generate_html_report(pipeline_outputs, run_dir=run_dir)
        tmp_pdf = os.path.join(tempfile.mkdtemp(), "report.pdf")
        if _render_html_to_pdf_via_browser(html, tmp_pdf):
            import fitz  # pymupdf
            from docx import Document as _Doc
            from docx.shared import Inches
            doc = _Doc()
            sec = doc.sections[0]
            sec.top_margin = sec.bottom_margin = Inches(0.3)
            sec.left_margin = sec.right_margin = Inches(0.3)
            usable_w = sec.page_width - sec.left_margin - sec.right_margin
            pdf = fitz.open(tmp_pdf)
            for i, page in enumerate(pdf):
                pix = page.get_pixmap(dpi=150)
                img_path = os.path.join(os.path.dirname(tmp_pdf), f"page_{i}.png")
                pix.save(img_path)
                if i > 0:
                    doc.add_page_break()
                doc.add_picture(img_path, width=usable_w)
            pdf.close()
            doc.save(str(output_path))
            return Path(output_path)
    except Exception:
        pass
    return _generate_docx_python(pipeline_outputs, output_path, run_dir)


# 모듈 독립 실행 테스트: reports/evaluation_result.json을 읽어 DOCX/PDF를 생성합니다.
if __name__ == "__main__":
    import json
    from config import REPORTS_DIR

    result_file = REPORTS_DIR / "evaluation_result.json"
    if not result_file.exists():
        raise SystemExit(f"[Error] {result_file} 이 없습니다. 먼저 python main.py 를 실행해 결과를 생성하세요.")

    results = json.loads(result_file.read_text(encoding="utf-8"))
    docx_out = generate_docx_report(results, REPORTS_DIR / "QA_최종_테스트_결과_보고서.docx")
    pdf_out = generate_pdf_report(results, REPORTS_DIR / "QA_최종_테스트_결과_보고서.pdf")
    print(f"[Success] DOCX 생성: {docx_out}")
    print(f"[Success] PDF 생성: {pdf_out}")
